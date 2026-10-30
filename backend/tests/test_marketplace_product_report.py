from io import BytesIO

from openpyxl import load_workbook

from app.db import SessionLocal
from app.models import Account, MarketplaceProductMapping
from .conftest import buy, sell


def test_product_report_consolidates_sales_stock_and_excel_across_marketplaces(client):
    with SessionLocal() as db:
        db.add(Account(id=4, company_id=1, name="WB shop", channel="Wildberries"))
        db.flush()
        db.add_all([
            MarketplaceProductMapping(company_id=1, account_id=1, channel="Kaspi",
                                      external_key="kaspi-card-1", product_id=1),
            MarketplaceProductMapping(company_id=1, account_id=4, channel="Wildberries",
                                      external_key="wb-card-1", product_id=1),
        ])
        db.commit()

    buy(client, ref="P-REPORT", qty=10, cost=1000)
    sell(client, ref="S-KASPI-REPORT", qty=2, price=3000, account=1)
    sell(client, ref="S-WB-REPORT", qty=3, price=2500, account=4)

    saved = client.put("/api/marketplace-product-report/settings", json={
        "company_id": 1, "product_id": 1,
        "unit_cost_tenge": 1000, "tax_rate_percent": 10,
    })
    assert saved.status_code == 200, saved.text

    response = client.get("/api/marketplace-product-report", params={
        "start": "2026-10-01", "end": "2026-10-31", "company_id": 1,
    })
    assert response.status_code == 200, response.text
    report = response.json()
    assert report["total"] == 1
    item = report["items"][0]
    assert item["sku"] == "A"
    assert item["sold_units"] == 5
    assert item["returned_units"] == 0
    assert item["available_stock"] == 5
    assert item["marketplace_card_count"] == 2
    assert {part["channel"] for part in item["marketplace_breakdown"]} == {"Kaspi", "Wildberries"}
    assert item["unit_cost_tenge"] == 1000
    assert item["tax_rate_percent"] == 10
    assert item["cogs"] == 500000
    assert item["tax"] == 1350

    workbook_response = client.get("/api/marketplace-product-report.xlsx", params={
        "start": "2026-10-01", "end": "2026-10-31", "company_id": 1,
    })
    assert workbook_response.status_code == 200, workbook_response.text
    workbook = load_workbook(BytesIO(workbook_response.content), data_only=False)
    assert workbook.sheetnames == ["Отчёт по товарам", "Разрез по магазинам"]
    assert workbook["Отчёт по товарам"].max_row == 5
    assert workbook["Разрез по магазинам"].max_row == 3
