from datetime import date, timedelta
from io import BytesIO

from openpyxl import load_workbook

from app import wb_weekly
from app.db import SessionLocal
from app.models import Account, Batch, MarketplaceRecord, MarketplaceProductMapping, MarketplaceProductCost, Product, Return, Sale, User, now


def test_period_fetch_uses_bounded_week_and_rrd_cursor(monkeypatch):
    calls = []
    monkeypatch.setattr(wb_weekly.time, "sleep", lambda _: None)
    def fake_request(url, token, payload):
        calls.append((url, payload))
        if payload["rrdId"] == 0:
            return 200, [{"rrdId": 8, "reportId": 2}]
        return 204, {}
    monkeypatch.setattr(wb_weekly, "_request", fake_request)
    rows = wb_weekly.fetch_week("private", date(2026, 9, 28), date(2026, 10, 4))
    assert [r["rrdId"] for r in rows] == [8]
    assert len(calls) == 2
    assert all(c[1]["period"] == "weekly" and c[1]["dateFrom"] == "2026-09-28" and c[1]["dateTo"] == "2026-10-04" for c in calls)
    assert calls[1][1]["rrdId"] == 8


def test_fetch_daily_period_and_month_validation(monkeypatch):
    calls = []
    monkeypatch.setattr(wb_weekly.time, "sleep", lambda _: None)
    def fake_request(url, token, payload):
        calls.append(payload)
        return (200, [{"rrdId": 101}]) if payload["rrdId"] == 0 else (204, {})
    monkeypatch.setattr(wb_weekly, "_request", fake_request)
    rows = wb_weekly.fetch_period("private", date(2026, 9, 1), date(2026, 9, 30), "daily")
    assert [r["rrdId"] for r in rows] == [101]
    assert len(calls) == 2
    assert all(p["period"] == "daily" and p["dateFrom"] == "2026-09-01" and p["dateTo"] == "2026-09-30" for p in calls)
    assert wb_weekly._validate_report_period("monthly", date(2026, 9, 1), date(2026, 9, 30)) == "daily"
    assert wb_weekly._validate_report_period("daily", date(2026, 10, 10), date(2026, 10, 10)) == "daily"
    today = wb_weekly.finance_today()
    if today.day > 1:
        assert wb_weekly._validate_report_period("monthly", today.replace(day=1), today - timedelta(days=1)) == "daily"


def test_profit_is_based_on_payout_and_deducts_costs_once():
    assert wb_weekly.payout_profit("700", "400", "50", "10", "20", "5", "15") == 200
    # The independent commission/acquiring fields and retail/SPP fields are intentionally not inputs.


def test_monthly_finance_sync_uses_daily_api_and_separate_period_storage(client, monkeypatch):
    from app.models import Account
    calls = []
    def fake_fetch(token, start, end, api_period):
        calls.append((start, end, api_period))
        return [{"rrdId": 501, "currency": "KZT", "vendorCode": "M-1", "nmId": 501,
                 "docTypeName": "Продажа", "quantity": 2, "forPay": "1200", "rrDt": "2026-09-14"}]
    monkeypatch.setattr(wb_weekly, "fetch_period", fake_fetch)
    with SessionLocal() as db:
        account = Account(company_id=1, name="WB monthly", channel="Wildberries", mode="api")
        db.add(account); db.commit(); db.refresh(account)
        state = wb_weekly.sync_period(db, account, "token", "monthly", date(2026, 9, 1), date(2026, 9, 30))
        report = wb_weekly.build_report(db, account, date(2026, 9, 1), date(2026, 9, 30), period="monthly")
        assert calls == [(date(2026, 9, 1), date(2026, 9, 30), "daily")]
        assert state["status"] == report["status"] == "готов"
        assert report["period"] == "monthly"
        assert report["totals"]["payout"] == "1200"
        # Same account/date may also have a separate daily or weekly report without key collision.
        assert not db.query(MarketplaceRecord).filter(
            MarketplaceRecord.account_id == account.id,
            MarketplaceRecord.source == wb_weekly.SOURCE,
            MarketplaceRecord.external_id == "row:2026-09-01:501",
        ).first()


def test_xlsx_totals_use_full_report_including_unallocated_operations():
    report = {
        "period": "monthly", "start": "2026-09-01", "end": "2026-09-30",
        "week_start": "2026-09-01", "week_end": "2026-09-30", "currency": "KZT",
        "row_count": 2, "status": "готов", "sales_qty": 1, "returns_qty": 0,
        "unallocated_operations": 1, "cogs_complete": False,
        "totals": {"retail":"0", "realization":"100", "payout":"90", "commission":"10",
                   "wb_reward":"10", "commission_vat":"0", "acquiring":"0",
                   "direct_logistics":"15", "reverse_logistics":"0", "penalty":"0",
                   "storage":"0", "other":"0", "cogs":None, "profit":None},
        "products": [{"article":"P-1", "name":"Товар", "sku":"", "currency":"KZT",
                      "sales_qty":1,"returns_qty":0,"retail":"0","realization":"100",
                      "payout":"90","commission":"10","wb_reward":"10","commission_vat":"0",
                      "acquiring":"0","direct_logistics":"0","reverse_logistics":"0","penalty":"0",
                      "storage":"0","other":"0","cogs":None,"profit":None}],
        "operations": [],
    }
    book = load_workbook(BytesIO(wb_weekly.create_xlsx(report)), data_only=True)
    assert book.sheetnames == ["Финансовый итог", "Продажи по товарам", "Себестоимость", "Расходы Wildberries", "Возвраты", "Прибыль и маржинальность", "Сверка с Wildberries"]
    assert book["Финансовый итог"]["B5"].value == 90
    assert book["Расходы Wildberries"]["B6"].value == 15  # direct logistics row; standalone delivery remains report-level
    book.close()


def test_weekly_sync_is_idempotent_and_xlsx_is_real_workbook(client, monkeypatch):
    start, end = date(2026, 9, 28), date(2026, 10, 4)
    rows = [{"rrdId": 901, "reportId": 42, "dateFrom": str(start), "dateTo": str(end),
        "currency": "KZT", "nmId": 123, "vendorCode": "ABC-1", "sku": "4870000000001", "title": "Товар",
        "docTypeName": "Продажа", "quantity": 2, "retailPriceWithDisc": "1000.00", "retailAmount": "950.00", "forPay": "700.00",
        "commissionAmount": "100.00", "ppvzSalesCommission": "12", "deliveryService": "50.00", "penalty": "10.00", "paidStorage": "5.00",
        "acquiringFee": "20.00", "rrDt": "2026-10-01"}]
    monkeypatch.setattr(wb_weekly, "fetch_week", lambda *args: rows)
    with SessionLocal() as db:
        account = Account(company_id=1, name="WB test", channel="Wildberries", mode="api")
        db.add(account); db.commit(); db.refresh(account)
        first = wb_weekly.sync_week(db, account, "private", start, end)
        second = wb_weekly.sync_week(db, account, "private", start, end)
        assert first["status"] == second["status"] == "готов"
        stored = db.query(MarketplaceRecord).filter(MarketplaceRecord.account_id == account.id,
            MarketplaceRecord.source == wb_weekly.SOURCE, MarketplaceRecord.external_id.like(f"row:{start}:%")).all()
        assert len(stored) == 1
        report = wb_weekly.build_report(db, account, start, end)
        product = report["products"][0]
        assert product["payout"] == "700.00"
        assert product["direct_logistics"] is None  # delivery without product identity stays report-level
        assert product["profit"] is None  # missing actual FIFO cost; no invented profit
        xlsx = wb_weekly.create_xlsx(report)
        book = load_workbook(BytesIO(xlsx), data_only=False)
        assert book.sheetnames == ["Финансовый итог", "Продажи по товарам", "Себестоимость", "Расходы Wildberries", "Возвраты", "Прибыль и маржинальность", "Сверка с Wildberries"]
        assert book["Финансовый итог"]["B5"].value == 700.0
        assert book["Себестоимость"]["F2"].value == "Не заполнена/проверить возврат"
        assert book["Сверка с Wildberries"]["B2"].value == 700.0
        book.close()


def test_wb_commission_amount_is_currency_but_percent_is_not(client, monkeypatch):
    from app import wb_weekly
    from app.db import SessionLocal
    from app.models import Account
    start, end = date(2026, 9, 28), date(2026, 10, 4)
    monkeypatch.setattr(wb_weekly, "fetch_week", lambda *args: [{
        "rrdId": 1901, "currency": "KZT", "nmId": 1001, "docTypeName": "Продажа",
        "quantity": 1, "forPay": "880", "ppvzSalesCommission": "23.74", "commissionPercent": 24,
        "vw": "23.74", "vwNds": "4.45",
    }])
    with SessionLocal() as db:
        account = Account(company_id=1, name="WB percent", channel="Wildberries", mode="api")
        db.add(account); db.commit(); db.refresh(account)
        wb_weekly.sync_week(db, account, "token", start, end)
        report = wb_weekly.build_report(db, account, start, end)
        assert report["totals"]["commission"] == "23.74"
        assert report["totals"]["wb_reward"] == "23.74"
        assert report["totals"]["commission_vat"] == "4.45"
        assert report["totals"]["payout"] == "880"


def test_weekly_retail_does_not_invent_a_unit_for_zero_quantity(client, monkeypatch):
    start, end = date(2026, 9, 28), date(2026, 10, 4)
    monkeypatch.setattr(wb_weekly, "fetch_week", lambda *args: [
        {"rrdId": 2901, "currency": "KZT", "vendorCode": "A", "docTypeName": "Логистика",
         "quantity": 0, "retailPriceWithDisc": "999", "deliveryRub": "15"},
        {"rrdId": 2902, "currency": "KZT", "vendorCode": "A", "docTypeName": "Возврат",
         "quantity": 1, "retailPriceWithDisc": "100"},
    ])
    with SessionLocal() as db:
        account = Account(company_id=1, name="WB zero quantity", channel="Wildberries", mode="api")
        db.add(account); db.commit(); db.refresh(account)
        wb_weekly.sync_week(db, account, "token", start, end)
        report = wb_weekly.build_report(db, account, start, end)
        assert report["totals"]["retail"] == "-100"
        assert report["totals"]["direct_logistics"] == "15"


def test_productless_wb_costs_remain_report_level(client, monkeypatch):
    start, end = date(2026, 9, 28), date(2026, 10, 4)
    monkeypatch.setattr(wb_weekly, "fetch_week", lambda *args: [
        {"rrdId": 3901, "currency": "KZT", "vendorCode": "A-1", "nmId": 1001,
         "docTypeName": "Продажа", "quantity": 1, "forPay": "800"},
        {"rrdId": 3902, "currency": "KZT", "docTypeName": "Логистика", "quantity": 0,
         "deliveryService": "50"},
    ])
    with SessionLocal() as db:
        account = Account(company_id=1, name="WB unallocated cost", channel="Wildberries", mode="api")
        db.add(account); db.commit(); db.refresh(account)
        wb_weekly.sync_week(db, account, "token", start, end)
        report = wb_weekly.build_report(db, account, start, end)
        assert len(report["products"]) == 1
        assert report["products"][0]["article"] == "A-1"
        assert report["totals"]["direct_logistics"] == "50"
        assert report["unallocated_operations"] == 1
        assert report["totals"]["profit"] is None


def test_fifo_cost_matches_wb_sale_and_return_ids_not_settlement_dates(client, monkeypatch):
    from app.marketplace_pipeline import reference as wb_reference
    start, end = date(2026, 9, 28), date(2026, 10, 4)
    monkeypatch.setattr(wb_weekly, "fetch_week", lambda *args: [
        {"rrdId": 4901, "currency": "KZT", "vendorCode": "A-1", "nmId": 1001, "srid": "sale-77",
         "docTypeName": "Продажа", "quantity": 1, "forPay": "900", "ppvzSalesCommission": "80", "rrDt": "2026-10-01"},
        {"rrdId": 4902, "currency": "KZT", "vendorCode": "A-1", "nmId": 1001, "srid": "sale-77",
         "docTypeName": "Возврат", "quantity": 1, "forPay": "-900", "rrDt": "2026-10-02"},
        {"rrdId": 4903, "currency": "KZT", "docTypeName": "Логистика", "quantity": 0, "deliveryService": "20"},
    ])
    with SessionLocal() as db:
        account = Account(company_id=1, name="WB identity COGS", channel="Wildberries", mode="api")
        db.add(account); db.flush()
        db.add(MarketplaceProductMapping(company_id=1, account_id=account.id, channel="Wildberries",
            external_key="product:1001", product_id=1))
        sale = Sale(company_id=1, account_id=account.id, product_id=1,
            reference=wb_reference(account, "wb-sale:sale-77"), date=date(2026, 9, 20), quantity=1,
            revenue=90000, discount=0, cogs=40000, cogs_pending=False, commission=0, logistics=0, advertising=0)
        db.add(sale); db.flush()
        db.add(Return(company_id=1, sale_id=sale.id,
            reference=wb_reference(account, "realization:4902", "return"), date=date(2026, 10, 2),
            quantity=1, revenue=90000, cogs=40000, inventory_status="restocked"))
        db.commit()
        wb_weekly.sync_week(db, account, "token", start, end)
        report = wb_weekly.build_report(db, account, start, end)
        assert report["cogs_complete"] is True
        assert report["totals"]["payout"] == "0"
        assert report["totals"]["cogs"] == "0"
        assert report["totals"]["profit"] == "-20"


def test_selected_finance_report_projects_sale_for_fifo_profit(client, monkeypatch):
    start, end = date(2026, 9, 28), date(2026, 10, 4)
    monkeypatch.setattr(wb_weekly, "fetch_week", lambda *args: [{
        "rrdId": 5901, "reportId": 42, "currencyName": "KZT", "vendorCode": "ERP-A",
        "nmId": 2001, "srid": "finance-sale-59", "docTypeName": "Продажа", "quantity": 2,
        "retailAmount": "2000", "forPay": "1500", "ppvzSalesCommission": "100",
        "deliveryService": "100", "saleDt": "2026-10-01", "rrDt": "2026-10-03",
    }])
    with SessionLocal() as db:
        user = db.get(User, 1)
        account = Account(company_id=1, name="WB Finance projection", channel="Wildberries", mode="api")
        db.add(account); db.flush()
        db.add(MarketplaceProductMapping(company_id=1, account_id=account.id, channel="Wildberries",
            external_key="product:2001", product_id=1))
        db.add(Batch(company_id=1, product_id=1, purchase_id=None, date=date(2026, 9, 1),
            quantity=2, remaining=2, total_cost=80000, remaining_cost=80000))
        db.commit()
        state = wb_weekly.sync_week(db, account, "token", start, end, project_user=user)
        report = wb_weekly.build_report(db, account, start, end)
        posted = db.query(Sale).filter(Sale.account_id == account.id).one()
        assert state["ledger_projection"] == {"rows": 1, "projected": 1, "pending": 0}
        assert posted.reference.endswith("wb-sale:finance-sale-59")
        assert report["cogs_complete"] is True
        assert report["totals"]["payout"] == "1500"
        assert report["totals"]["cogs"] == "800"
        assert report["totals"]["profit"] == "600"


def test_dated_cost_shows_each_week_and_late_return_reduces_next_week(client, monkeypatch):
    week1, end1 = date(2026, 9, 21), date(2026, 9, 27)
    week2, end2 = date(2026, 9, 28), date(2026, 10, 4)
    reports = {
        week1: [
            {"rrdId":6901,"currencyName":"KZT","vendorCode":"COST-1","nmId":7001,"srid":"cost-sale-1","docTypeName":"Продажа","quantity":1,"forPay":"6000","saleDt":"2026-09-22"},
            {"rrdId":6902,"currencyName":"KZT","vendorCode":"COST-1","nmId":7001,"srid":"cost-sale-2","docTypeName":"Продажа","quantity":1,"forPay":"6000","saleDt":"2026-09-23"},
        ],
        week2: [
            {"rrdId":6903,"currencyName":"KZT","vendorCode":"COST-1","nmId":7001,"srid":"cost-sale-1","docTypeName":"Возврат","quantity":1,"forPay":"-5000","rrDt":"2026-09-29"},
        ],
    }
    monkeypatch.setattr(wb_weekly, "fetch_week", lambda token, start, end: reports[start])
    with SessionLocal() as db:
        account=Account(company_id=1,name="WB cost history",channel="Wildberries",mode="api")
        product=Product(sku="WB-COST-HISTORY",name="Cost history item")
        db.add_all([account,product]);db.flush()
        db.add(MarketplaceProductMapping(company_id=1,account_id=account.id,channel="Wildberries",external_key="product:7001",product_id=product.id))
        db.add_all([
            MarketplaceProductCost(company_id=1,account_id=account.id,external_key="product:7001",nm_id="7001",seller_article="COST-1",unit_cost=500000,currency="KZT",effective_from=date(2026,9,1)),
            MarketplaceProductCost(company_id=1,account_id=account.id,external_key="product:7001",nm_id="7001",seller_article="COST-1",unit_cost=700000,currency="KZT",effective_from=date(2026,9,23)),
        ])
        db.commit()
        wb_weekly.sync_week(db,account,"token",week1,end1)
        first=wb_weekly.build_report(db,account,week1,end1)
        assert first["cogs_complete"] is True
        assert first["returns_qty"] == 0
        assert first["totals"]["cogs"] == "12000"
        assert first["totals"]["profit"] == "0"
        wb_weekly.sync_week(db,account,"token",week2,end2)
        returned=wb_weekly.build_report(db,account,week2,end2)
        assert returned["returns_qty"] == 1
        assert returned["cogs_complete"] is True
        assert returned["totals"]["cogs"] == "0"
        assert returned["totals"]["profit"] == "-5000"


