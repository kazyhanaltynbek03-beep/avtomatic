from datetime import date
from types import SimpleNamespace
from urllib.parse import urlparse, parse_qs
import pytest
from sqlalchemy import select,func
from app import connectors,canonical,importing,marketplace_pipeline as mp,schemas as s,finance as f
from app.models import Account,User,MarketplaceRecord,Sale,Return,CanonicalEvent,Batch,Movement,Expense
from app.db import SessionLocal
from .conftest import buy


def staged(db,acc,external,**payload):
    row=MarketplaceRecord(account_id=acc.id,company_id=acc.company_id,source=acc.channel.lower(),external_id=external,payload=payload)
    db.add(row);db.flush();return row


def test_kaspi_period_projection_uses_completion_and_return_dates(client):
    with SessionLocal() as db:
        acc = db.get(Account, 1)
        user = db.get(User, 1)
        row = staged(db, acc, 'CROSS-MONTH', record_type='order_entry',
                     order_status='COMPLETED', sku='A', quantity=1, total_price='1000',
                     order_creation_date='2026-09-30', completed_date='2026-10-02')
        result = mp.project_account(db, user, acc, date_from=date(2026, 10, 1), date_to=date(2026, 10, 31))
        sale = db.scalar(select(Sale).where(Sale.reference == mp.reference(acc, row.external_id)))
        assert sale is not None, result
        assert sale.date == date(2026, 10, 2)
        row.payload = {**row.payload, 'order_status': 'RETURNED', 'returned_date': '2026-11-03'}
        result = mp.project_account(db, user, acc, date_from=date(2026, 11, 1), date_to=date(2026, 11, 30))
        returned = db.scalar(select(Return).where(Return.sale_id == sale.id))
        assert returned is not None, result
        assert returned.date == date(2026, 11, 3)
        assert db.scalar(select(func.count()).select_from(Sale)) == 1


def test_marketplace_summary_rejects_inverted_period(client):
    response = client.get('/api/marketplace-summary', params={
        'channel': 'Kaspi', 'start': '2026-10-31', 'end': '2026-10-01',
    })
    assert response.status_code == 422


@pytest.mark.parametrize('channel',['Kaspi','Wildberries'])
def test_marketplace_summary_can_filter_one_store_without_changing_all_store_total(client,channel):
    with SessionLocal() as db:
        primary=db.get(Account,1); primary.channel=channel
        second=Account(id=4,company_id=1,name=f'{channel} второй магазин',channel=channel)
        db.add(second); db.flush()
        for account_id,reference,revenue in ((1,'store-one',10000),(4,'store-two',25000)):
            db.add(Sale(company_id=1,account_id=account_id,product_id=1,reference=reference,
                        date=date(2026,10,25),quantity=1,revenue=revenue,cogs=0,
                        cogs_pending=False,commission=0,logistics=0,advertising=0))
        db.commit()
    params={'channel':channel,'start':'2026-10-25','end':'2026-10-25'}
    all_stores=client.get('/api/marketplace-summary',params=params)
    one_store=client.get('/api/marketplace-summary',params={**params,'account_id':4})
    wrong_channel=client.get('/api/marketplace-summary',params={**params,'account_id':3})
    assert all_stores.status_code==200,all_stores.text
    assert one_store.status_code==200,one_store.text
    assert all_stores.json()['revenue']==35000
    assert one_store.json()['revenue']==25000
    assert wrong_channel.status_code==404


def test_kaspi_created_orders_filter_by_store(client):
    with SessionLocal() as db:
        first=db.get(Account,1); first.sync_status='done'
        second=Account(id=4,company_id=1,name='Kaspi второй магазин',channel='Kaspi',sync_status='done')
        db.add(second); db.flush()
        staged(db,first,'order:K-1',record_type='order',order_id='id-1',code='K-1',
               status='COMPLETED',creation_date='2026-10-25',total_price='1000')
        staged(db,second,'order:K-2',record_type='order',order_id='id-2',code='K-2',
               status='COMPLETED',creation_date='2026-10-25',total_price='2000')
        db.commit()
    response=client.get('/api/kaspi/orders-created',params={
        'start':'2026-10-25','end':'2026-10-25','account_id':4,
    })
    assert response.status_code==200,response.text
    data=response.json()
    assert data['total']==1
    assert data['items'][0]['account_id']==4


def test_kaspi_orders_created_report_filters_by_almaty_creation_date(client):
    with SessionLocal() as db:
        acc=db.get(Account,1); acc.channel='Kaspi'; acc.sync_status='done'
        staged(db,acc,'order:K-1',record_type='order',order_id='id-1',code='K-1',
               status='COMPLETED',creation_date='2026-10-25T00:00:00+05:00',total_price='1200')
        staged(db,acc,'entry:id-1:e-1',record_type='order_entry',order_id='id-1',
               order_status='COMPLETED',quantity=2,total_price='1200')
        db.commit()
    response=client.get('/api/kaspi/orders-created',params={'start':'2026-10-25','end':'2026-10-25'})
    assert response.status_code==200,response.text
    data=response.json()
    assert data['api_data_available'] is True
    assert data['total']==1
    assert data['items'][0]['created_date']=='2026-10-25'
    assert data['items'][0]['units']==2


def test_kaspi_summary_does_not_use_creation_date_as_issue_date(client, monkeypatch):
    monkeypatch.delenv('KASPI_COMMISSION_RATE',raising=False)
    with SessionLocal() as db:
        acc=db.get(Account,1); acc.channel='Kaspi'
        staged(db,acc,'creation-only',record_type='order_entry',order_status='COMPLETED',
               sku='A',quantity=1,total_price='1000',order_creation_date='2026-10-25')
        db.commit()
    response=client.get('/api/marketplace-summary',params={'channel':'Kaspi','start':'2026-10-25','end':'2026-10-25'})
    assert response.status_code==200,response.text
    data=response.json()
    assert data['api_revenue'] is None
    assert data['api_units'] is None
    assert data['commission_estimated']==0
    assert data['commission_rate'] is None


def test_report_skips_ignored_non_financial_staging_rows(client):
    with SessionLocal() as db:
        acc=db.get(Account,1);u=db.get(User,1)
        ignored=staged(db,acc,'ignored-confirmed-sale',record_type='sale',saleDt='2026-10-02',quantity=1)
        ignored.stage_status='ignored_non_financial'
        ignored.reason_code='NOT_FINANCIAL_OR_NOT_CONFIRMED'
        staged(db,acc,'received-confirmed-sale',record_type='sale',saleDt='2026-10-02',quantity=1)
        db.commit()

        report=f.report(db,u,date(2026,10,1),date(2026,10,31),1)

        assert report['pending_source_events']==1
        assert report['profit_complete'] is False


@pytest.mark.parametrize('channel',['Kaspi','Wildberries','Ozon'])
def test_normalized_api_sale_three_replays_metrics(client,channel):
    buy(client,qty=10,cost=350000)
    with SessionLocal() as db:
        acc=db.get(Account,1);acc.channel=channel;u=db.get(User,1)
        payload={
          'Kaspi':dict(record_type='order_entry',order_status='COMPLETED',sku='A',quantity=10,total_price='100000',order_creation_date='2026-10-02',completed_date='2026-10-02'),
          'Wildberries':dict(record_type='realization',sellerOperName='Продажа',vendorCode='A',nmId=42,sku='barcode',quantity=10,retailPriceWithDisc='10000',currency='KZT',saleDt='2026-10-02',rrDate='2026-10-02',srid='sale42',retailAmount='100000',ppvzForPay='84000',ppvzSalesCommission='12',commissionPercent=24,deliveryService='4000'),
          'Ozon':dict(record_type='posting',status='delivered',shipment_date='2026-10-02',products=[dict(offer_id='A',sku=42,quantity=10,price='10000',currency_code='KZT')])}[channel]
        row=staged(db,acc,'SRC1',**payload)
        if channel!='Wildberries':
            for name,amt in [('Marketplace commission',12000),('Marketplace logistics',4000)]:
                canonical.post_event(db,u,'expense',s.ExpenseIn(company_id=1,account_id=1,reference=name,date=date(2026,10,2),category=name,amount=amt*100))
        for _ in range(3):
            result=mp.project_account(db,u,acc);assert result['pending_erp']==0,result;db.commit()
        report=f.report(db,u,date(2026,10,1),date(2026,10,31),1)
        for k,v in dict(revenue=10000000,cogs=3500000,commission=1200 if channel=='Wildberries' else 1200000,logistics=400000,net_profit=6098800 if channel=='Wildberries' else 4900000,cash_flow=0).items():assert report[k]==v,(k,report[k])
        assert db.scalar(select(func.count()).select_from(Sale))==1
        assert db.scalar(select(func.sum(Batch.remaining)))==0
        assert row.status=='projected'


def test_kaspi_commission_estimate_is_in_central_report_and_detail(client, monkeypatch):
    monkeypatch.setenv('KASPI_COMMISSION_RATE','0.12')
    buy(client,qty=10,cost=350000)
    with SessionLocal() as db:
        acc=db.get(Account,1);acc.channel='Kaspi';u=db.get(User,1)
        staged(db,acc,'KASPI-COMMISSION',record_type='order_entry',order_status='COMPLETED',
               sku='A',quantity=10,total_price='100000',completed_date='2026-10-02')
        result=mp.project_account(db,u,acc);assert result['pending_erp']==0,result;db.commit()
        report=f.report(db,u,date(2026,10,1),date(2026,10,31),1)
        assert report['revenue']==10000000
        assert report['commission_actual']==0
        assert report['commission_estimated']==1200000
        assert report['commission']==1200000
        assert report['commission_rate']=='0.12'
        assert report['net_profit']==5300000
        detail=f.report_details(db,u,date(2026,10,1),date(2026,10,31),'commission',1)
        assert detail['expected']==-1200000
        assert detail['explained']==-1200000
        assert detail['rows'][0]['type']=='Расчётная комиссия Kaspi'
        assert report['cash_flow']==0


def test_kaspi_estimate_respects_explicit_commission_without_double_count(client):
    buy(client,qty=10,cost=350000)
    with SessionLocal() as db:
        acc=db.get(Account,1);acc.channel='Kaspi';u=db.get(User,1)
        staged(db,acc,'KASPI-EXPLICIT',record_type='order_entry',order_status='COMPLETED',
               sku='A',quantity=10,total_price='100000',completed_date='2026-10-02')
        canonical.post_event(db,u,'expense',s.ExpenseIn(company_id=1,account_id=1,
            reference='Kaspi fee',date=date(2026,10,2),category='Marketplace commission',amount=1200000))
        result=mp.project_account(db,u,acc);assert result['pending_erp']==0,result;db.commit()
        report=f.report(db,u,date(2026,10,1),date(2026,10,31),1)
        assert report['commission_actual']==1200000
        assert report['commission_estimated']==0
        assert report['commission']==1200000


def test_mapping_fifo_retry(client):
    with SessionLocal() as db:
        acc=db.get(Account,1);acc.channel='Wildberries'
        staged(db,acc,'sale:S1',record_type='sale',supplierArticle='ALIAS',nmId=11,saleID='S1',srid='R1',priceWithDisc=1000,date='2026-10-02',currency='KZT');db.commit()
    assert client.post('/api/accounts/1/reprocess').json()['pending_erp']==1
    inbox=client.get('/api/marketplace-records?status=pending').json()['items'][0]
    assert inbox['stage_status']=='pending_fifo'
    with SessionLocal() as db:
        sale=db.scalar(select(Sale));pid=sale.product_id
        assert sale.revenue==100000 and sale.cogs_pending
    response=client.post('/api/post/opening',json=dict(company_id=1,product_id=pid,reference='OPEN-ALIAS',date='2026-10-01',quantity=2,unit_cost=1000,evidence='Supplier invoice'))
    assert response.status_code==200,response.text
    with SessionLocal() as db:
        sale=db.scalar(select(Sale));assert not sale.cogs_pending and sale.cogs==1000
        movements=db.scalars(select(Movement).where(Movement.product_id==pid)).all()
        assert sum(x.quantity for x in movements)==1
        assert sum(x.cost for x in movements)==1000
    assert client.post('/api/accounts/1/reprocess').json()['projected']==0


def test_ozon_delivering_not_sale_and_multiline_atomic(client):
    buy(client,qty=2,cost=1000)
    with SessionLocal() as db:
        acc=db.get(Account,1);acc.channel='Ozon';u=db.get(User,1)
        x=dict(offer_id='A',sku=11,price=100,quantity=1,currency_code='KZT')
        row=staged(db,acc,'posting:X',record_type='posting',status='delivering',shipment_date='2026-10-02',products=[x])
        assert not mp.project_record(db,u,acc,row)[0] and row.status=='received'
        row.payload={**row.payload,'status':'delivered','products':[x,dict(offer_id='UNKNOWN',sku=12,price=None,quantity=1,currency_code='KZT')]}
        assert not mp.project_record(db,u,acc,row)[0]
        assert db.scalar(select(func.count()).select_from(Sale))==0
        assert db.scalar(select(func.sum(Batch.remaining)))==2
        assert db.scalar(select(func.count()).select_from(CanonicalEvent).where(CanonicalEvent.event_type=='sale'))==0


def test_ozon_explicit_posting_commission_is_an_expense_not_payout(client):
    buy(client,qty=1,cost=1000)
    with SessionLocal() as db:
        acc=db.get(Account,1);acc.channel='Ozon';u=db.get(User,1)
        row=staged(db,acc,'posting:commission',record_type='posting',status='delivered',shipment_date='2026-10-02',
            products=[dict(offer_id='A',sku=11,quantity=1,price=10000,currency_code='KZT')],
            financial_data={'products':[{'product_id':11,'commission_amount':1200,'commissions_currency_code':'KZT','payout':8800}]})
        assert mp.project_record(db,u,acc,row)[0]
        report=f.report(db,u,date(2026,10,1),date(2026,10,31),1)
        assert report['revenue']==1000000 and report['commission']==120000
        assert report['cash_flow']==0 # Ozon payout is an accrual, never fabricated cash.


@pytest.mark.parametrize('channel', ['Kaspi', 'Ozon'])
def test_confirmed_marketplace_return_reverses_existing_sale_once(client, channel):
    buy(client, qty=1, cost=1000)
    with SessionLocal() as db:
        acc=db.get(Account,1); acc.channel=channel; u=db.get(User,1)
        if channel=='Kaspi':
            sale_payload=dict(record_type='order_entry',order_status='COMPLETED',sku='A',quantity=1,total_price=10000,order_creation_date='2026-10-02',completed_date='2026-10-02')
            return_payload={**sale_payload,'order_status':'RETURNED','returned_date':'2026-10-03'}
        else:
            sale_payload=dict(record_type='posting',status='delivered',shipment_date='2026-10-02',products=[dict(offer_id='A',sku=11,quantity=1,price=10000,currency_code='KZT')])
            return_payload={**sale_payload,'status':'returned','cancelled_at':'2026-10-03'}
        row=staged(db,acc,'RETURN-ONE',**sale_payload)
        assert mp.project_record(db,u,acc,row)[0]
        row.payload=return_payload
        assert mp.project_record(db,u,acc,row)[0]
        # Replaying the provider's return must remain an idempotent canonical event.
        assert mp.project_record(db,u,acc,row)[0]
        report=f.report(db,u,date(2026,10,1),date(2026,10,31),1)
        assert report['revenue']==0 and report['units']==0
        assert db.scalar(select(func.count()).select_from(Sale))==1
        ret=db.scalar(select(Return))
        assert ret.inventory_status == "requested"
        assert db.scalar(select(func.sum(Batch.remaining)))==0
        ret.inventory_status = "restocked"
        assert f.restock_return_inventory(db,u,ret)==1
        assert f.restock_return_inventory(db,u,ret)==0
        assert db.scalar(select(func.sum(Batch.remaining)))==1
        assert db.scalar(select(func.count()).select_from(Movement).where(Movement.kind=="return"))==1


def test_kaspi_summary_uses_only_issued_sales_and_separates_delivery_costs(client, monkeypatch):
    monkeypatch.setenv('KASPI_COMMISSION_RATE','0.12')
    with SessionLocal() as db:
        acc=db.get(Account,1)
        staged(db,acc,'entry:issued:1',record_type='order_entry',order_status='COMPLETED',
               sku='A',quantity=1,total_price=100000,order_creation_date='2026-10-02',completed_date='2026-10-02')
        staged(db,acc,'entry:cancelled:2',record_type='order_entry',order_status='CANCELLED',
               sku='A',quantity=1,total_price=50000,order_creation_date='2026-10-02',completed_date='2026-10-02')
        db.add_all([
            Expense(company_id=1,account_id=1,reference='seller-delivery',date=date(2026,10,2),
                    category='Marketplace logistics',amount=100000),
            Expense(company_id=1,account_id=1,reference='return-delivery',date=date(2026,10,3),
                    category='Marketplace return logistics',amount=50000),
        ])
        db.commit()
    result=client.get('/api/marketplace-summary',params={'channel':'Kaspi','start':'2026-10-01','end':'2026-10-31'})
    assert result.status_code==200,result.text
    data=result.json()
    assert data['api_revenue']==10000000  # only the issued 100,000 ₸ order
    assert data['commission_estimated']==1200000  # 12% of issued gross only
    assert data['logistics']==100000  # seller delivery only; buyer delivery is not an expense
    assert data['return_logistics']==50000
    assert data['kaspi_estimated_proceeds']==8650000  # 100,000 - 12,000 - 1,000 - 500


def test_fx_revision_and_return_replay(client, monkeypatch):
    monkeypatch.setattr(mp.exchange_rates, "fetch_nbk_rate", lambda *a, **k: (_ for _ in ()).throw(Exception("offline")))
    buy(client,qty=3,cost=1000)
    with SessionLocal() as db:
        acc=db.get(Account,1);acc.channel='Wildberries';u=db.get(User,1)
        row=staged(db,acc,'sale:S1',record_type='sale',saleID='S1',srid='RID',supplierArticle='A',priceWithDisc='100',date='2026-10-02')
        assert mp.project_record(db,u,acc,row)==(False,'FX_REQUIRED:RUB')
        acc.settlement_currency='RUB';acc.settlement_fx_rate=55000
        from app.models import AccountExchangeRate
        db.add(AccountExchangeRate(account_id=acc.id,currency='RUB',rate_date=date(2026,10,2),rate_scaled=55000,source='manual_confirmed'));db.flush()
        assert mp.project_record(db,u,acc,row)[0]
        assert db.scalar(select(Sale.revenue))==55000
        row.payload={**row.payload,'priceWithDisc':'200'}
        assert mp.project_record(db,u,acc,row)==(False,'EVENT_REVISION_REQUIRES_RECONCILIATION')
        assert db.scalar(select(Sale.revenue))==55000
        ret=staged(db,acc,'sale:R1',record_type='sale',saleID='R1',srid='RID',supplierArticle='A',date='2026-10-03')
        assert mp.project_record(db,u,acc,ret)[1]=='RETURN_STOCK_CONFIRMATION_REQUIRED'
        ret.payload={**ret.payload,'returned_to_warehouse':True}
        for _ in range(3):assert mp.project_record(db,u,acc,ret)[0]
        assert db.scalar(select(func.sum(Batch.remaining)))==2
        pending=db.scalar(select(Return))
        assert pending is not None and pending.inventory_status == "requested"
        assert f.report(db,u,date(2026,10,1),date(2026,10,31),1)['revenue']==0


def test_mapping_company_permissions(client):
    with SessionLocal() as db:
        row=staged(db,db.get(Account,2),'entry2',record_type='order_entry',sku='X');rid=row.id;db.commit()
    auth=client.post('/api/auth/login',json={'email':'warehouse@test.kz','password':'Test-password-1!'}).json();client.headers['X-CSRF-Token']=auth['csrf']
    assert not client.get('/api/marketplace-records?company_id=1').json()['items']
    assert client.post(f'/api/marketplace-records/{rid}/mapping',json={'product_id':1,'external_keys':['sku:X']}).status_code==403
    assert client.put('/api/accounts/2/settlement',json={'currency':'RUB','fx_rate':55000,'effective_date':'2026-10-02'}).status_code==403


def test_wb_pagination_sale_return_and_stock_identity(monkeypatch):
    c=connectors.WildberriesConnector(SimpleNamespace(channel='Wildberries',mode='api'),{'token':'fixture'})
    c.statistics_page_size=2;c.stocks_page_size=2;calls=[]
    def request(url,method='GET',payload=None,interval=0):
        calls.append((url,payload))
        if 'orders/new' in url:return {'orders':[]}
        if '/supplier/orders' in url:return []
        if '/supplier/sales' in url:
            start=parse_qs(urlparse(url).query)['dateFrom'][0]
            if 'T' not in start:return [dict(saleID='S1',srid='SAME',lastChangeDate='2026-10-01T00:00:00'),dict(saleID='R1',srid='SAME',lastChangeDate='2026-10-02T00:00:00')]
            return [dict(saleID='S2',srid='OTHER',lastChangeDate='2026-10-03T00:00:00')]
        if 'sales-reports/detailed' in url:return [dict(rrdId=10,srid='x')] if payload['rrdId']==0 else {}
        if 'stocks-report/products/products' in url:return {'data':{'items':[dict(nmId=1,stockCount=9)]}}
        if 'get/cards/list' in url:return {'cards':[],'cursor':{'total':0}}
        raise AssertionError(url)
    monkeypatch.setattr(c,'_request',request);r=c.fetch()
    assert {'sale:S1','sale:R1','sale:S2','stock:1','realization:10'}<={x['external_id'] for x in r.records}
    assert not r.warnings and any(p and p.get('rrdId')==10 for _,p in calls)


def test_kaspi_window_pagination_and_merchant_sku(monkeypatch):
    calls=[]
    def request(url,headers,**kwargs):
        calls.append(url)
        if '/orderentries/' in url:return 200,{'data':{'id':'master1','attributes':{'code':'market-code'}}}
        if '/masterproducts/' in url:return 200,{'data':{'attributes':{'code':'A','name':'Alpha'}}}
        if '/entries?' in url:return 200,{'data':[{'id':'entry1','attributes':{'quantity':2,'totalPrice':100}}]}
        q=parse_qs(urlparse(url).query)
        assert int(q['filter[orders][creationDate][$le]'][0])-int(q['filter[orders][creationDate][$ge]'][0])<14*86400000
        return 200,{'data':[{'id':'order'+q['page[number]'][0],'attributes':{'status':'COMPLETED','creationDate':1}}],'meta':{'pageCount':2}}
    monkeypatch.setattr(connectors,'_json_request',request)
    connector=connectors.KaspiConnector(SimpleNamespace(channel='Kaspi',mode='api'),{'token':'x'})
    first=connector.fetch()
    r=connector.fetch(first.next_cursor)
    assert any('page%5Bnumber%5D=1' in x for x in calls)
    assert all(x['sku']=='A' for x in r.records if x['record_type']=='order_entry')
    assert not r.warnings


def test_xlsx_sheet_header_csv_preamble():
    from io import BytesIO
    from openpyxl import Workbook
    w=Workbook();w.active.title='Instructions';w.active.append(['Read me']);ws=w.create_sheet('Sales')
    ws.append(['Отчёт']);ws.append([]);ws.append(['Артикул','Дата','Количество','Цена']);ws.append(['A','02.10.2026',1,'1\u00a0000,50'])
    out=BytesIO();w.save(out);meta={}
    headers,rows=importing.parse_file(out.getvalue(),'report.xlsx',sheet='Sales',metadata=meta)
    assert meta['header_row']==3 and importing._money(rows[0]['Цена'])==100050
    assert importing._date(rows[0]['Дата'])=='2026-10-02'
    headers,rows=importing.parse_file('Отчёт\n\nАртикул;Дата;Количество;Цена\nA;02.10.2026;1;1000\n\n'.encode(),'report.csv')
    assert headers==['Артикул','Дата','Количество','Цена'] and len(rows)==1


@pytest.mark.parametrize('source',['manual','xlsx','csv','Kaspi','Wildberries','Ozon'])
def test_user_golden_31000_all_sources(client,source):
    """Native sale fixtures + equivalent normalized returns/expenses, integer tiyn."""
    from io import BytesIO
    from openpyxl import Workbook
    with SessionLocal() as db:
        u=db.get(User,1);acc=db.get(Account,1)
        acc.channel=source if source in {'Kaspi','Wildberries','Ozon'} else 'Offline'
        canonical.post_event(db,u,'purchase',s.PurchaseIn(company_id=1,reference='OPEN',supplier='Supplier',
            date=date(2026,10,1),allocation_method='quantity',extra_cost=3888889,lines=[s.PurchaseLine(product_id=1,quantity=10,unit_cost=0)]))
        # The 1/10 return leaves exactly 35,000 KZT FIFO cost after rounding.
        if source in {'xlsx','csv'}:
            headers=['reference','date','company_id','account_id','product_id','quantity','unit_price']
            values=['GOLD','2026-10-02','1','1','1','10','10000']
            if source=='xlsx':
                w=Workbook();w.active.append(headers);w.active.append(values);out=BytesIO();w.save(out);raw=out.getvalue()
            else:raw=(','.join(headers)+'\n'+','.join(values)+'\n').encode()
            cols,rows=importing.parse_file(raw,'gold.'+source)
            run,_=importing.preview(db,u,'sales',rows,{k:k for k in cols});assert not run.errors,run.errors
            importing.confirm(db,u,run)
        elif source=='manual':
            canonical.post_event(db,u,'sale',s.SaleIn(company_id=1,account_id=1,product_id=1,reference='GOLD',date=date(2026,10,2),quantity=10,unit_price=1000000))
        else:
            payload={
                'Kaspi':dict(record_type='order_entry',order_status='COMPLETED',sku='A',quantity=10,total_price=100000,order_creation_date='2026-10-02',completed_date='2026-10-02'),
                'Wildberries':dict(record_type='realization',sellerOperName='Продажа',vendorCode='A',quantity=10,retailPriceWithDisc=10000,currency='KZT',saleDt='2026-10-02',srid='GOLD'),
                'Ozon':dict(record_type='posting',status='delivered',shipment_date='2026-10-02',products=[dict(offer_id='A',sku=42,quantity=10,price=10000,currency_code='KZT')])}[source]
            row=staged(db,acc,'GOLD',**payload);assert mp.project_record(db,u,acc,row)[0],row.payload
        sale=db.scalar(select(Sale))
        canonical.post_event(db,u,'return',s.ReturnIn(sale_id=sale.id,reference='RETURN',date=date(2026,10,3),quantity=1))
        ret=db.scalar(select(Return).where(Return.reference == 'RETURN'))
        ret.inventory_status='restocked'
        f.restock_return_inventory(db,u,ret)
        for category,amount in [('Marketplace commission',12000),('Marketplace logistics',4000),('Реклама',5000),('Other',3000)]:
            canonical.post_event(db,u,'expense',s.ExpenseIn(company_id=1,account_id=1,reference=category,date=date(2026,10,3),category=category,amount=amount*100))
        report=f.report(db,u,date(2026,10,1),date(2026,10,31),1)
        for key,value in dict(revenue=9000000,returns=1000000,cogs=3500000,commission=1200000,logistics=400000,advertising=500000,net_profit=3100000,cash_flow=0).items():
            assert report[key]==value,(key,report[key])
        if source == 'Kaspi':
            assert report['commission_estimated'] == 0
            assert report['commission_estimate_suppressed'] is False

def test_wb_nonfinancial_realization_never_becomes_expense(client):
    with SessionLocal() as db:
        acc=db.get(Account,1);acc.channel='Wildberries';u=db.get(User,1)
        row=staged(db,acc,'realization:meta',record_type='realization',sellerOperName='Корректировка карточки товара',deduction=999999,deliveryService=888888,rrDate='2026-10-02',currency='KZT')
        ok,reason=mp.project_record(db,u,acc,row)
        assert not ok and row.status=='received'
        assert reason.startswith('WB_OPERATION_NOT_FINANCIAL_OR_UNCLASSIFIED')
        report=f.report(db,u,date(2026,10,1),date(2026,10,31),1)
        assert report['expenses']==0 and report['net_profit']==0

def test_wb_sale_does_not_double_book_logistics_field(client):
    buy(client,qty=1,cost=1000)
    with SessionLocal() as db:
        acc=db.get(Account,1);acc.channel='Wildberries';u=db.get(User,1)
        row=staged(db,acc,'realization:sale',record_type='realization',sellerOperName='Продажа',vendorCode='A',quantity=1,retailPriceWithDisc=10000,currency='KZT',saleDt='2026-10-02',rrDate='2026-10-02',srid='S1',retailAmount=10000,ppvzForPay=4800,deliveryService=4000,ppvzSalesCommission=12,commissionPercent=24)
        assert mp.project_record(db,u,acc,row)[0]
        report=f.report(db,u,date(2026,10,1),date(2026,10,31),1)
        assert report['revenue']==1000000 and report['commission']==1200 and report['logistics']==400000
        # Commission is booked from the explicit money field, not its percentage field.


def test_wb_negative_commission_and_vat_adjustments_post_as_signed_expenses(client):
    buy(client, qty=1, cost=1000)
    with SessionLocal() as db:
        acc=db.get(Account,1); acc.channel='Wildberries'; u=db.get(User,1)
        row=staged(db,acc,'realization:credit',record_type='realization',sellerOperName='Продажа',
            vendorCode='A',quantity=1,currency='KZT',saleDt='2026-10-02',rrDate='2026-10-02',
            srid='CREDIT-SRID',retailAmount='5000',forPay='5349.94',ppvzSalesCommission='-301.67',
            vwNds='-48.27',commissionPercent=-2.79)
        assert mp.project_record(db,u,acc,row)[0]
        report=f.report(db,u,date(2026,10,1),date(2026,10,31),1)
        assert report['revenue']==500000
        assert report['commission']==-34994
        assert report['net_profit']==533994
        signed=db.scalars(select(Expense).where(Expense.amount < 0)).all()
        assert sorted((e.category,e.amount) for e in signed)==[
            ('Marketplace commission',-30167),('Marketplace commission VAT',-4827)]

def test_wb_reprocess_preserves_historical_generated_expense(client):
    with SessionLocal() as db:
        acc=db.get(Account,1);acc.channel='Wildberries';u=db.get(User,1)
        row=staged(db,acc,'realization:badold',record_type='realization',sellerOperName='Корректировка карточки товара',deduction=100000,rrDate='2026-10-02',currency='KZT')
        ref=mp.reference(acc,row.external_id+':Marketplace deductions','expense')
        canonical.post_event(db,u,'expense',s.ExpenseIn(company_id=1,account_id=1,reference=ref,date=date(2026,10,2),category='Marketplace deductions',amount=10000000))
        assert f.report(db,u,date(2026,10,1),date(2026,10,31),1)['expenses']==10000000
        result=mp.project_account(db,u,acc)
        assert result['reconciled_false_expenses']==0
        assert f.report(db,u,date(2026,10,1),date(2026,10,31),1)['expenses']==10000000

def test_wb_realization_posts_revenue_without_fifo_and_without_precreated_product(client):
    """A real WB sale must reach the ERP even on a brand-new QOR account.
    Missing historical cost may delay COGS, never revenue/units.
    """
    with SessionLocal() as db:
        acc=db.get(Account,1); acc.channel='Wildberries'; u=db.get(User,1)
        row=staged(db,acc,'realization:9001',record_type='realization',sellerOperName='Продажа',
            sa_name='WB-NEW-SKU',subject_name='Реальный товар WB',nm_id=987654,quantity=2,
            retail_price_withdisc_rub='10000',currency='KZT',sale_dt='2026-10-02',rr_dt='2026-10-02',srid='REAL-SRID')
        result=mp.project_account(db,u,acc)
        assert result['pending_erp']==1,result
        sale=db.scalar(select(Sale))
        assert sale and sale.quantity==2 and sale.revenue==2000000 and sale.cogs_pending is True
        movements=db.scalars(select(Movement).where(Movement.product_id==sale.product_id)).all()
        assert sum(x.quantity for x in movements)==-2
        assert any(x.batch_id is None and x.kind=='sale_pending_fifo' for x in movements)
        assert row.status=='projected'
        rep=f.report(db,u,date(2026,10,1),date(2026,10,31),1)
        assert rep['revenue']==2000000 and rep['units']==2
        assert rep['net_profit'] is None and rep['profit_complete'] is False
