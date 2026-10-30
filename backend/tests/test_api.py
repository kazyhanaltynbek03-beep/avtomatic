import io, json
from openpyxl import Workbook, load_workbook
from .conftest import buy, sell, report
from app.db import SessionLocal
from app.models import User, Job, Account, Batch, Movement
from app.worker import run_one
from app import ai
from sqlalchemy import select, func


def switch(client, email):
    a = client.post(
        "/api/auth/login", json={"email": email, "password": "Test-password-1!"}
    )
    assert a.status_code == 200, a.text
    client.headers["X-CSRF-Token"] = a.json()["csrf"]


def test_auth_csrf_logout(client):
    csrf = client.headers.pop("X-CSRF-Token")
    assert (
        client.post("/api/data/products", json={"sku": "X", "name": "X"}).status_code
        == 403
    )
    client.headers["X-CSRF-Token"] = csrf
    assert client.post("/api/auth/logout").status_code == 200
    assert client.get("/api/data/products").status_code == 401


def test_rbac_company_scope(client):
    from app.db import SessionLocal
    from app.models import ProductCompany
    with SessionLocal() as db:
        db.add(ProductCompany(product_id=1, company_id=2))
        db.commit()
    buy(client, company=1)
    buy(client, ref="other", company=2)
    switch(client, "warehouse@test.kz")
    assert client.get("/api/report?start=2026-10-01&end=2026-10-31").status_code == 403
    assert (
        client.post(
            "/api/jobs", json={"kind": "assistant", "payload": {"question": "profit"}}
        ).status_code
        == 403
    )
    rows = client.get("/api/data/batches").json()["items"]
    assert {r["company_id"] for r in rows} == {1}
    assert client.get("/api/stock?company_id=2").status_code == 403
    assert client.post("/api/data/users", json={}).status_code == 403



def test_product_initial_receipt_is_atomic_and_rejects_missing_cost(client):
    payload = {
        "sku": "INITIAL-RECEIPT-TEST",
        "name": "Initial stock test",
        "purchase_cost": 125000,
        "initial_quantity": 4,
        "initial_receipt_date": "2026-10-12",
        "specifications": {"_qor_company_ids": [1]},
    }
    response = client.post("/api/data/products", json=payload)
    assert response.status_code == 200, response.text
    product = response.json()
    rows = client.get("/api/stock?company_id=1").json()
    stock = next(row for row in rows if row["product_id"] == product["id"])
    assert stock["quantity"] == 4
    assert stock["value"] == 500000
    with SessionLocal() as db:
        batch = db.scalar(select(Batch).where(Batch.product_id == product["id"]))
        assert batch is not None
        assert str(batch.date) == "2026-10-12"
        assert batch.quantity == 4
        movement_count = db.scalar(select(func.count()).select_from(Movement).where(
            Movement.product_id == product["id"]
        ))
        assert movement_count == 1

    invalid = {**payload, "sku": "INITIAL-RECEIPT-NO-COST", "purchase_cost": 0}
    rejected = client.post("/api/data/products", json=invalid)
    assert rejected.status_code == 422
    assert client.get("/api/data/products?search=INITIAL-RECEIPT-NO-COST").json()["total"] == 0


def test_master_crud_and_audit(client):
    row = client.post(
        "/api/data/products", json={"sku": "NEW", "name": "New product", "price": 100,
                                     "specifications": {"_qor_company_ids": [1]}}
    ).json()
    assert (
        client.put(
            "/api/data/products/" + str(row["id"]),
            json={"version": row["version"], "sku": "NEW", "name": "Changed", "price": 200},
        ).status_code
        == 200
    )
    assert client.get("/api/data/products?search=Changed").json()["total"] == 1
    assert client.delete("/api/data/products/" + str(row["id"])).status_code == 200
    assert client.get("/api/data/audit").json()["total"] >= 4


def test_posted_immutable(client):
    p = buy(client)
    assert (
        client.put("/api/data/purchases/" + str(p["id"]), json={"total": 1}).status_code
        == 405
    )
    assert client.delete("/api/data/purchases/" + str(p["id"])).status_code == 405


def test_import_preview_no_writes_commit_and_duplicates(client):
    payload = {
        "kind": "products",
        "rows": [{"code": "IMP", "title": "Imported", "price": "12.34"}],
        "mapping": {"sku": "code", "name": "title", "price": "price"},
        "defaults": {"company_id": 1},
    }
    p = client.post("/api/import/preview", json=payload)
    assert p.status_code == 200, p.text
    p = p.json()
    assert p["errors"] == []
    assert client.get("/api/data/products?search=IMP").json()["total"] == 0
    assert client.post(f"/api/import/{p['id']}/confirm").status_code == 200
    assert (
        client.get("/api/data/products?search=IMP").json()["items"][0]["price"] == 1234
    )
    assert client.post(f"/api/import/{p['id']}/confirm").status_code == 409
    assert client.post("/api/import/preview", json=payload).status_code == 409


def test_import_errors_block_commit(client):
    payload = {
        "kind": "products",
        "rows": [{"sku": "A", "name": "Duplicate"}, {"sku": "NEW", "name": "Valid"}],
        "mapping": {"sku": "sku", "name": "name"},
    }
    p = client.post("/api/import/preview", json=payload).json()
    assert p["errors"]
    assert client.get("/api/data/products?search=NEW").json()["total"] == 0
    assert client.post(f"/api/import/{p['id']}/confirm").status_code == 409


def test_csv_xlsx_parse_and_export(client):
    for ext, raw in [("csv", b"sku,name,price\nC,CSV,12.50\n")]:
        r = client.post("/api/import/parse", files={"file": ("data." + ext, raw)})
        assert r.status_code == 200
        assert r.json()["count"] == 1
    w = Workbook()
    w.active.append(["sku", "name"])
    w.active.append(["XLS", "Excel"])
    buf = io.BytesIO()
    w.save(buf)
    assert (
        client.post(
            "/api/import/parse", files={"file": ("data.xlsx", buf.getvalue())}
        ).json()["count"]
        == 1
    )
    res = client.get("/api/export/products?format=xlsx")
    assert res.status_code == 200
    wb = load_workbook(io.BytesIO(res.content))
    assert wb.active.max_row == 3


def test_import_sales_fifo_preview_rollback(client):
    buy(client)
    row = {
        "company_id": "1",
        "account_id": "1",
        "product_id": "1",
        "reference": "IMPS",
        "date": "2026-10-02",
        "quantity": "2",
        "unit_price": "30.00",
    }
    p = client.post(
        "/api/import/preview",
        json={"kind": "sales", "rows": [row], "mapping": {k: k for k in row}},
    ).json()
    assert p["errors"] == []
    assert report(client)["stock_value"] == 10000
    assert client.post(f"/api/import/{p['id']}/confirm").status_code == 200
    assert report(client)["revenue"] == 6000
    assert report(client)["stock_value"] == 8000


def test_demo_connector_no_financial_writes(client):
    with SessionLocal() as db:
        db.get(Account, 1).mode = "demo"
        db.commit()
    res = client.post("/api/accounts/1/sync-now").json()
    assert res["status"] == "done"
    assert res["source"] == "demo"
    assert report(client)["revenue"] == 0


def test_missing_ai_key_and_real_connector_fail_honestly(client, monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    j = client.post(
        "/api/jobs",
        json={
            "kind": "assistant",
            "payload": {
                "question": "Profit",
                "start": "2026-10-01",
                "end": "2026-10-31",
            },
        },
    ).json()
    run_one()
    res = client.get("/api/jobs/" + str(j["id"])).json()
    assert res["status"] == "failed"
    assert "ANTHROPIC_API_KEY" in res["error"]
    with SessionLocal() as db:
        db.get(Account, 1).mode = "api"
        db.commit()
    res = client.post("/api/accounts/1/sync-now")
    assert res.status_code == 422
    assert "credentials" in res.text.lower() or "token" in res.text.lower()


def test_ai_function_boundary_with_mock_provider(client, monkeypatch):
    buy(client)
    sell(client)
    calls = []

    def fake(inputs, instructions, tools=None, tool_choice=None):
        calls.append(inputs.copy())
        if len(calls) == 1:
            return {
                "output": [
                    {
                        "type": "function_call",
                        "name": "get_finance_summary",
                        "call_id": "t1",
                        "arguments": "{}",
                        "id": "fc1",
                        "status": "completed",
                    }
                ]
            }
        evidence = json.loads(inputs[-1]["output"])
        assert evidence["revenue"] == 9000
        return {
            "output": [
                {
                    "type": "message",
                    "content": [
                        {"type": "output_text", "text": "Показатели получены из ERP."}
                    ],
                }
            ]
        }

    monkeypatch.setattr(ai, "call_response", fake)
    j = client.post(
        "/api/jobs",
        json={
            "kind": "assistant",
            "payload": {
                "question": "Profit",
                "start": "2026-10-01",
                "end": "2026-10-31",
            },
        },
    ).json()
    run_one()
    res = client.get("/api/jobs/" + str(j["id"])).json()
    assert res["status"] == "done"
    assert "get_finance_summary" in res["result"]["tools_used"]
    assert len(calls) == 2


def test_ai_assistant_receives_chat_history_and_uploaded_pdf(client, monkeypatch):
    calls = []

    def fake(inputs, instructions, tools=None, tool_choice=None):
        calls.append(inputs)
        if len(calls) == 1:
            return {"output": [{
                "type": "function_call", "name": "get_finance_summary",
                "call_id": "finance-1", "arguments": "{}", "id": "call-1",
                "status": "completed",
            }]}
        return {"output": [{
            "type": "message",
            "content": [{"type": "output_text", "text": "Данные ERP проверены."}],
        }]}

    monkeypatch.setattr(ai, "call_response", fake)
    doc = client.post(
        "/api/documents/upload",
        data={"company_id": 1},
        files={"file": ("report.pdf", b"%PDF-1.4 test report", "application/pdf")},
    ).json()
    job = client.post("/api/jobs", json={
        "kind": "assistant",
        "payload": {
            "question": "Сравни это с ERP",
            "agent": "rick",
            "company_id": 1,
            "start": "2026-10-01",
            "end": "2026-10-22",
            "document_ids": [doc["id"]],
            "history": [
                {"role": "user", "content": "Сколько продаж?"},
                {"role": "assistant", "content": "Проверяю период."},
                {"role": "user", "content": "Сравни это с ERP"},
            ],
        },
    }).json()

    result = client.get("/api/jobs/" + str(job["id"])).json()
    assert result["status"] == "done", result
    assert result["result"]["answer"] == "Данные ERP проверены."
    initial_messages = [x for x in calls[0] if isinstance(x, dict) and x.get("role")]
    assert [x["role"] for x in initial_messages] == ["user", "assistant", "user"]
    final_content = initial_messages[-1]["content"]
    assert any(x.get("type") == "input_file" and x.get("filename") == "report.pdf" for x in final_content)
    assert len(calls) == 2


def test_niche_insufficient_data_and_score(client):
    r = client.post(
        "/api/data/niches",
        json={
            "name": "Test",
            "source": "Supplier report 2026",
            "observed_at": "2026-10-01",
            "data": {},
        },
    )
    assert r.status_code == 200
    assert (
        client.get("/api/data/niches").json()["items"][0]["analysis"]["score"] is None
    )
    assert (
        0
        <= ai.niche_score(
            {
                "monthly_demand": 500,
                "competitors": 10,
                "margin_percent": 30,
                "advertising_percent": 10,
                "seasonality_risk": 20,
                "logistics_percent": 8,
                "rating": 4,
                "return_percent": 5,
            }
        )["score"]
        <= 100
    )


def test_calculator_exact_result(client):
    r = client.post(
        "/api/calculator",
        json={
            "quantity": 10,
            "price": 100,
            "unit_cost": 40,
            "commission_percent": 10,
            "logistics": 20,
            "advertising": 30,
        },
    ).json()
    assert r["net_profit"] == "450.00"
    assert r["margin_percent"] == "45.00"


def test_channel_filter(client):
    buy(client)
    sell(client)
    sell(client, ref="B2B", account=3, customer_id=1)
    assert client.get("/api/data/sales?channel=Kaspi").json()["total"] == 1
    assert client.get("/api/data/sales?channel=Ozon").json()["total"] == 0


def test_document_access_and_invoice(client):
    doc = client.post(
        "/api/documents/upload",
        data={"company_id": 2},
        files={"file": ("a.pdf", b"%PDF-1.4 test", "application/pdf")},
    ).json()
    buy(client)
    s = sell(client)
    assert (
        "Счёт".encode() in client.get("/api/sales/" + str(s["id"]) + "/invoice").content
    )
    switch(client, "warehouse@test.kz")
    assert (
        client.get("/api/documents/" + str(doc["id"]) + "/download").status_code == 403
    )


def test_password_change_invalidates_sessions(client):
    assert (
        client.post(
            "/api/auth/password",
            json={"old": "Test-password-1!", "new": "New-password-123!"},
        ).status_code
        == 200
    )
    assert client.get("/api/auth/me").status_code == 401


def test_import_all_financial_types(client):
    def imp(kind, row):
        p = client.post(
            "/api/import/preview",
            json={"kind": kind, "rows": [row], "mapping": {k: k for k in row}},
        )
        assert p.status_code == 200, p.text
        p = p.json()
        assert p["errors"] == [], p
        r = client.post(f"/api/import/{p['id']}/confirm")
        assert r.status_code == 200, r.text

    imp(
        "purchases",
        {
            "company_id": "1",
            "reference": "IP",
            "supplier": "Supplier",
            "date": "2026-10-01",
            "product_id": "1",
            "quantity": "10",
            "unit_cost": "10",
            "extra_cost": "1",
            "paid": "50",
        },
    )
    imp(
        "stock",
        {
            "company_id": "1",
            "product_id": "2",
            "date": "2026-10-01",
            "quantity": "5",
            "unit_cost": "2",
            "reason": "Opening balance",
        },
    )
    imp(
        "expenses",
        {
            "company_id": "1",
            "reference": "IE",
            "date": "2026-10-01",
            "category": "Аренда",
            "amount": "10",
            "paid": "true",
        },
    )
    imp(
        "bank",
        {
            "company_id": "1",
            "reference": "IB",
            "date": "2026-10-01",
            "amount": "100",
            "category": "Вклад собственника",
        },
    )
    imp(
        "b2b",
        {
            "company_id": "1",
            "account_id": "3",
            "customer_id": "1",
            "product_id": "1",
            "reference": "ISB",
            "date": "2026-10-02",
            "quantity": "1",
            "unit_price": "30",
        },
    )
    imp(
        "reports",
        {
            "company_id": "1",
            "account_id": "1",
            "product_id": "1",
            "reference": "ISM",
            "date": "2026-10-02",
            "quantity": "1",
            "unit_price": "30",
        },
    )
    r = report(client)
    assert r["revenue"] == 6000
    assert r["b2b_debt"] == 3000
    assert r["cash"] == 4000


def test_marketplace_sync_not_available_via_legacy_jobs(client):
    res = client.post("/api/jobs", json={"kind": "sync", "payload": {"account_id": 1}})
    assert res.status_code == 410


def test_formula_injection_export(client):
    client.post("/api/data/products", json={"sku": "=1+1", "name": '=HYPERLINK("bad")',
                                             "specifications": {"_qor_company_ids": [1]}})
    res = client.get("/api/export/products?format=csv")
    assert "'=1+1" in res.text


def test_overpayment_rejected(client):
    buy(client)
    s = sell(client)
    r = client.post(
        "/api/post/cash",
        json={
            "company_id": 1,
            "sale_id": s["id"],
            "reference": "OVER",
            "date": "2026-10-03",
            "amount": 10000,
            "category": "Продажа",
        },
    )
    assert r.status_code == 422

def test_marketplace_credentials_can_be_replaced_without_409(client, monkeypatch):
    monkeypatch.setenv("CREDENTIAL_ENCRYPTION_KEY", "test-encryption-key-at-least-32-bytes")
    with SessionLocal() as db:
        acc = db.get(Account, 1)
        acc.channel = "Wildberries"
        acc.mode = "api"
        db.commit()
    r1 = client.put("/api/accounts/1/credentials", json={"token": "first-token-value"})
    assert r1.status_code == 200
    r2 = client.put("/api/accounts/1/credentials", json={"token": "second-token-value"})
    assert r2.status_code == 200
    assert r2.json()["configured"] is True


def test_legacy_marketplace_job_endpoint_is_gone(client):
    r = client.post("/api/jobs", json={"kind": "sync", "payload": {"account_id": 1}})
    assert r.status_code == 410
