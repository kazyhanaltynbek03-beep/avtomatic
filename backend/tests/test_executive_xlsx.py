import io
from openpyxl import load_workbook
from .conftest import buy, sell


def test_executive_xlsx_has_reconciled_financial_sheets(client):
    buy(client, qty=10, cost=1000)
    sell(client, qty=2, price=3000)
    r=client.get('/api/reports/executive.xlsx?start=2026-10-01&end=2026-10-31')
    assert r.status_code == 200
    assert r.content[:2] == b'PK'
    wb=load_workbook(io.BytesIO(r.content), data_only=True)
    assert {'Dashboard','P&L','Cash Flow','Продажи','Склад','Динамика','Контроль'} <= set(wb.sheetnames)
    ctl=wb['Контроль']
    assert [ctl.cell(i,5).value for i in range(5,9)] == ['PASS']*4
    pnl=wb['P&L']
    # backend stores tiyn; report displays KZT
    assert pnl['B5'].value == 60
    assert pnl['B6'].value == 20
    assert pnl['B7'].value == 40


def test_executive_xlsx_rejects_inverted_period(client):
    r=client.get('/api/reports/executive.xlsx?start=2026-10-31&end=2026-10-01')
    assert r.status_code == 422
