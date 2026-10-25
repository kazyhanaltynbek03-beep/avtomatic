import io
from openpyxl import Workbook, load_workbook
from fastapi import HTTPException
from app.importing import parse_file, export_bytes


def _xlsx(rows):
    wb=Workbook(); ws=wb.active
    for row in rows: ws.append(row)
    out=io.BytesIO(); wb.save(out); return out.getvalue()


def test_xlsx_import_rejects_formulas_instead_of_stale_calculated_values():
    raw=_xlsx([["sku","price"],["A-1","=100+200"]])
    try:
        parse_file(raw,"products.xlsx")
        assert False, "formula must be rejected"
    except HTTPException as exc:
        assert exc.status_code == 422
        assert "формул" in str(exc.detail).lower()


def test_xlsx_export_never_creates_formula_from_untrusted_text():
    raw=export_bytes([{"description":"=HYPERLINK(\"https://bad\",\"x\")"}], ["description"], "xlsx")
    wb=load_workbook(io.BytesIO(raw), data_only=False)
    cell=wb.active["A2"]
    assert cell.data_type != "f"
    assert str(cell.value).startswith("'=")
    wb.close()


def test_csv_export_neutralizes_formula_injection():
    raw=export_bytes([{"description":"+cmd"}], ["description"], "csv").decode("utf-8-sig")
    assert "'+cmd" in raw
