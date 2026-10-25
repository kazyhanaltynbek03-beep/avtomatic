"""Read-only WB weekly Excel reconciliation; never posts ledger entries."""
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from io import BytesIO
from openpyxl import load_workbook

COLUMNS = {
    'retail': 'Цена розничная',
    'retail_after_agreed_discount': 'Цена розничная с учетом согласованной скидки',
    'wb_realized': 'Вайлдберриз реализовал Товар (Пр)',
    'seller_for_pay': 'К перечислению Продавцу за реализованный Товар',
    'delivery': 'Услуги по доставке товара покупателю',
    'penalty': 'Общая сумма штрафов',
    'storage': 'Хранение',
    'deduction': 'Удержания',
    'acquiring': 'Компенсация платёжных услуг/Комиссия за интеграцию платёжных сервисов',
    'commission_before_agent': 'Вознаграждение с продаж до вычета услуг поверенного, без НДС',
    'commission_without_vat': 'Вознаграждение Вайлдберриз (ВВ), без НДС',
    'commission_vat': 'НДС с Вознаграждения Вайлдберриз',
    'commission_adjustment': 'Корректировка Вознаграждения Вайлдберриз (ВВ)',
    'returns_count': 'Количество возврата',
    'quantity': 'Кол-во',
}

def money(value):
    if value is None or value == '': return Decimal('0')
    if isinstance(value, bool): raise ValueError('Boolean is not money')
    try: return Decimal(str(value).replace(' ', '').replace('\u00a0','').replace(',', '.'))
    except InvalidOperation as exc: raise ValueError(f'Invalid numeric cell: {value!r}') from exc

def parse_report(data: bytes):
    if len(data) > 5 * 1024 * 1024: raise ValueError('File exceeds 5 MB')
    book = load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        sheet = book.active
        sheet.reset_dimensions()  # WB exports may contain incorrect worksheet dimensions
        iterator = sheet.iter_rows(values_only=True)
        headers = [str(v).strip() if v is not None else '' for v in next(iterator)]
        indexes = {h:i for i,h in enumerate(headers) if h}
        if 'Обоснование для оплаты' not in indexes or 'К перечислению Продавцу за реализованный Товар' not in indexes:
            raise ValueError('Not a recognized Wildberries detailed financial report')
        available = {k:indexes[v] for k,v in COLUMNS.items() if v in indexes}
        totals = defaultdict(lambda: defaultdict(Decimal))
        rows=[]
        for line, cells in enumerate(iterator, start=2):
            if line > 50002: raise ValueError('Report exceeds 50,000 rows')
            if not any(x is not None for x in cells): continue
            reason = str(cells[indexes['Обоснование для оплаты']] or '').strip()
            category = 'sale' if reason.lower() == 'продажа' else ('return' if 'возврат' in reason.lower() else 'other')
            amounts = {k:money(cells[i] if i < len(cells) else None) for k,i in available.items()}
            for k,value in amounts.items(): totals[category][k] += value
            rows.append({'line':line, 'category':category, 'reason':reason,
                         'srid':str(cells[indexes['Srid']]) if 'Srid' in indexes and indexes['Srid'] < len(cells) and cells[indexes['Srid']] is not None else None,
                         'amounts':{k:str(v) for k,v in amounts.items()}})
        return {'sheet':sheet.title,'row_count':len(rows),'currency':None,
                'currency_note':'В исходном XLSX нет отдельного поля валюты; не предполагается автоматически.',
                'available_fields':list(available),'missing_fields':[k for k in COLUMNS if k not in available],
                'by_category':{cat:{k:str(v) for k,v in values.items()} for cat,values in totals.items()},
                'totals':{k:str(sum((values.get(k,Decimal('0')) for values in totals.values()),Decimal('0'))) for k in available},
                'rows':rows[:100], 'rows_truncated':len(rows)>100,
                'note':'Read-only WB report preview. No ERP postings, exchange conversion, or inferred profit.'}
    finally: book.close()
