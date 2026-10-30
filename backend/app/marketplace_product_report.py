from collections import defaultdict
from datetime import date
from io import BytesIO

from fastapi import HTTPException
from sqlalchemy import select, func
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .models import Account, Batch, MarketplaceProductMapping, MarketplaceProductReportSetting, Product, ProductCompany, Return, Sale
from .security import require, scope


CHANNELS = {"Kaspi", "Wildberries", "Ozon"}


def _accounts(db, user, start, end, company_id=None, channel=None, account_id=None):
    if end < start:
        raise HTTPException(422, "Дата окончания раньше начальной")
    if channel and channel not in CHANNELS:
        raise HTTPException(422, "Неизвестный маркетплейс")
    query = select(Account).where(Account.channel.in_(CHANNELS), Account.active.is_(True))
    if user.role != "Owner":
        query = query.where(Account.company_id.in_(user.company_ids or []))
    if company_id is not None:
        scope(user, company_id)
        query = query.where(Account.company_id == company_id)
    if channel:
        query = query.where(Account.channel == channel)
    if account_id is not None:
        query = query.where(Account.id == account_id)
    return list(db.scalars(query))


def build_report(db, user, start, end, company_id=None, channel=None, account_id=None):
    accounts = _accounts(db, user, start, end, company_id, channel, account_id)
    account_map = {account.id: account for account in accounts}
    account_ids = set(account_map)
    if not account_ids:
        return {"start": str(start), "end": str(end), "items": [], "total": 0}

    sales = list(db.scalars(
        select(Sale).where(
            Sale.account_id.in_(account_ids),
            Sale.date >= start,
            Sale.date <= end,
        )
    ))
    return_pairs = list(db.execute(
        select(Return, Sale).join(Sale, Return.sale_id == Sale.id).where(
            Sale.account_id.in_(account_ids),
            Return.date >= start,
            Return.date <= end,
        )
    ))
    returns = [pair[0] for pair in return_pairs]
    return_sales = {ret.id: sale for ret, sale in return_pairs}
    product_ids = {sale.product_id for sale in sales}
    product_ids.update(sale.product_id for sale in return_sales.values())
    product_ids.discard(None)
    products = {
        product.id: product
        for product in db.scalars(select(Product).where(Product.id.in_(product_ids or {-1})))
    }
    company_ids = {account.company_id for account in accounts}
    settings = {
        (setting.company_id, setting.product_id): setting
        for setting in db.scalars(
            select(MarketplaceProductReportSetting).where(
                MarketplaceProductReportSetting.company_id.in_(company_ids or {-1}),
                MarketplaceProductReportSetting.product_id.in_(product_ids or {-1}),
            )
        )
    }
    sales_by_id = {sale.id: sale for sale in sales}
    keys = {(sale.company_id, sale.product_id, sale.account_id) for sale in sales}
    keys.update(
        (return_sales[ret.id].company_id, return_sales[ret.id].product_id,
         return_sales[ret.id].account_id)
        for ret in returns if ret.id in return_sales
    )
    grouped = {}
    for company, product_id, account_id_value in keys:
        product = products.get(product_id)
        account = account_map.get(account_id_value)
        if not product or not account:
            continue
        setting = settings.get((company, product_id))
        grouped[(company, product_id, account_id_value)] = {
            "company_id": company,
            "product_id": product.id,
            "sku": product.sku or "",
            "article": product.article or "",
            "product_name": product.name or "",
            "channel": account.channel,
            "account_id": account.id,
            "account_name": account.name,
            "sold_units": 0,
            "sold_revenue": 0,
            "returned_units": 0,
            "returned_revenue": 0,
            "restocked_returned_units": 0,
            "commission": 0,
            "logistics": 0,
            "advertising": 0,
            "unit_cost_cents": setting.unit_cost_cents if setting else None,
            "tax_rate_bps": setting.tax_rate_bps if setting else None,
        }

    for sale in sales:
        row = grouped.get((sale.company_id, sale.product_id, sale.account_id))
        if row is None:
            continue
        row["sold_units"] += int(sale.quantity or 0)
        row["sold_revenue"] += int(sale.revenue or 0)
        row["commission"] += int(sale.commission or 0)
        row["logistics"] += int(sale.logistics or 0)
        row["advertising"] += int(sale.advertising or 0)
    for ret in returns:
        sale = return_sales.get(ret.id)
        if sale is None:
            continue
        row = grouped.get((sale.company_id, sale.product_id, sale.account_id))
        if row is None:
            continue
        row["returned_units"] += int(ret.quantity or 0)
        row["returned_revenue"] += int(ret.revenue or 0)
        if ret.inventory_status == "restocked":
            row["restocked_returned_units"] += int(ret.quantity or 0)

    items = []
    for row in grouped.values():
        row["net_units"] = row["sold_units"] - row["returned_units"]
        row["net_revenue"] = row["sold_revenue"] - row["returned_revenue"]
        unit_cost = row["unit_cost_cents"]
        tax_rate = row["tax_rate_bps"]
        row["cogs"] = (
            (row["sold_units"] - row["restocked_returned_units"]) * unit_cost
            if unit_cost is not None else None
        )
        row["tax"] = (
            round(row["net_revenue"] * tax_rate / 10000)
            if tax_rate is not None else None
        )
        row["profit"] = (
            row["net_revenue"] - row["cogs"] - row["tax"]
            - row["commission"] - row["logistics"] - row["advertising"]
            if row["cogs"] is not None and row["tax"] is not None else None
        )
        row["unit_cost_tenge"] = (
            round(unit_cost / 100, 2) if unit_cost is not None else None
        )
        row["tax_rate_percent"] = (
            round(tax_rate / 100, 2) if tax_rate is not None else None
        )
        items.append(row)
    # The customer report is keyed by the internal SKU. Marketplace cards and
    # stores are detail beneath that product, never separate summary rows.
    batch_rows = db.execute(
        select(
            Batch.company_id, Batch.product_id,
            func.coalesce(func.sum(Batch.quantity), 0),
            func.coalesce(func.sum(Batch.total_cost), 0),
            func.coalesce(func.sum(Batch.remaining), 0),
        ).where(
            Batch.company_id.in_(company_ids or {-1}),
            Batch.product_id.in_(product_ids or {-1}),
        ).group_by(Batch.company_id, Batch.product_id)
    ).all()
    batch_stats = {
        (company, product_id): {
            "quantity": int(quantity or 0), "total_cost": int(total_cost or 0),
            "stock": int(stock or 0),
        }
        for company, product_id, quantity, total_cost, stock in batch_rows
    }
    combined = {}
    marketplace_groups = {}
    amount_fields = (
        "sold_units", "sold_revenue", "returned_units", "returned_revenue",
        "restocked_returned_units", "commission", "logistics", "advertising",
    )
    for item in items:
        key = (item["company_id"], item["product_id"])
        row = combined.get(key)
        if row is None:
            row = {**item, **{field: 0 for field in amount_fields}}
            row["marketplace_breakdown"] = []
            combined[key] = row
        for field in amount_fields:
            row[field] += int(item.get(field) or 0)
        breakdown_key = (key[0], key[1], item["channel"], item["account_id"])
        detail = marketplace_groups.get(breakdown_key)
        if detail is None:
            detail = {
                "channel": item["channel"], "account_id": item["account_id"],
                "account_name": item["account_name"],
                **{field: 0 for field in amount_fields},
            }
            marketplace_groups[breakdown_key] = detail
        for field in amount_fields:
            detail[field] += int(item.get(field) or 0)

    links = list(db.scalars(select(MarketplaceProductMapping).where(
        MarketplaceProductMapping.company_id.in_(company_ids or {-1}),
        MarketplaceProductMapping.product_id.in_(product_ids or {-1}),
        MarketplaceProductMapping.account_id.in_(account_ids or {-1}),
    )))
    linked_cards = defaultdict(list)
    for link in links:
        account = account_map.get(link.account_id)
        if account:
            linked_cards[(link.company_id, link.product_id)].append({
                "channel": link.channel, "store": account.name,
                "external_id": link.external_key,
            })

    products_for_price = products
    for key, row in combined.items():
        stats = batch_stats.get(key, {})
        product = products_for_price.get(key[1])
        row["available_stock"] = stats.get("stock", 0)
        row["selling_price_tenge"] = round((product.price or 0) / 100, 2) if product else 0
        media = (product.media or []) if product else []
        first_media = media[0] if media else ""
        row["photo"] = (first_media.get("url") or first_media.get("src") or "") if isinstance(first_media, dict) else str(first_media or "")
        unit_cost = row.get("unit_cost_cents")
        if unit_cost is None and stats.get("quantity", 0) > 0:
            unit_cost = round(stats["total_cost"] / stats["quantity"])
            row["unit_cost_cents"] = unit_cost
        tax_rate = row.get("tax_rate_bps")
        row["net_units"] = row["sold_units"] - row["returned_units"]
        row["net_revenue"] = row["sold_revenue"] - row["returned_revenue"]
        row["cogs"] = (
            (row["sold_units"] - row["restocked_returned_units"]) * unit_cost
            if unit_cost is not None else None
        )
        row["tax"] = (
            round(row["net_revenue"] * tax_rate / 10000)
            if tax_rate is not None else None
        )
        row["profit"] = (
            row["net_revenue"] - row["cogs"] - row["tax"]
            - row["commission"] - row["logistics"] - row["advertising"]
            if row["cogs"] is not None and row["tax"] is not None else None
        )
        row["unit_cost_tenge"] = round(unit_cost / 100, 2) if unit_cost is not None else None
        row["tax_rate_percent"] = round(tax_rate / 100, 2) if tax_rate is not None else None
        row["marketplace_breakdown"] = sorted(
            (detail for detail_key, detail in marketplace_groups.items()
             if detail_key[:2] == key),
            key=lambda detail: (detail["channel"], detail["account_name"].casefold()),
        )
        row["marketplace_cards"] = sorted(
            linked_cards.get(key, []), key=lambda link: (link["channel"], link["store"].casefold(), link["external_id"])
        )
        row["marketplace_card_count"] = len(row["marketplace_cards"])
        row["marketplace_names"] = ", ".join(
            f'{detail["channel"]} · {detail["account_name"]}'
            for detail in row["marketplace_breakdown"]
        )
        row["channel"] = " + ".join(sorted({detail["channel"] for detail in row["marketplace_breakdown"]}))
        row["account_name"] = ", ".join(sorted({detail["account_name"] for detail in row["marketplace_breakdown"]}))
    consolidated = list(combined.values())
    consolidated.sort(key=lambda row: row["product_name"].casefold())
    return {"start": str(start), "end": str(end), "items": consolidated, "total": len(consolidated)}


def save_setting(db, user, company_id, product_id, unit_cost_tenge, tax_rate_percent):
    scope(user, company_id)
    require(user, "finance", True)
    product = db.get(Product, product_id)
    if not product:
        raise HTTPException(404, "Товар не найден")
    owned = db.scalar(select(ProductCompany.id).where(
        ProductCompany.product_id == product_id,
        ProductCompany.company_id == company_id,
    ))
    if not owned:
        raise HTTPException(404, "Товар не принадлежит выбранной компании")
    if unit_cost_tenge is not None and unit_cost_tenge < 0:
        raise HTTPException(422, "Себестоимость не может быть отрицательной")
    if tax_rate_percent is not None and not 0 <= tax_rate_percent <= 100:
        raise HTTPException(422, "Ставка налога должна быть от 0 до 100%")
    setting = db.scalar(select(MarketplaceProductReportSetting).where(
        MarketplaceProductReportSetting.company_id == company_id,
        MarketplaceProductReportSetting.product_id == product_id,
    ))
    if setting is None:
        setting = MarketplaceProductReportSetting(
            company_id=company_id, product_id=product_id
        )
        db.add(setting)
    setting.unit_cost_cents = (
        round(unit_cost_tenge * 100) if unit_cost_tenge is not None else None
    )
    setting.tax_rate_bps = (
        round(tax_rate_percent * 100) if tax_rate_percent is not None else None
    )
    db.flush()
    return {
        "company_id": company_id,
        "product_id": product_id,
        "unit_cost_cents": setting.unit_cost_cents,
        "tax_rate_bps": setting.tax_rate_bps,
    }


def create_xlsx(report):
    book = Workbook()
    sheet = book.active
    sheet.title = "Отчёт по товарам"
    sheet.append(["Отчёт по товарам маркетплейсов"])
    sheet.append(["Период", f'{report["start"]} — {report["end"]}'])
    sheet.append([])
    headers = [
        "Артикул", "SKU", "Товар", "Маркетплейс", "Магазин",
        "Продано, шт.", "Возвращено, шт.", "Итого, шт.",
        "Выручка, ₸", "Возвраты, ₸", "Чистый доход, ₸",
        "Себестоимость за шт., ₸", "Себестоимость, ₸",
        "Налог, %", "Налог, ₸", "Комиссия, ₸", "Логистика, ₸", "Реклама, ₸",
        "Прибыль после расходов, ₸",
    ]
    sheet.append(headers)
    for cell in sheet[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="176A52")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in report["items"]:
        amount = lambda key: round(row[key] / 100, 2) if row.get(key) is not None else None
        sheet.append([
            row["article"], row["sku"], row["product_name"], row["channel"],
            row["account_name"], row["sold_units"], row["returned_units"],
            row["net_units"], amount("sold_revenue"), amount("returned_revenue"),
            amount("net_revenue"), row["unit_cost_tenge"], amount("cogs"),
            row["tax_rate_percent"], amount("tax"), amount("commission"),
            amount("logistics"), amount("advertising"), amount("profit"),
        ])
    sheet.freeze_panes = "A5"
    sheet.auto_filter.ref = sheet.dimensions
    widths = [18, 18, 36, 18, 24, 14, 16, 14, 18, 18, 18, 22, 18, 12, 16, 16, 16, 16, 24]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.row_dimensions[4].height = 42
    for cells in sheet.iter_rows(min_row=5):
        for column in (9, 10, 11, 12, 13, 15, 16, 17, 18, 19):
            cells[column - 1].number_format = '#,##0.00'
        cells[13].number_format = '0.00"%"'
    output = BytesIO()
    book.save(output)
    return output.getvalue()
