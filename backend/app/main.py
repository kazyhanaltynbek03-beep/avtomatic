from dataclasses import asdict
import os, secrets, json, hashlib, uuid, io, base64, logging, threading, time
import hmac
import re
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Literal
from datetime import datetime, timedelta, timezone, date
from zoneinfo import ZoneInfo
from decimal import Decimal
from fastapi import (
    FastAPI,
    BackgroundTasks,
    Depends,
    HTTPException,
    Request,
    Response,
    UploadFile,
    File,
    Form,
    Query,
)
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from fastapi.exceptions import RequestValidationError
from fastapi.staticfiles import StaticFiles
from fastapi.encoders import jsonable_encoder
from sqlalchemy import select, func, or_, and_, text, update, cast, String, case
from sqlalchemy.exc import IntegrityError
from pydantic import BaseModel, Field, ValidationError
from .db import get_db, SessionLocal
from .models import *
from .security import *
from . import schemas as s, finance as f, importing, ai, action_engine, reporting_xlsx, marketplace_pipeline, canonical, wb_weekly, marketplace_costs, marketplace_returns, marketplace_matching
from .connectors import AccessRequired, connector, load_saved_credentials
from cryptography.fernet import Fernet, InvalidToken
from .product_studio import assess_performance, next_experiment
from . import kaspi_integrity

app = FastAPI(
    title="QOR Business ERP",
    version="1.0.0",
    docs_url="/api/docs",
    openapi_url="/api/openapi.json",
)
logger = logging.getLogger(__name__)
SYNC_RUNNING_TIMEOUT_SECONDS = max(60, int(os.getenv("MARKETPLACE_SYNC_TIMEOUT_SECONDS", "1800")))
KASPI_JOB_RECOVERY_GRACE_SECONDS = min(
    SYNC_RUNNING_TIMEOUT_SECONDS,
    max(60, int(os.getenv("KASPI_JOB_RECOVERY_GRACE_SECONDS", "300"))),
)
KASPI_API_HISTORY_DAYS = max(1, int(os.getenv("KASPI_API_HISTORY_DAYS", "370")))
CLAUDE_SKILL_INSTALL_LOCK = threading.Lock()


def sync_status_for_display(account, at=None):
    """Hide abandoned running states after the last persisted sync heartbeat."""
    status = account.sync_status or "never"
    error = account.sync_error or ""
    if status.lower() != "running":
        return status, error
    heartbeat = account.last_sync
    if heartbeat is not None and heartbeat.tzinfo is None:
        heartbeat = heartbeat.replace(tzinfo=timezone.utc)
    reference = at or now()
    if reference.tzinfo is None:
        reference = reference.replace(tzinfo=timezone.utc)
    if heartbeat is None or heartbeat < reference - timedelta(seconds=SYNC_RUNNING_TIMEOUT_SECONDS):
        return "failed", "SYNC_TIMEOUT: синхронизация давно не обновлялась; запустите её повторно"
    return status, error


def _kaspi_issued_scan_start(today_local):
    """Configured backfill horizon, not a claim about Kaspi retention limits."""
    return today_local - timedelta(days=KASPI_API_HISTORY_DAYS)


def _kaspi_period_sync_key(report_mode, start, end):
    return kaspi_integrity.period_key(report_mode, start, end)


def _require_complete_kaspi_report(summary):
    if summary.get("sync_in_progress") or summary.get("sync_status") == "running":
        raise HTTPException(409, "Kaspi ещё синхронизируется. Дождитесь завершения и скачайте отчёт повторно.")
    if summary.get("sync_status") != "done" or not summary.get("source_complete"):
        raise HTTPException(409, "Полная синхронизация Kaspi не подтверждена. Проверьте статус и повторите выгрузку.")


def _resume_kaspi_jobs_once():
    """Recover durable API-read jobs after a restart, with a cross-process lease."""
    tasks = []
    with SessionLocal() as db:
        accounts = list(db.scalars(select(Account).where(
            Account.channel == "Kaspi", Account.active.is_(True), Account.mode == "api")))
        for account in accounts:
            settings = dict(account.integration_settings or {})
            job = dict(settings.get("kaspi_period_job") or {})
            heartbeat = account.last_sync
            if heartbeat is not None and heartbeat.tzinfo is None:
                heartbeat = heartbeat.replace(tzinfo=timezone.utc)
            if not job and account.sync_status == "running" and heartbeat and heartbeat > now()-timedelta(seconds=SYNC_RUNNING_TIMEOUT_SECONDS):
                continue
            # Optional, explicitly targeted one-time backfill. Configuration is
            # non-secret; API credentials remain in the existing encrypted store.
            target_ids = {int(x) for x in os.getenv("KASPI_BACKFILL_ACCOUNT_IDS", "").split(",") if x.strip().isdigit()}
            configured_start, configured_end = os.getenv("KASPI_BACKFILL_START"), os.getenv("KASPI_BACKFILL_END")
            if account.id in target_ids and configured_start and configured_end:
                start_day, end_day = date.fromisoformat(configured_start), date.fromisoformat(configured_end)
                if end_day < start_day or (end_day-start_day).days > 366:
                    raise ValueError("Invalid configured Kaspi backfill period")
                target_key = _kaspi_period_sync_key("issued", start_day, end_day)
                current_is_target = job.get("start") == str(start_day) and job.get("end") == str(end_day)
                target_complete = bool((settings.get("kaspi_period_sync") or {}).get(target_key))
                current_is_active = job.get("status") in {"running", "retry"}
                if not current_is_target and not target_complete:
                    if current_is_active:
                        handoff_grace = max(60, int(os.getenv("KASPI_TARGET_HANDOFF_GRACE_SECONDS", "60")))
                        if heartbeat and heartbeat > now()-timedelta(seconds=handoff_grace):
                            continue
                        # Pause the unrelated current period at its last durable
                        # cursor, then prioritize the explicitly configured
                        # customer period. Lease replacement prevents the old
                        # worker from committing a second page after this point.
                        paused = list(settings.get("kaspi_period_paused_jobs") or [])
                        if not any(x.get("start") == job.get("start") and x.get("end") == job.get("end")
                                   for x in paused):
                            paused.append({**job, "status":"paused", "paused_at":now().isoformat()})
                        settings["kaspi_period_paused_jobs"] = paused[-20:]
                    owner = db.scalar(select(User).where(User.active.is_(True), User.role == "Owner").order_by(User.id))
                    if owner:
                        # The saved cursor for this exact period is reused.
                        job = {"start":str(start_day),"end":str(end_day),"mode":"issued","user_id":owner.id,
                               "status":"retry","retry_at":now().isoformat(),"batches":0,"records":0}
            if job.get("status") not in {"running", "retry"}:
                continue
            if job["status"] == "running" and heartbeat and heartbeat > now()-timedelta(seconds=KASPI_JOB_RECOVERY_GRACE_SECONDS):
                continue
            if job["status"] == "retry" and datetime.fromisoformat(job["retry_at"]) > now():
                continue
            lease = "KASPI_PERIOD_SYNC:" + uuid.uuid4().hex
            job["status"] = "running"
            settings["kaspi_period_job"] = job
            claimed = db.execute(update(Account).where(
                Account.id == account.id, Account.sync_error == account.sync_error,
                Account.last_sync == account.last_sync if account.last_sync else Account.last_sync.is_(None),
            ).values(sync_status="running", sync_error=lease, last_sync=now(), integration_settings=settings))
            if claimed.rowcount == 1:
                tasks.append((account.id, job["user_id"], date.fromisoformat(job["start"]),
                              date.fromisoformat(job["end"]), [], job["mode"], job["mode"] == "issued", lease))
        db.commit()
    for task in tasks:
        _background_kaspi_period_sync(*task)


def _kaspi_job_loop():
    while True:
        try:
            _resume_kaspi_jobs_once()
        except Exception:
            logger.exception("Kaspi recovery scan failed")
        time.sleep(60)


@app.on_event("startup")
def start_kaspi_job_recovery():
    if not os.getenv("PYTEST_CURRENT_TEST"):
        threading.Thread(target=_kaspi_job_loop, name="kaspi-recovery", daemon=True).start()
WB_DEFAULT_SYNC_SECTIONS = ("new_orders", "order", "sale")
WB_ALLOWED_SYNC_SECTIONS = frozenset(("new_orders", "order", "sale", "realization", "stocks", "products", "order_statuses"))


def _wb_weekly_loop():
    """Check the most recently completed week on a bounded server-side cadence."""
    from . import wb_weekly
    while True:
        try:
            start, end = wb_weekly.previous_week()
            # Fetch only the last completed week. Older periods are user-selected;
            # an automatic catch-up loop would spend the seller's strict Finance API quota.
            with SessionLocal() as db:
                accounts = db.scalars(select(Account).where(Account.channel == "Wildberries", Account.mode == "api", Account.active.is_(True)).order_by(Account.id)).all()
                system_owner = db.scalar(select(User).where(User.role == "Owner", User.active.is_(True)).order_by(User.id))
                for account in accounts:
                    try:
                        credentials = load_saved_credentials(db, account.id)
                        if not credentials.get("token"):
                            continue
                        wb_weekly.sync_week(db, account, credentials["token"], start, end,
                                            project_user=system_owner, trigger="scheduler")
                    except Exception as exc:
                        db.rollback()
                        logger.warning("WB weekly check skipped for account %s (%s)", account.id, type(exc).__name__)
        except Exception:
            logger.exception("WB weekly scheduler cycle failed")
        time.sleep(max(300, int(os.getenv("WB_WEEKLY_CHECK_SECONDS", "1800"))))


@app.on_event("startup")
def start_wb_weekly_scheduler():
    # Each web worker has a small daemon loop; the report row lease/DB lock makes
    # concurrent checks safe across Render workers and repeated deployments.
    thread = threading.Thread(target=_wb_weekly_loop, name="wb-weekly-reports", daemon=True)
    thread.start()


@app.middleware("http")
async def security_headers(request, call_next):
    if int(request.headers.get("content-length", "0") or 0) > 6 * 1024 * 1024:
        return JSONResponse({"detail": "Максимум 6 МБ"}, status_code=413)
    response = await call_next(request)
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["Referrer-Policy"] = "same-origin"
    response.headers["Cache-Control"] = (
        "no-store" if request.url.path.startswith("/api") else "no-cache"
    )
    response.headers["Content-Security-Policy"] = (
        "default-src 'self'; img-src 'self' data: https:; style-src 'self' 'unsafe-inline'; script-src 'self'; connect-src 'self'; frame-ancestors 'none'; base-uri 'self'"
    )
    return response


@app.exception_handler(IntegrityError)
async def integrity_handler(request, exc):
    return JSONResponse(
        status_code=409,
        content={
            "detail": "Дубликат записи либо запись используется в других документах"
        },
    )


@app.exception_handler(ValidationError)
async def validation_handler(request, exc):
    return JSONResponse(
        status_code=422,
        content={"code": "VALIDATION_ERROR", "message": "Проверьте заполнение полей.", "detail": "Проверьте заполнение полей."},
    )


@app.exception_handler(RequestValidationError)
async def request_validation_handler(request, exc):
    """Keep Pydantic internals out of API responses and browser error surfaces."""
    return JSONResponse(
        status_code=422,
        content={"code": "VALIDATION_ERROR", "message": "Проверьте параметры запроса.", "detail": "Проверьте параметры запроса."},
    )


def dump(obj):
    if isinstance(obj, dict):
        return jsonable_encoder(obj)
    result = jsonable_encoder({
        c.name: getattr(obj, c.name)
        for c in obj.__table__.columns
        if c.name not in ("password_hash", "token_hash", "path", "rows")
    })
    if isinstance(obj,Sale) and obj.cogs_pending: result["cogs"]=None
    if isinstance(obj, Product):
        specs = obj.specifications or {}
        result["description"] = specs.get("description", "")
        result["barcode"] = specs.get("barcode", "")
        result["purchase_cost"] = int(specs.get("purchase_cost") or 0)
        result["custom_fields"] = specs.get("custom_fields") or {}
        result["marketplace_stock_allocations"] = specs.get("marketplace_stock_allocations") or {}
    if isinstance(obj, Account):
        result.pop("integration_settings", None)
    return result


class Login(BaseModel):
    email: str = Field(max_length=200)
    password: str = Field(max_length=128)


@app.post("/api/auth/login")
def login(p: Login, request: Request, response: Response, db=Depends(get_db)):
    identity = hashlib.sha256(
        (
            p.email.lower() + ":" + (request.client.host if request.client else "")
        ).encode()
    ).hexdigest()
    count = db.scalar(
        select(func.count())
        .select_from(LoginAttempt)
        .where(
            LoginAttempt.identity == identity,
            LoginAttempt.created_at > now() - timedelta(minutes=15),
        )
    )
    if count >= 8:
        raise HTTPException(429, "Слишком много попыток. Повторите через 15 минут")
    db.add(LoginAttempt(identity=identity))
    db.commit()
    u = db.scalar(select(User).where(User.email == p.email.lower()))
    if not u or not u.active or not passwords.verify(p.password, u.password_hash):
        raise HTTPException(401, "Неверный логин или пароль")
    token = secrets.token_urlsafe(48)
    db.add(
        Session(
            token_hash=digest(token),
            user_id=u.id,
            expires_at=now() + timedelta(hours=12),
        )
    )
    audit(db, u, "login", "users", u.id)
    db.commit()
    response.set_cookie(
        "erp_session",
        token,
        httponly=True,
        secure=os.getenv("COOKIE_SECURE", "true").lower() == "true",
        samesite="strict",
        max_age=43200,
        path="/",
    )
    return {"user": dump(u), "csrf": token, "permissions": list(PERMISSIONS[u.role])}


@app.get("/api/auth/me")
def me(request: Request, u=Depends(current_user)):
    return {
        "user": dump(u),
        "csrf": request.cookies.get("erp_session"),
        "permissions": list(PERMISSIONS[u.role]),
    }


@app.post("/api/auth/logout")
def logout(
    request: Request, response: Response, u=Depends(current_user), db=Depends(get_db)
):
    sess = db.scalar(
        select(Session).where(
            Session.token_hash == digest(request.cookies["erp_session"])
        )
    )
    if sess:
        db.delete(sess)
    db.commit()
    response.delete_cookie("erp_session")
    return {"ok": True}


@app.post("/api/auth/password")
def change_password(payload: dict, u=Depends(current_user), db=Depends(get_db)):
    if not passwords.verify(str(payload.get("old", "")), u.password_hash):
        raise HTTPException(403, "Неверный текущий пароль")
    new = str(payload.get("new", ""))
    if len(new) < 12 or len(new) > 128:
        raise HTTPException(422, "Пароль от 12 до 128 символов")
    u.password_hash = passwords.hash(new)
    for sess in db.scalars(select(Session).where(Session.user_id == u.id)):
        db.delete(sess)
    audit(db, u, "change_password", "users", u.id)
    db.commit()
    return {"ok": True}




DOC_KIND_LABELS = {
    "contract": "Договор", "invoice": "Счёт", "waybill": "Накладная",
    "act": "Акт", "commercial_offer": "Коммерческое предложение",
    "payment_invoice": "Счёт на оплату",
}

@app.get("/api/generated-documents/{doc_id}/pdf")
def generated_document_pdf(doc_id: int, u=Depends(current_user), db=Depends(get_db)):
    require(u, "documents")
    doc = db.get(GeneratedDocument, doc_id)
    if not doc:
        raise HTTPException(404, "Документ не найден")
    scope(u, doc.company_id)
    company = db.get(Company, doc.company_id)
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib import colors
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
    except ImportError:
        raise HTTPException(503, "PDF renderer unavailable")
    font = "Helvetica"
    for candidate in ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"):
        if os.path.exists(candidate):
            pdfmetrics.registerFont(TTFont("QORUnicode", candidate)); font = "QORUnicode"; break
    buf = io.BytesIO()
    pdf = SimpleDocTemplate(buf, pagesize=A4, rightMargin=42, leftMargin=42, topMargin=42, bottomMargin=42, title=f"{DOC_KIND_LABELS.get(doc.kind, doc.kind)} {doc.number}")
    styles = getSampleStyleSheet(); styles["Title"].fontName = font; styles["Normal"].fontName = font
    styles["Title"].alignment = TA_CENTER
    story = [Paragraph(DOC_KIND_LABELS.get(doc.kind, doc.kind), styles["Title"]), Spacer(1, 12)]
    rows = [
        ["Номер", doc.number], ["Дата", str(doc.date)], ["Компания", company.name if company else str(doc.company_id)],
        ["БИН/ИИН", company.bin if company else ""], ["Контрагент", doc.counterparty or "—"],
        ["Сумма", f"{doc.amount/100:,.2f} ₸".replace(",", " ")], ["Статус", doc.status],
    ]
    table = Table(rows, colWidths=[120, 350])
    table.setStyle(TableStyle([("FONTNAME", (0,0), (-1,-1), font), ("GRID", (0,0), (-1,-1), .4, colors.grey), ("BACKGROUND", (0,0), (0,-1), colors.whitesmoke), ("VALIGN", (0,0), (-1,-1), "TOP"), ("PADDING", (0,0), (-1,-1), 7)]))
    story += [table, Spacer(1, 16)]
    payload = doc.payload or {}
    if payload:
        story.append(Paragraph("Дополнительные данные", styles["Normal"])); story.append(Spacer(1, 6))
        extra = [[str(k), str(v)] for k,v in payload.items() if v not in (None, "", [], {})]
        if extra:
            t2 = Table(extra, colWidths=[150, 320]); t2.setStyle(TableStyle([("FONTNAME", (0,0), (-1,-1), font), ("GRID", (0,0), (-1,-1), .3, colors.lightgrey), ("VALIGN", (0,0), (-1,-1), "TOP"), ("PADDING", (0,0), (-1,-1), 6)])); story.append(t2)
    story += [Spacer(1, 24), Paragraph("Сформировано в QOR AI Business ERP. Перед подписанием документ должен быть проверен ответственным сотрудником.", styles["Normal"])]
    pdf.build(story); buf.seek(0)
    safe = "".join(c for c in doc.number if c.isalnum() or c in "-_") or str(doc.id)
    return StreamingResponse(buf, media_type="application/pdf", headers={"Content-Disposition": f'attachment; filename="QOR-{safe}.pdf"'})


@app.get("/api/health")
def health(db=Depends(get_db)):
    db.execute(select(1))
    try:
        from alembic.config import Config
        from alembic.script import ScriptDirectory
        expected_revision = ScriptDirectory.from_config(Config("alembic.ini")).get_current_head()
        revision = db.execute(text("SELECT version_num FROM alembic_version")).scalar_one_or_none()
    except Exception:
        expected_revision = revision = None
    if revision != expected_revision:
        return JSONResponse(
            status_code=503,
            content={"status": "degraded", "database": "ok", "migration": "incompatible"},
        )
    return {"status": "ok", "database": "ok", "migration": "compatible"}


MODELS = {
    "companies": (Company, "companies", s.CompanyIn),
    "products": (Product, "catalog", s.ProductIn),
    "accounts": (Account, "accounts", s.AccountIn),
    "customers": (Customer, "sales", s.CustomerIn),
    "purchases": (Purchase, "inventory", None),
    "batches": (Batch, "inventory", None),
    "movements": (Movement, "inventory", None),
    "sales": (Sale, "sales", None),
    "returns": (Return, "sales", None),
    "expenses": (Expense, "finance", None),
    "cash": (CashEntry, "finance", None),
    "documents": (Document, "documents", None),
    "audit": (Audit, "audit", None),
    "users": (User, "users", s.UserIn),
    "jobs": (Job, "documents", None),
    "imports": (ImportRun, "imports", None),
    "niches": (Niche, "marketing", s.NicheIn),
    "counterparties": (Counterparty, "finance", s.CounterpartyIn),
    "bank_transactions": (BankTransaction, "finance", s.BankTransactionIn),
    "tax_rules": (TaxRule, "finance", s.TaxRuleIn),
    "generated_documents": (GeneratedDocument, "documents", s.GeneratedDocumentIn),
}


def model_info(entity):
    if entity not in MODELS:
        raise HTTPException(404, "Раздел не найден")
    return MODELS[entity]


def access_query(db, u, entity, search="", company_id=None, account_id=None):
    model, cap, _ = model_info(entity)
    require(u, cap)
    q = scope_query(select(model), model, u)
    if entity == "companies" and u.role != "Owner":
        q = q.where(Company.id.in_(u.company_ids))
    if entity == "users" and u.role != "Owner":
        q = q.where(User.role != "Owner")
    if entity == "imports":
        q = q.where(ImportRun.user_id == u.id)
    if entity == "jobs":
        q = q.where(Job.user_id == u.id)
    # Archived marketplace accounts stay available in the account manager for
    # reconnecting, while their old sales and raw events leave active views.
    if entity == "sales":
        q = q.where(Sale.account_id.in_(select(Account.id).where(Account.active.is_(True))))
    elif entity == "returns":
        q = q.join(Sale, Return.sale_id == Sale.id).where(
            Sale.account_id.in_(select(Account.id).where(Account.active.is_(True)))
        )
    elif entity == "expenses":
        q = q.where(or_(Expense.account_id.is_(None), Expense.account_id.in_(select(Account.id).where(Account.active.is_(True)))))
    if company_id and hasattr(model, "company_id"):
        scope(u, company_id)
        q = q.where(model.company_id == company_id)
    if account_id and hasattr(model, "account_id"):
        q = q.where(model.account_id == account_id)
    if search:
        cols = [
            getattr(model, k)
            for k in (
                "name",
                "sku",
                "reference",
                "description",
                "email",
                "action",
                "supplier",
            )
            if hasattr(model, k)
        ]
        if cols:
            search_terms = [c.ilike("%" + search[:100] + "%") for c in cols]
            if entity == "products":
                search_terms.append(cast(Product.specifications, String).ilike("%" + search[:100] + "%"))
            q = q.where(or_(*search_terms))
    return q


@app.get("/api/marketplace-records")
def marketplace_records(
    channel: str | None = None,
    company_id: int | None = None,
    account_id: int | None = None,
    status: str | None = None,
    offset: int = Query(0, ge=0),
    limit: int = Query(50, ge=1, le=500),
    u=Depends(current_user),
    db=Depends(get_db),
):
    """Return idempotent raw/staged records received from marketplace APIs.

    Marketplace ORDER/POSTING records are intentionally not booked as financial
    sales until a confirmed sale/payment/return event with sufficient accounting
    data exists. This endpoint makes successful sync visible instead of hiding it
    behind the financial sales table.
    """
    q = select(MarketplaceRecord).join(Account, Account.id == MarketplaceRecord.account_id).where(Account.active.is_(True))
    if u.role != "Owner":
        q = q.where(MarketplaceRecord.company_id.in_(u.company_ids or []))
    if company_id is not None:
        scope(u, company_id)
        q = q.where(MarketplaceRecord.company_id == company_id)
    if channel:
        q = q.where(Account.channel == channel)
    if account_id is not None:
        q = q.where(MarketplaceRecord.account_id == account_id)
    if status:
        q = q.where(or_(MarketplaceRecord.status.startswith("pending"),MarketplaceRecord.stage_status.in_(["pending_fifo","pending_fx","pending_mapping"])) if status == "pending" else MarketplaceRecord.status == status)
    count = db.scalar(select(func.count()).select_from(q.subquery())) or 0
    rows=[]
    for rec in db.scalars(q.order_by(MarketplaceRecord.id.desc()).offset(offset).limit(limit)):
        payload = rec.payload or {}
        # Presentation is enriched from the canonical product mapping. Marketplace
        # sales payloads (notably WB statistics) often carry sellerArticle/nmId
        # but omit a human product name even though the catalogue endpoint has it.
        # Do not force users to re-map data that QOR already mapped during projection.
        display_product = None
        display_keys = marketplace_pipeline.identifiers(payload)
        if display_keys:
            mapped_ids = db.scalars(select(MarketplaceProductMapping.product_id).where(
                MarketplaceProductMapping.account_id == rec.account_id,
                MarketplaceProductMapping.company_id == rec.company_id,
                MarketplaceProductMapping.external_key.in_(display_keys)
            )).all()
            unique_ids = list(dict.fromkeys(mapped_ids))
            if len(unique_ids) == 1:
                display_product = db.get(Product, unique_ids[0])
        inferred_quantity = payload.get("quantity")
        if inferred_quantity is None and rec.source == "wildberries" and payload.get("record_type") == "sale":
            inferred_quantity = 1
        rows.append({
            "id": rec.id, "account_id": rec.account_id, "company_id": rec.company_id,
            "source": rec.source, "external_id": rec.external_id, "status": rec.status,
            "stage_status":rec.stage_status,"reason_code":rec.reason_code,
            "created_at": rec.created_at,
            "marketplace_status": payload.get("status") or payload.get("state") or payload.get("substatus") or "received",
            "reference": payload.get("code") or payload.get("posting_number") or payload.get("rid") or rec.external_id,
            "amount": next((payload[k] for k in ("total_price", "retailPriceWithDisc", "priceWithDisc", "finishedPrice", "amount", "retail_amount") if payload.get(k) is not None), None),
            "product_name": payload.get("product_name") or payload.get("name") or payload.get("title") or payload.get("subjectName") or ((payload.get("products") or [{}])[0].get("name") if payload.get("products") else None) or (display_product.name if display_product else None),
            "sku": payload.get("supplierArticle") or payload.get("vendorCode") or payload.get("offer_id") or payload.get("sku") or payload.get("nmId") or ((payload.get("products") or [{}])[0].get("offer_id") if payload.get("products") else None),
            "currency": payload.get("currency") or payload.get("currency_code") or ((payload.get("products") or [{}])[0].get("currency_code") if payload.get("products") else None) or ("KZT" if rec.source == "kaspi" else "RUB"),
            "quantity": inferred_quantity or sum(int(x.get("quantity") or 0) for x in (payload.get("products") or [])) or None,
            "projection_reason": payload.get("_qor_projection_reason", ""),
            "external_keys": list(dict.fromkeys(marketplace_pipeline.identifiers(payload) + [key for x in payload.get("products", []) for key in marketplace_pipeline.identifiers(x)])),
            "payload": payload,
        })
    return {"items": jsonable_encoder(rows), "total": count}


@app.get("/api/marketplace-orders")
def marketplace_orders_report(channel: str = Query(...), company_id: int | None = None,
    account_id: int | None = None, status_filter: str = "all",
    limit: int = Query(100, ge=1, le=500), offset: int = Query(0, ge=0),
    u=Depends(current_user), db=Depends(get_db)):
    """Marketplace order lifecycle view; values are operational, not booked cash."""
    require(u, "analytics")
    if channel not in ("Kaspi", "Wildberries"): raise HTTPException(422, "Поддерживаются Kaspi и Wildberries")
    aq=select(Account).where(Account.channel==channel,Account.active.is_(True))
    if u.role!="Owner": aq=aq.where(Account.company_id.in_(u.company_ids or []))
    if company_id is not None: scope(u,company_id); aq=aq.where(Account.company_id==company_id)
    if account_id is not None: aq=aq.where(Account.id==account_id)
    accounts=list(db.scalars(aq)); aids=[a.id for a in accounts]
    if not aids: return {"items":[],"total":0,"summary":{},"stores":[]}
    types=("order","order_entry") if channel=="Kaspi" else ("assembly_order","new_order","realization")
    records=list(db.scalars(select(MarketplaceRecord).where(MarketplaceRecord.account_id.in_(aids),
        MarketplaceRecord.payload["record_type"].as_string().in_(types)).order_by(MarketplaceRecord.id.desc())))
    accs={a.id:a for a in accounts}; rows=[]
    if channel=="Kaspi":
        orders={}; lines={}
        for rec in records:
            p=rec.payload or {}; oid=str(p.get("order_id") or "")
            if p.get("record_type")=="order": orders[(rec.account_id,oid)]=(rec,p)
            elif p.get("record_type")=="order_entry": lines.setdefault((rec.account_id,oid),[]).append(p)
        names={"APPROVED_BY_BANK":"Ожидает продавца","ACCEPTED_BY_MERCHANT":"Принят продавцом","COMPLETED":"Выдан покупателю","CANCELLED":"Отменён","CANCELLING":"Отмена в процессе","KASPI_DELIVERY_RETURN_REQUESTED":"Запрошен возврат","RETURNED":"Возвращён"}
        for (aid,oid),(rec,p) in orders.items():
            st=str(p.get("status") or ""); state=str(p.get("state") or "")
            label=names.get(st,st or "Статус не передан")
            if st=="ACCEPTED_BY_MERCHANT" and state: label += " · "+("Kaspi доставка" if state=="KASPI_DELIVERY" else "доставка продавца" if state=="DELIVERY" else state)
            group="returns" if st in ("CANCELLING","KASPI_DELIVERY_RETURN_REQUESTED","RETURNED") else "canceled" if st=="CANCELLED" else "delivered" if st=="COMPLETED" else "in_delivery" if st=="ACCEPTED_BY_MERCHANT" else "waiting"
            for line in lines.get((aid,oid)) or [None]:
                l=line or {}; rows.append({"account_id":aid,"store":accs[aid].name,"order_id":oid,"date":p.get("creation_date"),"product":l.get("product_name") or "—","sku":l.get("sku") or "","quantity":l.get("quantity"),"amount":l.get("total_price") or p.get("total_price"),"status":label,"group":group,"erp_status":"Проведено в ERP" if rec.status=="projected" else ("Ожидает проведения в ERP" if st=="COMPLETED" else "Не проводится до завершения заказа"),"finance_status":"Kaspi API заказа не подтверждает денежное зачисление","finance_amount":None})
    else:
        orders={}; finance={}
        for rec in records:
            p=rec.payload or {}; typ=p.get("record_type")
            if typ in ("assembly_order","new_order"):
                key=(rec.account_id,str(p.get("id") or p.get("rid") or ""))
                if typ=="assembly_order" or key not in orders: orders[key]=(rec,p)
            elif typ=="realization":
                k=str(p.get("srid") or p.get("rid") or "")
                if k: finance[(rec.account_id,k)]=p
        sup={"new":"Новый заказ","confirm":"Подтверждён, сборка","complete":"Передан в доставку","cancel":"Отменён продавцом","cancel_carrier":"Отменён перевозчиком"}
        wb={"waiting":"Ожидает передачи WB","sorted":"Сортируется","sold":"Выкуплен","canceled":"Отменён","canceled_by_client":"Отменён клиентом","declined_by_client":"Отклонён клиентом","defect":"Брак","ready_for_pickup":"В ПВЗ, ждёт покупателя","postponed_delivery":"Доставка перенесена","accepted_by_carrier":"Принят перевозчиком","sent_to_carrier":"Передан перевозчику","canceled_by_carrier":"Отменён перевозчиком"}
        for (aid,oid),(rec,p) in orders.items():
            s=str(p.get("supplierStatus") or ""); w=str(p.get("wbStatus") or ""); label=wb.get(w) or sup.get(s) or "Статус пока не передан"
            group="returns" if w=="defect" else "canceled" if "cancel" in (w or s) else "delivered" if w in ("sold","ready_for_pickup") else "in_delivery" if w in ("sorted","accepted_by_carrier","sent_to_carrier","postponed_delivery") or s=="complete" else "waiting"
            fp=finance.get((aid,str(p.get("rid") or "")))
            rows.append({"account_id":aid,"store":accs[aid].name,"order_id":str(p.get("id") or oid),"date":p.get("createdAt"),"product":p.get("article") or "—","sku":", ".join(map(str,p.get("skus") or [])),"quantity":1,"amount":None,"status":label,"group":group,"erp_status":"Заказ не считается продажей до подтверждения финансовой строки","finance_status":"Строка отчёта WB найдена" if fp else ("Ожидает строки финансового отчёта WB" if group=="delivered" else "Нет финансовой строки"),"finance_amount":(fp.get("forPay") or fp.get("ppvzForPay") or fp.get("retailAmount")) if fp else None})
    summary={k:sum(1 for r in rows if r["group"]==k) for k in ("in_delivery","delivered","waiting","canceled","returns")}
    if status_filter!="all": rows=[r for r in rows if r["group"]==status_filter]
    rows.sort(key=lambda r:str(r.get("date") or ""),reverse=True); total=len(rows)
    return {"items":rows[offset:offset+limit],"total":total,"summary":summary,"stores":[{"id":a.id,"name":a.name,"sync_status":a.sync_status,"last_sync":a.last_sync} for a in accounts]}


@app.get("/api/marketplace-products")
def marketplace_products(channel: str, company_id: int | None = None, account_id: int | None = None, start: date | None = None, end: date | None = None, limit: int = Query(100, ge=1, le=500), u=Depends(current_user), db=Depends(get_db)):
    """Fast catalogue cards. Avoid N+1 queries and never recalculate the whole ERP report."""
    if channel not in {"Kaspi", "Wildberries", "Ozon"}:
        raise HTTPException(422, "Неподдерживаемый маркетплейс")
    aq = select(Account).where(Account.channel == channel, Account.active.is_(True))
    if u.role != "Owner": aq = aq.where(Account.company_id.in_(u.company_ids or []))
    if company_id is not None:
        scope(u, company_id); aq = aq.where(Account.company_id == company_id)
    accounts = list(db.scalars(aq)); account_ids = [a.id for a in accounts]
    account_by_id = {a.id: a for a in accounts}
    if account_id is not None:
        if account_id not in account_by_id: raise HTTPException(404, "Кабинет маркетплейса не найден")
        account_ids = [account_id]
        account_by_id = {account_id: account_by_id[account_id]}
    if not account_ids: return {"items": [], "total": 0}

    mappings = list(db.scalars(select(MarketplaceProductMapping).where(MarketplaceProductMapping.account_id.in_(account_ids))))
    by_product = {}
    account_keys = {}
    for m in mappings:
        by_product.setdefault(m.product_id, set()).add(m.external_key)
        account_keys.setdefault((m.product_id, m.account_id), set()).add(m.external_key)
    product_ids = list(by_product)[:limit]
    if not product_ids:
        # Kaspi can return products attached to orders even when the seller has
        # not linked them to an ERP product yet. Keep those visible as observed
        # marketplace items instead of showing an empty catalogue.
        if channel == "Kaspi":
            observed = {}
            raw_q = select(MarketplaceRecord.account_id, MarketplaceRecord.payload).where(
                MarketplaceRecord.account_id.in_(account_ids), MarketplaceRecord.source == "kaspi")
            for account_id, payload in db.execute(raw_q):
                p = payload or {}
                if p.get("record_type") != "order_entry":
                    continue
                try:
                    day = marketplace_pipeline._day(p.get("order_creation_date"))
                    if start is not None and day < start: continue
                    if end is not None and day > end: continue
                except (ValueError, TypeError, KeyError):
                    day = None
                attrs = p.get("merchant_product_attributes") or p.get("product_attributes") or {}
                external = str(p.get("external_product_id") or p.get("sku") or p.get("product_name") or p.get("external_id"))
                images = marketplace_pipeline._media_urls(p)
                row = observed.setdefault(external, {"id":"kaspi:"+external,"sku":p.get("sku"),"article":p.get("sku"),
                    "name":p.get("product_name") or attrs.get("name") or "Kaspi товар",
                    "brand":attrs.get("brand"),"category":attrs.get("category"),"image_url":images[0] if images else None,
                    "images":images,"account_id":account_id,"account_name":account_by_id[account_id].name,"currency":getattr(account_by_id[account_id], "currency", None) or "KZT","sold_units":0,"revenue":0,
                    "lifetime_sold_units":0,"lifetime_revenue":0,"unconverted_sales":0,"cogs":0,"commission":0,
                    "logistics":0,"advertising":0,"profit":None,"stock_quantity":None,"specifications":{},
                    "source":"Kaspi order items","orders":[]})
                try:
                    qty = int(p.get("quantity") or 0)
                    amount = float(p.get("total_price") or 0)
                except (TypeError, ValueError):
                    qty, amount = 0, 0
                row["lifetime_sold_units"] += qty
                row["lifetime_revenue"] += amount
                row["orders"].append({"order_id":p.get("order_id"),"status":p.get("order_status"),
                    "date":day.isoformat() if day else None,"quantity":qty,"total_price":amount})
                if (start is None or day is None or day >= start) and (end is None or day is None or day <= end):
                    row["sold_units"] += qty
                    row["revenue"] += amount
            rows = sorted(observed.values(), key=lambda x:(x["revenue"],x["sold_units"]), reverse=True)[:limit]
            return {"items":jsonable_encoder(rows),"total":len(observed),"catalog_scope":"products appearing in synced orders"}
        return {"items": [], "total": 0}
    products = {p.id: p for p in db.scalars(select(Product).where(Product.id.in_(product_ids)))}
    sales = list(db.scalars(select(Sale).where(Sale.product_id.in_(product_ids), Sale.account_id.in_(account_ids))))
    sales_by_product = {}
    for sale in sales: sales_by_product.setdefault(sale.product_id, []).append(sale)

    # Build external-key -> product once, then scan WB sale records once. This
    # replaces product_count × raw_record_count work from the old endpoint.
    key_to_pid = {}
    for pid, keys in by_product.items():
        for key in keys: key_to_pid[str(key)] = pid
    wb_stats = {pid: {"pq": 0, "pr": 0, "lq": 0, "lr": 0, "bad": 0} for pid in product_ids}
    if channel == "Wildberries":
        raw_q = select(MarketplaceRecord.account_id, MarketplaceRecord.payload).where(
            MarketplaceRecord.account_id.in_(account_ids), MarketplaceRecord.source == "wildberries")
        for account_id, payload in db.execute(raw_q):
            p = payload or {}
            if str(p.get("record_type") or "").lower() != "sale": continue
            if str(p.get("saleID") or p.get("saleId") or "").upper().startswith("R") or p.get("isStorno"): continue
            pid = next((key_to_pid.get(str(k)) for k in marketplace_pipeline.identifiers(p) if str(k) in key_to_pid), None)
            if pid is None or pid not in wb_stats: continue
            st = wb_stats[pid]
            try:
                day = marketplace_pipeline._day(p.get("date"))
                amount = p.get("priceWithDisc") if p.get("priceWithDisc") is not None else p.get("finishedPrice")
                value = marketplace_pipeline._cents(amount, account_by_id[account_id], p.get("currency") or "RUB", db, p.get("date"))
            except (marketplace_pipeline.Pending, ValueError, TypeError, KeyError):
                st["bad"] += 1; continue
            st["lq"] += 1; st["lr"] += value
            if (start is None or day >= start) and (end is None or day <= end):
                st["pq"] += 1; st["pr"] += value

    out = []
    for pid in product_ids:
        product = products.get(pid)
        if not product: continue
        all_sales = sales_by_product.get(pid, [])
        period_sales = [x for x in all_sales if (start is None or x.date >= start) and (end is None or x.date <= end)]
        sale_costs, _, pending_ids = marketplace_costs.resolve_cost_overrides(db, period_sales, [], account_by_id)
        cogs = sum(sale_costs.get(x.id, x.cogs) for x in period_sales); commission = sum(x.commission for x in period_sales)
        logistics = sum(x.logistics for x in period_sales); advertising = sum(x.advertising for x in period_sales)
        if channel == "Wildberries":
            st = wb_stats[pid]; period_qty, period_revenue = st["pq"], st["pr"]
            lifetime_qty, lifetime_revenue, unconverted = st["lq"], st["lr"], st["bad"]
        else:
            period_qty = sum(x.quantity for x in period_sales); period_revenue = sum(x.revenue for x in period_sales)
            lifetime_qty = sum(x.quantity for x in all_sales); lifetime_revenue = sum(x.revenue for x in all_sales); unconverted = 0
        sku_specs = product.specifications or {}
        sku_unit_cost = int(sku_specs.get("purchase_cost") or 0)
        legacy_cost = marketplace_costs.latest_cost(db, account_ids[0], account_keys.get((pid, account_ids[0]), set()), end or business_today())
        effective_history = [x for x in (sku_specs.get("purchase_cost_history") or []) if isinstance(x, dict)]
        cost_day = end or business_today()
        dated_sku_costs = []
        for item in effective_history:
            try:
                effective = date.fromisoformat(str(item.get("effective_from")))
                unit_cost = int(item.get("unit_cost"))
                if effective <= cost_day:
                    dated_sku_costs.append((effective, unit_cost))
            except (AttributeError, TypeError, ValueError):
                continue
        sku_cost_at_period = max(dated_sku_costs, default=(None, sku_unit_cost), key=lambda item: item[0] or date.min)[1]
        sku_cost_history_applies = bool(dated_sku_costs)
        legacy_cost_applies = not sku_cost_history_applies and sku_unit_cost <= 0
        out.append({"id": product.id, "sku": product.sku, "article": product.article, "name": product.name, "brand": product.brand, "category": product.category,
            "image_url": (product.media or [None])[0], "images": product.media or [], "price": product.price, "currency": product.currency,
            "sold_units": period_qty, "revenue": period_revenue, "lifetime_sold_units": lifetime_qty, "lifetime_revenue": lifetime_revenue, "unconverted_sales": unconverted,
            "cogs": cogs, "commission": commission, "logistics": logistics, "advertising": advertising,
            "profit": None if any(x.id in pending_ids for x in period_sales) else sum(x.revenue for x in period_sales)-cogs-commission-logistics-advertising,
            "account_id": account_ids[0],
            # SKU purchase cost is the normal entry point. Legacy account-specific
            # history remains a read-only fallback for earlier sales.
            "unit_cost": (sku_cost_at_period if sku_cost_at_period > 0 else None) if sku_cost_history_applies else (legacy_cost.unit_cost if legacy_cost_applies and legacy_cost else (sku_unit_cost or None)),
            "cost_source": "SKU" if sku_cost_history_applies or sku_unit_cost else ("история кабинета" if legacy_cost_applies and legacy_cost else None),
            "cost_effective_from": str(max(dated_sku_costs, key=lambda item: item[0])[0]) if dated_sku_costs else (str(legacy_cost.effective_from) if legacy_cost_applies and legacy_cost else None),
            "stock_quantity": sum(int(v.get("quantity") or 0) for k,v in ((product.specifications or {}).get("marketplace_stock") or {}).items() if k.startswith(channel+":")),
            "specifications": product.specifications or {}})
    out.sort(key=lambda x: (x["revenue"], x["sold_units"]), reverse=True)
    return {"items": jsonable_encoder(out), "total": len(out)}



@app.put("/api/marketplace-products/cost")
def marketplace_product_cost_save(payload: dict, u=Depends(current_user), db=Depends(get_db)):
    require(u, "finance", True)
    try:
        account_id = int(payload.get("account_id"))
        product_id = int(payload.get("product_id"))
        unit_cost = Decimal(str(payload.get("unit_cost"))).quantize(Decimal("0.01"))
        effective_from = date.fromisoformat(str(payload.get("effective_from")))
    except (TypeError, ValueError, ArithmeticError):
        raise HTTPException(422, "Укажите кабинет, себестоимость и дату начала действия")
    if not unit_cost.is_finite() or unit_cost <= 0 or unit_cost > Decimal("1000000000000"):
        raise HTTPException(422, "Себестоимость должна быть больше нуля")
    account = db.get(Account, account_id)
    product = db.get(Product, product_id)
    if not account or account.channel not in {"Kaspi", "Wildberries", "Ozon"} or not product:
        raise HTTPException(404, "Товар или кабинет маркетплейса не найден")
    scope(u, account.company_id)
    mappings = db.scalars(select(MarketplaceProductMapping).where(
        MarketplaceProductMapping.account_id == account.id,
        MarketplaceProductMapping.company_id == account.company_id,
        MarketplaceProductMapping.product_id == product.id,
    )).all()
    if not mappings:
        raise HTTPException(422, "Сначала свяжите карточку Kaspi с товаром ERP")
    amount = int(unit_cost * 100)
    for mapping in mappings:
        row = db.scalar(select(MarketplaceProductCost).where(
            MarketplaceProductCost.account_id == account.id,
            MarketplaceProductCost.external_key == mapping.external_key,
            MarketplaceProductCost.effective_from == effective_from,
        ))
        if row:
            row.unit_cost, row.updated_by = amount, u.id
        else:
            db.add(MarketplaceProductCost(
                company_id=account.company_id, account_id=account.id, external_key=mapping.external_key,
                product_name=product.name, unit_cost=amount, currency="KZT",
                effective_from=effective_from, updated_by=u.id,
            ))
    db.flush()
    audit(db, u, "update", "marketplace_product_costs", product.id, {
        "account_id": account.id, "product_id": product.id, "unit_cost": str(unit_cost),
        "effective_from": str(effective_from), "external_keys": [m.external_key for m in mappings],
    })
    db.commit()
    return {"saved": True, "account_id": account.id, "product_id": product.id,
            "unit_cost": str(unit_cost), "effective_from": str(effective_from)}


def _kaspi_commission_estimate(sales, expenses, api_bases, configured_rate=None):
    """Estimate missing Kaspi commission only from completed API order entries.

    A normal ERP Sale does not prove that Kaspi marked the order as issued.
    Explicit commission data for the same account/day takes precedence.
    """
    from decimal import Decimal, InvalidOperation

    raw_rate = configured_rate if configured_rate is not None else os.getenv("KASPI_COMMISSION_RATE")
    if raw_rate in (None, ""):
        return 0, None, False
    try:
        rate = Decimal(str(raw_rate))
        if not rate.is_finite() or rate < 0 or rate > 1:
            raise InvalidOperation
    except (InvalidOperation, ValueError, TypeError):
        return 0, None, False
    basis_points = int(rate * 10000)
    explicit_days = {
        (expense.account_id, expense.date)
        for expense in expenses
        if expense.category in {"Marketplace commission", "Marketplace commission VAT"}
    }
    sale_commission_days = {
        (sale.account_id, sale.date) for sale in sales if sale.commission > 0
    }
    actual_accounts = {account_id for account_id, _day in explicit_days | sale_commission_days}
    skipped = explicit_days | sale_commission_days
    estimated = 0
    suppressed = False
    for key, base in api_bases.items():
        if key in skipped:
            continue
        amount = (max(0, int(base or 0)) * basis_points + 5000) // 10000
        if key[0] in actual_accounts:
            suppressed = suppressed or amount > 0
        else:
            estimated += amount
    return estimated, str(rate), suppressed


def _return_statement_mapping(headers):
    aliases = {
        "external_return_id": ("return_id", "id возврата", "номер возврата", "rrd_id", "rrdid", "идентификатор операции", "operation id"),
        "external_order_id": ("order_id", "order", "заказ", "номер заказа", "номер исходного заказа", "srid"),
        "financial_return_date": ("return date", "дата возврата", "дата операции", "rrd_date", "rrddate", "дата выплаты"),
        "amount": ("refund", "refund amount", "сумма возврата", "сумма операции", "amount", "forpay"),
        "sku": ("sku", "артикул продавца", "vendorcode", "vendor code", "артикул"),
        "product_name": ("товар", "название товара", "наименование", "product name"),
        "quantity": ("количество", "quantity", "шт."),
        "reason": ("причина", "причина возврата", "reason"),
        "currency": ("валюта", "currency"),
    }
    normalized = {str(h).strip().lower().replace(" ", "_"): h for h in headers}
    result = {}
    for field, variants in aliases.items():
        for header in headers:
            key = str(header).strip().lower()
            if key in variants or key.replace(" ", "_") in variants:
                result[field] = header
                break
        if field not in result:
            for variant in variants:
                if variant.replace(" ", "_") in normalized:
                    result[field] = normalized[variant.replace(" ", "_")]
                    break
    return result


@app.post("/api/marketplace-returns/financial-statement/preview")
async def preview_return_financial_statement(account_id: int = Form(...), file: UploadFile = File(...),
                                              u=Depends(current_user), db=Depends(get_db)):
    require(u, "finance")
    account = db.get(Account, account_id)
    if not account:
        raise HTTPException(404, "Магазин не найден")
    scope(u, account.company_id)
    if account.channel not in {"Kaspi", "Wildberries", "Ozon"}:
        raise HTTPException(422, "Выписку нужно связать с аккаунтом маркетплейса")
    if account.settlement_currency != "KZT":
        raise HTTPException(422, "Импорт суммы пока поддерживает только выписку в KZT")
    raw = await file.read()
    headers, rows = importing.parse_file(raw, file.filename or "statement.xlsx")
    return {"headers": headers, "sample": rows[:8], "row_count": len(rows),
            "suggested_mapping": _return_statement_mapping(headers),
            "note": "Предпросмотр не меняет данные. Финансовый возврат появится только после явного импорта с датой и суммой."}


@app.post("/api/marketplace-returns/financial-statement/import")
async def import_return_financial_statement(account_id: int = Form(...), mapping: str = Form(...),
                                            file: UploadFile = File(...), u=Depends(current_user),
                                            db=Depends(get_db)):
    require(u, "finance", True)
    account = db.get(Account, account_id)
    if not account:
        raise HTTPException(404, "Магазин не найден")
    scope(u, account.company_id)
    if account.channel not in {"Kaspi", "Wildberries", "Ozon"}:
        raise HTTPException(422, "Выписку нужно связать с аккаунтом маркетплейса")
    if account.settlement_currency != "KZT":
        raise HTTPException(422, "Импорт суммы пока поддерживает только выписку в KZT")
    try:
        column_map = json.loads(mapping)
    except (ValueError, TypeError):
        raise HTTPException(422, "Настройте сопоставление колонок")
    required_fields = {"external_return_id", "external_order_id", "financial_return_date", "amount"}
    if not isinstance(column_map, dict) or not required_fields.issubset(column_map) or any(not column_map.get(k) for k in required_fields):
        raise HTTPException(422, "Обязательны колонки: ID возврата, номер заказа, дата финансовой операции и сумма")
    raw = await file.read()
    headers, source_rows = importing.parse_file(raw, file.filename or "statement.xlsx")
    if any(value not in headers for value in column_map.values()):
        raise HTTPException(422, "Сопоставление содержит колонку, которой нет в файле")
    from decimal import Decimal, InvalidOperation
    imported = updated = skipped = 0
    errors = []
    for row_number, source_row in enumerate(source_rows, start=2):
        mapped = {field: source_row.get(column) for field, column in column_map.items()}
        operation_id = str(mapped.get("external_return_id") or "").strip()
        order_id = str(mapped.get("external_order_id") or "").strip()
        day = importing._date(mapped.get("financial_return_date"))
        try:
            operation_day = date.fromisoformat(str(day))
            amount_decimal = Decimal(str(mapped.get("amount") or "").replace("\u00a0", "").replace("\u202f", "").replace(" ", "").replace(",", "."))
            amount = int((abs(amount_decimal) * 100).quantize(Decimal("1")))
            quantity = int(str(mapped.get("quantity") or "1").strip())
        except (ValueError, TypeError, InvalidOperation, OverflowError):
            errors.append({"row": row_number, "message": "Не удалось прочитать дату, сумму или количество"})
            continue
        if not operation_id or not order_id or amount <= 0 or quantity <= 0:
            errors.append({"row": row_number, "message": "Нужны ID возврата, номер заказа, положительная сумма и количество"})
            continue
        row_currency = str(mapped.get("currency") or "KZT").strip().upper()
        if row_currency not in {"KZT", "₸", "ТЕНГЕ", "ТГ"}:
            errors.append({"row": row_number, "message": "Поддерживается только выписка в KZT; конвертация валют не выполняется"})
            continue
        sku = str(mapped.get("sku") or "").strip()
        product_name = str(mapped.get("product_name") or "").strip()
        source = "financial_statement_import"
        # Prefer linking to a single existing status observation by exact order
        # and SKU. Ambiguous or absent matches remain independent, auditable rows.
        existing = db.scalar(select(MarketplaceReturn).where(
            MarketplaceReturn.account_id == account.id,
            MarketplaceReturn.raw_data["financial_statement_return_id"].as_string() == operation_id,
        ))
        if existing is None:
            existing = db.scalar(select(MarketplaceReturn).where(
                MarketplaceReturn.account_id == account.id,
                MarketplaceReturn.source == source,
                MarketplaceReturn.external_return_id == operation_id,
            ))
        if existing is None:
            candidate_q = select(MarketplaceReturn).where(
                MarketplaceReturn.account_id == account.id,
                MarketplaceReturn.external_order_id == order_id,
                MarketplaceReturn.financial_status != "confirmed",
            )
            if sku: candidate_q = candidate_q.where(MarketplaceReturn.sku == sku)
            candidates = list(db.scalars(candidate_q.limit(2)))
            if len(candidates) == 1:
                existing = candidates[0]
        product_mapping = db.scalar(select(MarketplaceProductMapping).where(
            MarketplaceProductMapping.account_id == account.id,
            MarketplaceProductMapping.external_key.in_([sku, f"sku:{sku}"]),
        )) if sku else None
        product = db.get(Product, product_mapping.product_id) if product_mapping else None
        media = (product.media or []) if product else []
        if existing is None:
            existing = MarketplaceReturn(company_id=account.company_id, account_id=account.id,
                marketplace=account.channel, source=source, external_return_id=operation_id,
                external_order_id=order_id, detected_date=datetime.now(ZoneInfo("Asia/Almaty")).date(),
                accounted_status="unverified", return_status="financial_return_confirmed")
            db.add(existing)
            db.flush()
            db.add(MarketplaceReturnAudit(return_id=existing.id, user_id=u.id,
                action="financial_statement_imported", details={"source_file": (file.filename or "")[:160], "operation_id": operation_id}))
            imported += 1
        else:
            updated += 1
            db.add(MarketplaceReturnAudit(return_id=existing.id, user_id=u.id,
                action="financial_statement_matched", details={"source_file": (file.filename or "")[:160], "operation_id": operation_id}))
        existing.external_order_id = order_id
        existing.product_id = product.id if product else existing.product_id
        existing.product_name = product_name or (product.name if product else existing.product_name or "")
        existing.sku = sku or (product.sku if product else existing.sku or "")
        existing.image_url = (media[0] if media and isinstance(media[0], str) else existing.image_url or "")
        existing.quantity = quantity
        existing.financial_return_date = operation_day
        existing.amount = amount
        existing.currency = "KZT"
        existing.financial_status = "confirmed"
        existing.last_synced_at = datetime.now(timezone.utc)
        existing.raw_data = {**(existing.raw_data or {}), "financial_statement": True,
                             "financial_statement_return_id": operation_id, "file_name": (file.filename or "")[:160]}
    if errors and imported == 0 and updated == 0:
        raise HTTPException(422, {"message":"Ни одна строка не импортирована", "errors":errors[:20]})
    audit(db, u, "import_marketplace_returns", "marketplace_returns", details={
        "account_id":account.id, "source_file":(file.filename or "")[:160],
        "imported":imported, "updated":updated, "skipped":skipped, "errors":len(errors),
    })
    db.commit()
    return {"ok":True, "imported":imported, "updated":updated, "skipped":skipped,
            "errors":errors[:50], "message":"Данные записаны только в журнал возвратов; существующие финансовые отчёты не изменены."}


@app.post("/api/marketplace-returns/{marketplace_return_id}/inventory-match")
def match_marketplace_return_inventory(marketplace_return_id: int, u=Depends(current_user), db=Depends(get_db)):
    """Read-only match to an existing ERP return; never creates stock or financial entries."""
    require(u, "inventory")
    observation = db.get(MarketplaceReturn, marketplace_return_id)
    if not observation:
        raise HTTPException(404, "Возврат маркетплейса не найден")
    scope(u, observation.company_id)
    account = db.get(Account, observation.account_id)
    if not account or account.company_id != observation.company_id:
        raise HTTPException(409, "Магазин возврата не найден")
    matched = marketplace_returns.match_inventory_return(db, observation, account)
    if not matched:
        return {
            "matched": False,
            "message": "Не удалось однозначно сопоставить возврат с продажей и возвратом ERP. Остаток не изменён.",
        }
    return {
        "matched": True,
        "return_id": matched["id"],
        "status": matched["status"],
        "quantity": matched["quantity"],
        "message": "Найден возврат ERP. Изменение остатка выполняется только через подтверждение статуса.",
    }


@app.get("/api/marketplace-returns")
def get_marketplace_returns(
    marketplace: str = "all", company_id: int | None = None, account_id: int | None = None,
    start: date | None = None, end: date | None = None, search: str = "",
    return_status: str | None = None, financial_status: str | None = None,
    unknown_only: bool = False,
    offset: int = Query(0, ge=0), limit: int = Query(30, ge=1, le=100),
    u=Depends(current_user), db=Depends(get_db),
):
    """Separate return observations. This endpoint never reads or posts to P&L."""
    require(u, "sales")
    if marketplace not in {"all", "Kaspi", "Wildberries", "Ozon"}:
        raise HTTPException(422, "Неизвестный маркетплейс")
    if start and end and end < start:
        raise HTTPException(422, "Дата окончания раньше даты начала")
    if company_id is not None:
        scope(u, company_id)
    q = select(MarketplaceReturn, Account, Product).join(
        Account, MarketplaceReturn.account_id == Account.id
    ).outerjoin(Product, MarketplaceReturn.product_id == Product.id)
    if u.role != "Owner":
        q = q.where(MarketplaceReturn.company_id.in_(u.company_ids or []))
    if company_id is not None:
        q = q.where(MarketplaceReturn.company_id == company_id)
    if account_id is not None:
        account = db.get(Account, account_id)
        if not account:
            raise HTTPException(404, "Магазин не найден")
        scope(u, account.company_id)
        q = q.where(MarketplaceReturn.account_id == account_id)
    if marketplace != "all":
        q = q.where(MarketplaceReturn.marketplace == marketplace)
    if return_status:
        if return_status not in {"requested", "returned", "rejected", "financial_operation_reported", "financial_return_confirmed"}:
            raise HTTPException(422, "Неизвестный статус возврата")
        q = q.where(MarketplaceReturn.return_status == return_status)
    if financial_status:
        if financial_status not in {"confirmed", "unconfirmed"}:
            raise HTTPException(422, "Неизвестный финансовый статус возврата")
        q = q.where(MarketplaceReturn.financial_status == financial_status)
    if unknown_only:
        q = q.where(or_(MarketplaceReturn.financial_return_date.is_(None), MarketplaceReturn.amount.is_(None)))
    return_event_day = func.coalesce(MarketplaceReturn.financial_return_date, MarketplaceReturn.request_date)
    if start:
        q = q.where(or_(return_event_day >= start, return_event_day.is_(None)))
    if end:
        q = q.where(or_(return_event_day <= end, return_event_day.is_(None)))
    term = search.strip()[:120]
    if term:
        needle = f"%{term}%"
        q = q.where(or_(MarketplaceReturn.external_order_id.ilike(needle),
                        MarketplaceReturn.sku.ilike(needle), MarketplaceReturn.product_name.ilike(needle),
                        Account.name.ilike(needle)))
    rows = db.execute(q.order_by(MarketplaceReturn.detected_date.desc(), MarketplaceReturn.id.desc())
                      .offset(offset).limit(limit)).all()
    # Summary values are computed over the complete filtered set, independent of pagination.
    base = select(MarketplaceReturn).join(Account, MarketplaceReturn.account_id == Account.id)
    if u.role != "Owner": base = base.where(MarketplaceReturn.company_id.in_(u.company_ids or []))
    if company_id is not None: base = base.where(MarketplaceReturn.company_id == company_id)
    if account_id is not None: base = base.where(MarketplaceReturn.account_id == account_id)
    if marketplace != "all": base = base.where(MarketplaceReturn.marketplace == marketplace)
    if return_status: base = base.where(MarketplaceReturn.return_status == return_status)
    if financial_status: base = base.where(MarketplaceReturn.financial_status == financial_status)
    if unknown_only: base = base.where(or_(MarketplaceReturn.financial_return_date.is_(None), MarketplaceReturn.amount.is_(None)))
    base_event_day = func.coalesce(MarketplaceReturn.financial_return_date, MarketplaceReturn.request_date)
    if start: base = base.where(or_(base_event_day >= start, base_event_day.is_(None)))
    if end: base = base.where(or_(base_event_day <= end, base_event_day.is_(None)))
    if term:
        needle = f"%{term}%"
        base = base.where(or_(MarketplaceReturn.external_order_id.ilike(needle), MarketplaceReturn.sku.ilike(needle),
                              MarketplaceReturn.product_name.ilike(needle), Account.name.ilike(needle)))
    scoped_returns = base.subquery()
    agg = db.execute(select(
        func.count(),
        func.sum(case((scoped_returns.c.return_status == "requested", 1), else_=0)),
        func.sum(case((scoped_returns.c.financial_status == "confirmed", 1), else_=0)),
        func.sum(case(((scoped_returns.c.financial_status == "confirmed") & scoped_returns.c.amount.is_not(None), scoped_returns.c.amount), else_=0)),
        func.sum(case((scoped_returns.c.return_status == "rejected", 1), else_=0)),
        func.sum(case((scoped_returns.c.accounted_status == "yes", scoped_returns.c.amount), else_=0)),
        func.sum(case((scoped_returns.c.accounted_status == "no", scoped_returns.c.amount), else_=0)),
        func.sum(case((scoped_returns.c.accounted_status == "unverified", 1), else_=0)),
        func.sum(case(((scoped_returns.c.financial_return_date.is_(None)) | scoped_returns.c.amount.is_(None), 1), else_=0)),
    ).select_from(scoped_returns)).one()
    count = agg[0] or 0
    summary = {
        "requests": agg[1] or 0,
        "confirmed_count": agg[2] or 0,
        "confirmed_amount": agg[3] or 0,
        "pending_count": count - (agg[2] or 0),
        "rejected_count": agg[4] or 0,
        "accounted_amount": agg[5] or 0,
        "not_accounted_amount": agg[6] or 0,
        "unverified_accounting_count": agg[7] or 0,
        "unknown_count": agg[8] or 0,
        "existing_report_total": None, "adjusted_total": None, "test_mode": True,
    }
    integrations_q = select(Account).where(Account.channel.in_(["Kaspi", "Wildberries", "Ozon"]))
    if u.role != "Owner": integrations_q = integrations_q.where(Account.company_id.in_(u.company_ids or []))
    if company_id is not None: integrations_q = integrations_q.where(Account.company_id == company_id)
    if marketplace != "all": integrations_q = integrations_q.where(Account.channel == marketplace)
    integrations = [{"account_id": a.id, "marketplace": a.channel, "store": a.name,
                     "active": a.active, "mode": a.mode} for a in db.scalars(integrations_q)]
    return {"items": [marketplace_returns.serialize_return(r, a, p) for r, a, p in rows],
            "total": count, "offset": offset, "limit": limit, "summary": summary,
            "integrations": integrations}


@app.get("/api/marketplace-record-details")
def marketplace_record_details(
    channel: str, metric: str, start: date, end: date,
    company_id: int | None = None, account_id: int | None = None,
    u=Depends(current_user), db=Depends(get_db),
):
    """Read-only, account-scoped detail rows from already stored marketplace API records."""
    require(u, "analytics")
    if channel not in {"Kaspi", "Wildberries"}:
        raise HTTPException(422, "Подробный API-отчёт доступен для Kaspi и Wildberries")
    if metric not in {"sales", "returns", "logistics", "settlement"}:
        raise HTTPException(422, "Неизвестный показатель API-отчёта")
    if end < start:
        raise HTTPException(422, "Дата окончания раньше даты начала")
    aq = select(Account).where(Account.channel == channel, Account.active.is_(True))
    if u.role != "Owner": aq = aq.where(Account.company_id.in_(u.company_ids or []))
    if company_id is not None:
        scope(u, company_id)
        aq = aq.where(Account.company_id == company_id)
    if account_id is not None:
        aq = aq.where(Account.id == account_id)
    accounts = {a.id: a for a in db.scalars(aq)}
    if account_id is not None and account_id not in accounts:
        raise HTTPException(404, "Магазин выбранного маркетплейса не найден")
    if not accounts:
        return {"channel": channel, "metric": metric, "start": str(start), "end": str(end), "items": [],
                "note": "Нет активных API-магазинов в выбранной компании."}

    source_names = {"kaspi": ["kaspi"], "wildberries": ["wildberries", "wildberries_weekly"]}[channel.lower()]
    records = db.scalars(select(MarketplaceRecord).where(
        MarketplaceRecord.account_id.in_(accounts), MarketplaceRecord.source.in_(source_names)
    ).order_by(MarketplaceRecord.id.desc())).all()
    kaspi_order_codes = {}
    if channel == "Kaspi":
        kaspi_order_codes = {str((r.payload or {}).get("order_id")): str((r.payload or {}).get("code") or "")
                             for r in records if str((r.payload or {}).get("record_type") or "").lower() == "order"}
    output = []
    for rec in records:
        p = rec.payload or {}
        kind = str(p.get("record_type") or "").lower()
        event_date = None
        status = ""
        operation = ""
        reference = rec.external_id or ""
        product_name = str(p.get("product_name") or p.get("productName") or p.get("name") or p.get("title") or p.get("subjectName") or "")
        sku = str(p.get("sku") or p.get("supplierArticle") or p.get("vendorCode") or p.get("sa_name") or p.get("nmId") or "")
        quantity = p.get("quantity")
        amount = None
        currency = str(p.get("currency") or p.get("currency_name") or p.get("currencyName") or ("KZT" if channel == "Kaspi" else "RUB")).upper()
        include = False
        if channel == "Kaspi":
            status = str(p.get("order_status") or p.get("status") or "").upper()
            order_id = str(p.get("order_code") or p.get("code") or kaspi_order_codes.get(str(p.get("order_id") or "")) or p.get("order_id") or "")
            reference = order_id or reference
            if kind == "order_entry":
                returned = status in {"RETURNED", "KASPI_DELIVERY_RETURN_REQUESTED"}
                if metric == "returns" and returned:
                    include = True
                    event_date = p.get("returned_date") or p.get("return_date") or p.get("request_date")
                    operation = "Заявка на возврат" if status == "KASPI_DELIVERY_RETURN_REQUESTED" else "Возвращённый заказ"
                    amount = None
                elif metric == "sales" and status == "COMPLETED":
                    include = True
                    event_date = p.get("completed_date") or p.get("completion_date") or p.get("completionDate")
                    operation = "Продажа по дате выдачи"
                    amount = p.get("total_price") or p.get("totalPrice")
                elif metric == "logistics" and kind == "order_entry" and (p.get("delivery_cost_for_seller") is not None or p.get("deliveryCostForSeller") is not None or p.get("deliveryCost") is not None):
                    include = True
                    event_date = p.get("completed_date") or p.get("completionDate")
                    operation = "Доставка по данным заказа API"
                    amount = p.get("deliveryCostForSeller") or p.get("deliveryCost")
                    quantity = None
                elif metric == "settlement":
                    include = False
            elif kind in {"delivery_charge", "return_delivery_charge"} and metric == "logistics":
                include = True
                if kind == "return_delivery_charge":
                    event_date = p.get("returned_date") or p.get("return_date")
                else:
                    event_date = p.get("completed_date") or p.get("completion_date")
                operation = "Обратная доставка" if kind == "return_delivery_charge" else "Доставка"
                amount = p.get("amount")
                status = str(p.get("order_status") or status).upper()
                quantity = None
        else:
            operation = " ".join(str(p.get(key) or "") for key in (
                "docTypeName", "doc_type_name", "supplierOperName", "supplier_oper_name",
                "sellerOperName", "seller_oper_name")).strip()
            lower_operation = operation.lower()
            is_return = ("возврат" in lower_operation or "return" in lower_operation)
            is_sale = ("продаж" in lower_operation or lower_operation == "sale")
            event_date = p.get("rrdDate") or p.get("rr_dt") or p.get("rrDate") or p.get("saleDt") or p.get("sale_dt") or p.get("date")
            reference = str(p.get("rrdId") or p.get("rrd_id") or p.get("srid") or p.get("saleID") or reference)
            if metric == "returns" and kind in {"realization", ""} and is_return:
                include = True
            elif metric == "sales" and ((kind == "sale" and not str(p.get("saleID") or "").upper().startswith("R") and not p.get("isStorno")) or (kind in {"realization", ""} and is_sale)):
                include = True
                amount = p.get("priceWithDisc") if kind == "sale" else (p.get("retailAmount") or p.get("retail_amount"))
            elif metric == "logistics" and kind in {"realization", ""} and any(p.get(key) not in (None, "", 0, "0") for key in ("deliveryService", "deliveryRub", "delivery_rub", "rebillLogisticCost", "rebill_logistic_cost")):
                include = True
                amount = p.get("deliveryService") or p.get("deliveryRub") or p.get("delivery_rub") or p.get("rebillLogisticCost") or p.get("rebill_logistic_cost")
            elif metric == "settlement" and kind in {"realization", ""}:
                include = True
                amount = p.get("forPay") or p.get("ppvzForPay") or p.get("ppvz_for_pay")
        if not include:
            continue
        try:
            day = marketplace_pipeline._day(event_date, "Asia/Almaty" if channel == "Kaspi" else None) if event_date not in (None, "") else None
        except (ValueError, TypeError, KeyError, marketplace_pipeline.Pending):
            day = None
        if day is not None and not start <= day <= end:
            continue
        # Kaspi does not expose an actual refund amount/date in this order API.
        # Keep unknown event dates visible but never imply that order value is refund value.
        if channel == "Kaspi" and metric == "returns":
            amount = None
        output.append({"id": rec.id, "account_id": rec.account_id, "store": accounts[rec.account_id].name,
            "event_date": str(day) if day else None, "operation": operation or status or kind or "Операция API",
            "reference": reference, "product_name": product_name or "Товар не сопоставлен", "sku": sku,
            "quantity": quantity, "status": status or p.get("status") or "—", "amount": amount,
            "currency": currency, "source": rec.source})
    output.sort(key=lambda item: (item["event_date"] or "", item["id"]), reverse=True)
    note = ""
    if channel == "Kaspi" and metric == "returns":
        note = "Kaspi API подтверждает статус и состав заказа. Фактическую дату и сумму возврата API не отдаёт; они показаны как неподтверждённые до сверки с выпиской. Возвраты без даты остаются видимыми."
    elif channel == "Wildberries" and metric == "returns":
        note = "Возвраты взяты из финансовых строк WB по типу операции. Суммы показаны в исходной валюте отчёта; банковское зачисление здесь не подтверждается."
    elif metric == "settlement":
        note = "Начисление по строкам API-отчёта, а не подтверждение банковской выплаты."
    return {"channel": channel, "metric": metric, "start": str(start), "end": str(end), "items": output, "note": note}


@app.get("/api/marketplace-report.xlsx")
def marketplace_report_xlsx(channel: str, start: date, end: date,
                            company_id: int | None = None, account_id: int | None = None,
                            u=Depends(current_user), db=Depends(get_db)):
    """Export a separate API-sourced marketplace report without changing ERP calculations."""
    require(u, "analytics")
    if channel not in {"Kaspi", "Wildberries"}:
        raise HTTPException(422, "Отчёт API доступен для Kaspi и Wildberries")
    if end < start:
        raise HTTPException(422, "Дата окончания раньше даты начала")
    if channel == "Wildberries":
        # The official WB Finance export is the source of truth for this channel.
        # Keep the custom range in its isolated report namespace.
        period = "custom"
        _wb_finance_period_args(period, start, end)
        account = _wb_week_account(db, u, account_id)
        from .wb_weekly import build_report, create_xlsx
        report = build_report(db, account, start, end, period=period)
        if report.get("status") != "готов" or not report.get("row_count"):
            raise HTTPException(409, "Финансовый отчёт WB за период ещё не получен")
        raw = create_xlsx(report)
        filename = f"WB-API-{start}-{end}.xlsx"
    else:
        # Reuse the existing, read-only Kaspi API summary and detail selectors.
        # These are copies of current API/ERP fields; this export has no new
        # commissions, tax, return or profit formulas.
        summary = marketplace_summary(channel, start, end, company_id, account_id, u, db)
        _require_complete_kaspi_report(summary)
        sales = marketplace_record_details(channel, "sales", start, end, company_id, account_id, u, db)
        returns = marketplace_record_details(channel, "returns", start, end, company_id, account_id, u, db)
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from io import BytesIO
        book = Workbook(); overview = book.active; overview.title = "Итоги API"
        api_available = summary.get("api_data_available") is True
        overview.append(["Kaspi · отчёт по данным API"])
        overview.append(["Период", f"{start} — {end}"])
        overview.append(["Статус синхронизации", summary.get("sync_status")])
        overview.append(["Последняя синхронизация", str(summary.get("last_sync") or "—")])
        overview.append(["Данные API подтверждены", "Да" if api_available else "Нет"])
        overview.append(["Показатель", "Значение", "Источник / смысл"])
        for label, key, note in [
            ("Продажи API, до возвратов", "api_gross_revenue", "Подтверждённые строки Kaspi API"),
            ("Продажи API, за вычетом возвратов", "api_revenue", "Поле сводки Kaspi API"),
            ("Продано, шт.", "api_units", "Количество из строк Kaspi API"),
            ("Возвраты API", "api_returns", "Фактическая сумма возврата не предоставляется API"),
            ("Возвраты, шт.", "api_return_units", "Количество из строк Kaspi API"),
            ("Строки без конвертации", "api_unconverted", "Требуют проверки даты/суммы"),
        ]:
            value = summary.get(key) if api_available else None
            if key in {"api_gross_revenue", "api_revenue", "api_returns"} and value is not None:
                value = value / 100
            overview.append([label, value, note])
        sale_ws = book.create_sheet("Продажи API")
        headers = ["Дата выдачи", "Магазин", "ID заказа", "Товар", "SKU", "Количество", "Сумма API", "Валюта", "Статус"]
        sale_ws.append(headers)
        for item in sales.get("items", []):
            sale_ws.append([item.get("event_date"), item.get("store"), item.get("reference"), item.get("product_name"), item.get("sku"), item.get("quantity"), item.get("amount"), item.get("currency"), item.get("status")])
        return_ws = book.create_sheet("Возвраты API")
        return_ws.append(["Дата возврата API", "Магазин", "ID заказа", "Товар", "SKU", "Количество", "Сумма возврата", "Валюта", "Статус"])
        for item in returns.get("items", []):
            return_ws.append([item.get("event_date"), item.get("store"), item.get("reference"), item.get("product_name"), item.get("sku"), item.get("quantity"), item.get("amount"), item.get("currency"), item.get("status")])
        note_ws = book.create_sheet("Примечания")
        note_ws.append(["Показатель", "Источник", "Ограничение"])
        note_ws.append(["Суммы и продажи", "Kaspi API и действующая сводка QOR", "Экспорт не пересчитывает финансовые показатели"])
        note_ws.append(["Возвраты", "Статусы и позиции Kaspi API", "Kaspi API не подтверждает фактическую дату и сумму возврата; пустые суммы намеренно оставлены пустыми"])
        for ws in book.worksheets:
            ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions if ws.max_row > 1 else None
            header_rows = (1, 6) if ws.title == "Итоги API" else (1,)
            for row_no in header_rows:
                for cell in ws[row_no]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="135B45")
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
            for column in ws.columns:
                width = min(44, max(14, max(len(str(c.value or "")) for c in list(column)[:200]) + 2))
                ws.column_dimensions[column[0].column_letter].width = width
        out = BytesIO(); book.save(out); raw = out.getvalue()
        filename = f"Kaspi-API-{start}-{end}.xlsx"
    return Response(raw, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{filename}"'})


@app.post("/api/marketplace-returns/discover")
def discover_marketplace_returns(payload: dict, tasks: BackgroundTasks,
                                 u=Depends(current_user), db=Depends(get_db)):
    """Start an explicit background scan of cached return source rows."""
    require(u, "accounts", True)
    try:
        account_id = int(payload.get("account_id"))
    except (TypeError, ValueError):
        raise HTTPException(422, "Укажите магазин")
    account = db.get(Account, account_id)
    if not account:
        raise HTTPException(404, "Магазин не найден")
    scope(u, account.company_id)
    if account.channel not in {"Kaspi", "Wildberries"}:
        raise HTTPException(409, "Автоматическое обнаружение для этого маркетплейса пока недоступно")
    if not account.active or account.mode != "api":
        raise HTTPException(409, "Для проверки возвратов нужна активная API-интеграция")
    state = db.scalar(select(MarketplaceReturnSyncState).where(MarketplaceReturnSyncState.account_id == account_id))
    started_at = state.started_at if state else None
    if started_at is not None and started_at.tzinfo is None:
        started_at = started_at.replace(tzinfo=timezone.utc)
    if state and state.status == "running" and started_at and started_at > now() - timedelta(minutes=20):
        return {"status": "running", "account_id": account_id, "already_running": True}
    tasks.add_task(marketplace_returns.run_discovery, account_id)
    return {"status": "queued", "account_id": account_id}


@app.get("/api/marketplace-returns/sync-status")
def marketplace_returns_sync_status(company_id: int | None = None,
                                    u=Depends(current_user), db=Depends(get_db)):
    require(u, "sales")
    q = select(MarketplaceReturnSyncState, Account).join(Account, Account.id == MarketplaceReturnSyncState.account_id)
    if u.role != "Owner": q = q.where(Account.company_id.in_(u.company_ids or []))
    if company_id is not None: scope(u, company_id); q = q.where(Account.company_id == company_id)
    return {"items": [{"account_id": a.id, "marketplace": a.channel, "store": a.name,
                       "status": state.status, "scanned": state.scanned, "found": state.found,
                       "error": state.error, "started_at": state.started_at, "completed_at": state.completed_at}
                      for state, a in db.execute(q)]}


@app.get("/api/marketplace-summary")
def marketplace_summary(channel: str, start: date, end: date, company_id: int | None = None,
                        account_id: int | None = None, u=Depends(current_user), db=Depends(get_db)):
    """Fast marketplace page summary; bounded to the selected channel and period."""
    require(u, "analytics")
    if end < start:
        raise HTTPException(422, "Дата окончания раньше даты начала")
    if channel not in {"Kaspi", "Wildberries", "Ozon"}: raise HTTPException(422, "Неподдерживаемый маркетплейс")
    aq = select(Account).where(Account.channel == channel, Account.active.is_(True))
    if u.role != "Owner": aq = aq.where(Account.company_id.in_(u.company_ids or []))
    if company_id is not None: scope(u, company_id); aq = aq.where(Account.company_id == company_id)
    accounts = list(db.scalars(aq))
    if account_id is not None:
        selected = next((account for account in accounts if account.id == account_id), None)
        if selected is None:
            raise HTTPException(404, "Магазин выбранного маркетплейса не найден")
        accounts = [selected]
    ids = [a.id for a in accounts]; by_id = {a.id:a for a in accounts}
    display_sync = [sync_status_for_display(a) for a in accounts]
    sync_in_progress = any(status == 'running' for status, _error in display_sync)
    sync_status = ('failed' if any(x[0] == 'failed' for x in display_sync) else
                   'running' if any(x[0] == 'running' for x in display_sync) else
                   'partial' if any(x[0] == 'partial' for x in display_sync) else
                   'done' if any(x[0] == 'done' for x in display_sync) else 'never')
    sync_errors = [error for status, error in display_sync if error and status != 'running']
    last_sync = max((a.last_sync for a in accounts if a.last_sync), default=None)
    if not ids:
        return {'channel':channel,'start':str(start),'end':str(end),'revenue':0,'units':0,'api_gross_revenue':None,'api_revenue':None,'api_units':None,'api_returns':None,'api_return_units':None,'api_unconverted':0,'api_data_available':False,'sync_status':'not_connected','sync_errors':[],'last_sync':None,'pending_sales':0,'pending_reasons':{},'financial_complete':True,'cogs':0,'costs':0,'commission':0,'logistics':0,'return_logistics':0,'kaspi_estimated_proceeds':None,'advertising':0,'returns':0,'return_units':0,'settlement_accrued':0,'settlement_rows':0,'settlement_unconverted':0,'net_profit':0,'margin':None,'sync_in_progress':False,'source_complete':True}

    sales = list(db.scalars(select(Sale).where(Sale.account_id.in_(ids), Sale.date >= start, Sale.date <= end)))
    period_returns = list(db.scalars(select(Return).join(Sale, Return.sale_id == Sale.id).where(Sale.account_id.in_(ids), Return.date >= start, Return.date <= end)))
    sale_cogs, return_cogs, pending_ids = marketplace_costs.resolve_cost_overrides(db, sales, period_returns, by_id)
    return_revenue = sum(x.revenue for x in period_returns); return_units = sum(x.quantity for x in period_returns)
    revenue = sum(x.revenue for x in sales) - return_revenue; units = sum(x.quantity for x in sales) - return_units
    restocked_returns = [x for x in period_returns if x.inventory_status == "restocked"]
    cogs = sum(sale_cogs.get(x.id, x.cogs) for x in sales) - sum(return_cogs.get(x.id, x.cogs) for x in restocked_returns); commission = sum(x.commission for x in sales)
    logistics = sum(x.logistics for x in sales); advertising = sum(x.advertising for x in sales)
    pending = sum(1 for x in sales if x.id in pending_ids)
    # Marketplace expenses may be posted independently from a Sale.
    expenses = list(db.scalars(select(Expense).where(Expense.account_id.in_(ids), Expense.date >= start, Expense.date <= end)))
    commission += sum(x.amount for x in expenses if x.category in {"Marketplace commission", "Marketplace commission VAT"})
    commission_estimated = 0
    commission_rate = None
    commission_estimate_suppressed = False
    logistics += sum(x.amount for x in expenses if x.category == "Marketplace logistics")
    return_logistics = sum(x.amount for x in expenses if x.category == "Marketplace return logistics")
    advertising += sum(x.amount for x in expenses if x.category == "Реклама")
    other = sum(x.amount for x in expenses if x.category not in {"Marketplace commission","Marketplace commission VAT","Marketplace logistics","Marketplace return logistics","Реклама"})

    api_gross_revenue=api_revenue=api_units=api_unconverted=api_returns=api_return_units=settlement_accrued=settlement_rows=settlement_unconverted=0
    raw_kaspi_commission_bases = {}
    api_data_available = False
    raw_q = select(MarketplaceRecord.account_id, MarketplaceRecord.payload).where(MarketplaceRecord.account_id.in_(ids), MarketplaceRecord.source == channel.lower())
    raw_rows = list(db.execute(raw_q))
    for account_id, payload in raw_rows:
        p=payload or {}; typ=str(p.get('record_type') or '').lower(); src=channel.lower()
        if src == 'kaspi' and typ == 'order_entry':
            status_now = str(p.get('order_status') or '').upper()
            date_field = (p.get('returned_date') or p.get('return_date')) if status_now == 'RETURNED' else p.get('completed_date')
            try:
                api_data_available = api_data_available or start <= marketplace_pipeline._day(date_field, 'Asia/Almaty') <= end
            except marketplace_pipeline.Pending:
                pass
        try:
            if src=='wildberries' and typ=='sale' and not str(p.get('saleID') or '').upper().startswith('R') and not p.get('isStorno'):
                dt=p.get('date'); day=marketplace_pipeline._day(dt)
                if start<=day<=end:
                    api_units += 1; amount=p.get('priceWithDisc') if p.get('priceWithDisc') is not None else p.get('finishedPrice')
                    try: api_revenue += marketplace_pipeline._cents(amount,by_id[account_id],p.get('currency') or 'RUB',db,dt)
                    except marketplace_pipeline.Pending: api_unconverted += 1
            elif src=='wildberries' and typ=='realization':
                dt=p.get('rrDate') or p.get('rr_dt') or p.get('saleDt') or p.get('sale_dt'); day=marketplace_pipeline._day(dt)
                if start<=day<=end:
                    value=next((p.get(k) for k in ('forPay','ppvzForPay','ppvz_for_pay') if p.get(k) is not None), None)
                    if value is not None:
                        try: settlement_accrued += marketplace_pipeline._cents(wb_weekly.signed_wb_sale_amount(p, value),by_id[account_id],p.get('currency') or p.get('currency_name') or 'RUB',db,dt); settlement_rows += 1
                        except marketplace_pipeline.Pending: settlement_unconverted += 1
            elif src=='ozon' and typ=='posting' and str(p.get('status') or '').lower() in {'delivered','completed'}:
                dt=p.get('delivered_at') or p.get('shipment_date') or p.get('created_at'); day=marketplace_pipeline._day(dt)
                if start<=day<=end:
                    for x in p.get('products') or []:
                        qty=int(x.get('quantity') or 0); api_units+=qty
                        try: api_revenue += marketplace_pipeline._cents(marketplace_pipeline._number(x.get('price'))*qty,by_id[account_id],x.get('currency_code') or p.get('currency_code') or 'RUB',db,dt)
                        except marketplace_pipeline.Pending: api_unconverted += 1
            elif src=='kaspi' and typ=='order_entry':
                status=str(p.get('order_status') or '').upper()
                returned=status == 'RETURNED'
                if status=='COMPLETED' or returned:
                    sale_dt=p.get('completed_date')
                    sale_day=marketplace_pipeline._day(sale_dt, 'Asia/Almaty')
                    qty=int(p.get('quantity') or 0)
                    try:
                        sale_amount=marketplace_pipeline._cents(p.get('total_price'),by_id[account_id],'KZT',db,sale_dt)
                        if start<=sale_day<=end:
                            api_gross_revenue += sale_amount
                            api_revenue += sale_amount
                            api_units += qty
                            if status=='COMPLETED':
                                key=(account_id,sale_day)
                                raw_kaspi_commission_bases[key]=raw_kaspi_commission_bases.get(key,0)+sale_amount
                        if returned:
                            return_dt=p.get('returned_date') or p.get('return_date')
                            return_day=marketplace_pipeline._day(return_dt, 'Asia/Almaty')
                            if start<=return_day<=end:
                                api_returns += sale_amount
                                api_return_units += qty
                                api_revenue -= sale_amount
                                api_units -= qty
                    except marketplace_pipeline.Pending:
                        api_unconverted += 1
        except (marketplace_pipeline.Pending, ValueError, TypeError, KeyError):
            continue
    if channel == "Kaspi":
        commission_estimated, commission_rate, commission_estimate_suppressed = _kaspi_commission_estimate(
            sales, expenses, raw_kaspi_commission_bases
        )
        commission += commission_estimated
    pending_return_inspection_units = sum(x.quantity for x in period_returns if x.inventory_status in {"requested", "approved", "in_transit", "received", "inspection"})
    profit = None if pending or pending_return_inspection_units else revenue-cogs-commission-logistics-return_logistics-advertising-other
    kaspi_estimated_proceeds = api_revenue - commission - logistics - return_logistics if channel == "Kaspi" else None
    quality = kaspi_integrity.source_quality(accounts, raw_rows, start, end) if channel == "Kaspi" else None
    source_complete = not sync_in_progress and not sync_errors and not api_unconverted
    if quality:
        source_complete = source_complete and quality["coverage_complete"] and quality["data_complete"]
    api_unknown = channel == 'Kaspi'  and not api_data_available and sync_status not in {'done','partial'}
    return {'channel':channel,'start':str(start),'end':str(end),'revenue':revenue,'units':units,'api_gross_revenue':None if api_unknown else api_gross_revenue,'api_revenue':None if api_unknown else api_revenue,'api_units':None if api_unknown else api_units,'api_returns':None if api_unknown else api_returns,'api_return_units':None if api_unknown else api_return_units,'api_unconverted':api_unconverted,'api_data_available':api_data_available or sync_status in {'done','partial'},'sync_status':sync_status,'sync_errors':sync_errors,'last_sync':last_sync.isoformat() if last_sync else None,
        'pending_sales':pending,'pending_return_inspection_units':pending_return_inspection_units,'pending_reasons':{'COGS_PENDING':pending} if pending else {},'financial_complete':not pending and not api_unconverted and not pending_return_inspection_units and (source_complete if channel == 'Kaspi' else True),
        'cogs':cogs,'costs':commission+logistics+return_logistics+advertising+other,'commission':commission,'commission_estimated':commission_estimated,'commission_rate':commission_rate,'commission_estimate_suppressed':commission_estimate_suppressed,'logistics':logistics,'return_logistics':return_logistics,'kaspi_estimated_proceeds':kaspi_estimated_proceeds,'advertising':advertising,'returns':return_revenue,'return_units':return_units,
        'settlement_accrued':settlement_accrued if settlement_rows else None,'settlement_rows':settlement_rows,'settlement_unconverted':settlement_unconverted,'net_profit':profit,
        'margin':round(profit/revenue*100,2) if revenue and profit is not None else None,'sync_in_progress':sync_in_progress,'source_complete':source_complete,'source_quality':quality}



@app.get("/api/expenses/operations")
def expense_operations(start: date, end: date, company_id: int | None = None,
                       channel: str | None = None,
                       limit: int = Query(default=200, ge=1, le=500),
                       offset: int = Query(default=0, ge=0),
                       u=Depends(current_user), db=Depends(get_db)):
    """Return individual posted expense rows and sale-level marketplace charges."""
    require(u, "finance")
    if end < start:
        raise HTTPException(422, "Дата окончания раньше даты начала")
    allowed_channels = {"Kaspi", "Wildberries", "Ozon", "Offline", "B2B"}
    if channel and channel not in allowed_channels:
        raise HTTPException(422, "Неизвестный канал расходов")
    if company_id is not None:
        scope(u, company_id)

    account_q = scope_query(select(Account), Account, u)
    if company_id is not None:
        account_q = account_q.where(Account.company_id == company_id)
    accounts = {account.id: account for account in db.scalars(account_q)}
    account_ids = set(accounts)

    expenses_q = scope_query(select(Expense), Expense, u).where(
        Expense.date >= start, Expense.date <= end)
    if company_id is not None:
        expenses_q = expenses_q.where(Expense.company_id == company_id)
    expenses = list(db.scalars(expenses_q))

    sales_q = scope_query(select(Sale), Sale, u).where(
        Sale.date >= start, Sale.date <= end, Sale.account_id.in_(account_ids or {-1}))
    if company_id is not None:
        sales_q = sales_q.where(Sale.company_id == company_id)
    sales = list(db.scalars(sales_q))
    product_ids = {sale.product_id for sale in sales}
    products = {product.id: product for product in db.scalars(
        select(Product).where(Product.id.in_(product_ids or {-1})))}

    category_labels = {
        "Marketplace commission": "Комиссия маркетплейса",
        "Marketplace commission VAT": "НДС с комиссии",
        "Marketplace acquiring": "Эквайринг",
        "Marketplace logistics": "Логистика",
        "Marketplace return logistics": "Обратная логистика",
        "Marketplace storage": "Хранение",
        "Marketplace acceptance": "Приёмка",
        "Marketplace penalties": "Штрафы",
        "Marketplace deductions": "Удержания",
        "Marketplace other expenses": "Прочие расходы маркетплейса",
        "Реклама": "Реклама",
    }
    rows = []
    for expense in expenses:
        account = accounts.get(expense.account_id)
        row_channel = account.channel if account else "Offline"
        if channel and row_channel != channel:
            continue
        rows.append({
            "key": f"expense:{expense.id}", "date": str(expense.date),
            "channel": row_channel, "account_name": account.name if account else "—",
            "category": category_labels.get(expense.category, expense.category or "Прочие расходы"),
            "amount": expense.amount, "record_type": "expense",
            "record_id": expense.id, "reference": expense.reference or "—",
            "product_id": None, "product_name": "", "sku": "",
            "quantity": None, "status": "Проведено",
            "description": expense.description or expense.project or "",
        })
    sale_categories = (
        ("commission", "Комиссия маркетплейса"),
        ("logistics", "Логистика"),
        ("advertising", "Реклама"),
    )
    for sale in sales:
        account = accounts.get(sale.account_id)
        row_channel = account.channel if account else "Offline"
        if channel and row_channel != channel:
            continue
        product = products.get(sale.product_id)
        for field, category in sale_categories:
            amount = int(getattr(sale, field, 0) or 0)
            if not amount:
                continue
            rows.append({
                "key": f"sale:{sale.id}:{field}", "date": str(sale.date),
                "channel": row_channel, "account_name": account.name if account else "—",
                "category": category, "amount": amount, "record_type": "sale_charge",
                "record_id": sale.id, "reference": sale.reference or "—",
                "product_id": sale.product_id, "product_name": product.name if product else "",
                "sku": product.sku if product else "", "quantity": sale.quantity,
                "status": "Проведено", "description": "",
            })
    rows.sort(key=lambda row: (row["date"], row["record_id"], row["key"]), reverse=True)
    total = len(rows)
    page = rows[offset:offset + limit]
    return {"start": str(start), "end": str(end), "items": page,
            "total": total, "limit": limit, "offset": offset}

@app.get("/api/expenses/summary")
def expense_summary(start: date, end: date, company_id: int | None = None,
                    u=Depends(current_user), db=Depends(get_db)):
    """Summarize posted expenses by marketplace/account and readable category."""
    require(u, "finance")
    if end < start:
        raise HTTPException(422, "Дата окончания раньше даты начала")
    expense_q = scope_query(select(Expense), Expense, u).where(
        Expense.date >= start, Expense.date <= end)
    if company_id is not None:
        scope(u, company_id)
        expense_q = expense_q.where(Expense.company_id == company_id)
    expense_rows = list(db.scalars(expense_q))

    # Inactive/archived integrations keep their historical expense attribution.
    account_q = scope_query(select(Account), Account, u)
    if company_id is not None:
        account_q = account_q.where(Account.company_id == company_id)
    accounts = {account.id: account for account in db.scalars(account_q)}

    # Commission, logistics and advertising can be posted on a sale instead of
    # as a separate Expense. Include them here using the same arithmetic as the
    # profit report, so the expense page and accounting do not disagree.
    sale_q = scope_query(select(Sale), Sale, u).where(Sale.date >= start, Sale.date <= end,
        Sale.account_id.in_(accounts.keys() or [-1]))
    if company_id is not None:
        sale_q = sale_q.where(Sale.company_id == company_id)
    sales = list(db.scalars(sale_q))

    category_labels = {
        "Marketplace commission": "Комиссия маркетплейса",
        "Marketplace commission VAT": "НДС с комиссии",
        "Marketplace acquiring": "Эквайринг",
        "Marketplace logistics": "Логистика",
        "Marketplace return logistics": "Обратная логистика",
        "Marketplace storage": "Хранение",
        "Marketplace acceptance": "Приёмка",
        "Marketplace penalties": "Штрафы",
        "Marketplace deductions": "Удержания",
        "Marketplace other expenses": "Прочие расходы маркетплейса",
        "Реклама": "Реклама",
        "Налоги": "Налоги",
        "Аренда": "Аренда",
        "Зарплата": "Зарплата",
    }
    by_channel = {}
    by_category = {}
    def add(account_id, category, amount, source="posted"):
        amount = int(amount or 0)
        if not amount:
            return
        account = accounts.get(account_id)
        channel = account.channel if account else "Offline"
        account_name = account.name if account else "Офлайн / без магазина"
        channel_label = {"Kaspi": "Kaspi", "Wildberries": "Wildberries", "Ozon": "Ozon",
                         "Offline": "Офлайн", "B2B": "B2B"}.get(channel, channel)
        channel_key = (channel, account_id)
        channel_row = by_channel.setdefault(channel_key, {"channel": channel, "channel_label": channel_label,
            "account_id": account_id, "account_name": account_name, "total": 0,
            "posted_total": 0, "estimated_total": 0, "categories": {}, "category_breakdown": {}})
        channel_row["total"] += amount
        channel_row["estimated_total" if source == "estimate" else "posted_total"] += amount
        label = category_labels.get(category, category or "Прочие расходы")
        channel_row["categories"][label] = channel_row["categories"].get(label, 0) + amount
        category_detail = channel_row["category_breakdown"].setdefault(label, {
            "posted_total": 0, "estimated_total": 0, "total": 0})
        category_detail["estimated_total" if source == "estimate" else "posted_total"] += amount
        category_detail["total"] += amount
        category_row = by_category.setdefault(label, {"category": label, "total": 0,
            "posted_total": 0, "estimated_total": 0})
        category_row["total"] += amount
        category_row["estimated_total" if source == "estimate" else "posted_total"] += amount

    for row in expense_rows:
        add(row.account_id, row.category, row.amount)
    for sale in sales:
        for category, value in (("Marketplace commission", sale.commission),
                                ("Marketplace logistics", sale.logistics),
                                ("Реклама", sale.advertising)):
            add(sale.account_id, category, value)

    commission_estimated = 0
    commission_rate = None
    commission_warning = False
    if accounts:
        commission_estimates, commission_rate, commission_warning = f.kaspi_commission_estimates(
            db, u, start, end, company_id, None, sales, expense_rows, accounts)
        for (account_id, _day), amount in commission_estimates.items():
            commission_estimated += amount
            add(account_id, "Marketplace commission", amount, "estimate")
    categories = sorted(by_category.values(), key=lambda row: (-row["total"], row["category"]))
    channels = sorted(by_channel.values(), key=lambda row: (-row["total"], row["channel_label"], row["account_name"]))
    return {"start": str(start), "end": str(end), "total": sum(row["total"] for row in categories),
        "posted_total": sum(row.amount for row in expense_rows) + sum(s.commission+s.logistics+s.advertising for s in sales),
        "commission_estimated": commission_estimated, "commission_rate": commission_rate,
        "commission_estimate_warning": commission_warning, "channels": channels, "categories": categories,
        "records": len(expense_rows) + len(sales)}


@app.get("/api/kaspi/orders-created")
def kaspi_orders_created(start: date, end: date, company_id: int | None = None,
                         account_id: int | None = None,
                         limit: int = Query(default=100, ge=1, le=500),
                         offset: int = Query(default=0, ge=0),
                         u=Depends(current_user), db=Depends(get_db)):
    """List locally cached Kaspi orders whose creationDate is inside the selected range."""
    require(u, "analytics")
    if end < start:
        raise HTTPException(422, "Дата окончания раньше даты начала")
    aq = select(Account).where(Account.channel == "Kaspi", Account.active.is_(True))
    if u.role != "Owner": aq = aq.where(Account.company_id.in_(u.company_ids or []))
    if company_id is not None: scope(u, company_id); aq = aq.where(Account.company_id == company_id)
    accounts = list(db.scalars(aq))
    if account_id is not None:
        selected = next((account for account in accounts if account.id == account_id), None)
        if selected is None:
            raise HTTPException(404, "Kaspi-магазин не найден")
        accounts = [selected]
    ids = [a.id for a in accounts]
    if not ids:
        return {"start":str(start),"end":str(end),"items":[],"total":0,"orders_amount":None,"api_data_available":False,
                "sync_status":"not_connected","sync_errors":[],"last_sync":None}
    rows = db.scalars(select(MarketplaceRecord).where(
        MarketplaceRecord.account_id.in_(ids), MarketplaceRecord.source == "kaspi")).all()
    orders = {}
    for row in rows:
        p = row.payload or {}
        if p.get("record_type") == "order":
            try: created = marketplace_pipeline._day(p.get("creation_date"), "Asia/Almaty")
            except marketplace_pipeline.Pending: continue
            if start <= created <= end:
                orders[(row.account_id, str(p.get("order_id") or row.external_id.removeprefix("order:")))] = {
                    "order_id":str(p.get("order_id") or row.external_id.removeprefix("order:")),
                    "account_id":row.account_id,"created_date":str(created),
                    "status":str(p.get("status") or p.get("order_status") or "unknown"),
                    "amount":p.get("total_price"),"units":0}
    for row in rows:
        p = row.payload or {}
        if p.get("record_type") != "order_entry": continue
        key = (row.account_id, str(p.get("order_id") or ""))
        if key in orders:
            try: orders[key]["units"] += int(p.get("quantity") or 0)
            except (ValueError, TypeError): pass
    items = sorted(orders.values(), key=lambda x:(x["created_date"],x["order_id"]), reverse=True)
    display_sync = [sync_status_for_display(a) for a in accounts]
    statuses = [x[0] for x in display_sync]
    sync_state = "failed" if "failed" in statuses else "running" if "running" in statuses else "done" if "done" in statuses else "partial" if "partial" in statuses else "never"
    verified = bool(items) or sync_state in {"done", "partial"}
    return {"start":str(start),"end":str(end),"items":items[offset:offset+limit],"total":len(items),
            "orders_amount":sum((Decimal(str(x["amount"])) for x in items if x["amount"] not in (None,"")),Decimal("0")) if verified else None,
            "api_data_available":verified,"sync_status":sync_state,
            "sync_errors":[error for status,error in display_sync if error and status != "running"],
            "last_sync":max((a.last_sync for a in accounts if a.last_sync),default=None).isoformat() if any(a.last_sync for a in accounts) else None}


@app.post("/api/wildberries/weekly-excel-preview")
async def wb_weekly_excel_preview(file: UploadFile = File(...), u=Depends(current_user)):
    """Preview official weekly WB report without modifying any financial records."""
    require(u, "analytics")
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(422, "Нужен оригинальный отчёт WB в формате .xlsx")
    from .wb_excel_reconciliation import parse_report
    try:
        return parse_report(await file.read(5 * 1024 * 1024 + 1))
    except (ValueError, StopIteration, OSError, KeyError) as exc:
        raise HTTPException(422, str(exc)) from exc


@app.get("/api/wildberries/reconciliation")
def wildberries_reconciliation(start: date, end: date, company_id: int | None = None,
                               limit: int = Query(default=100, ge=1, le=500),
                               offset: int = Query(default=0, ge=0),
                               u=Depends(current_user), db=Depends(get_db)):
    """Explain WB realization rows without pretending that settlement equals bank payout.

    The ledger retains native WB amounts/currency; currency conversion and
    reconciliation to a WB report must be done for the same report period.
    """
    require(u, "analytics")
    if end < start:
        raise HTTPException(422, "Дата окончания раньше даты начала")
    aq = select(Account.id).where(Account.channel == "Wildberries")
    if u.role != "Owner":
        aq = aq.where(Account.company_id.in_(u.company_ids or []))
    if company_id is not None:
        scope(u, company_id)
        aq = aq.where(Account.company_id == company_id)
    ids = list(db.scalars(aq))
    if not ids:
        return {"start": str(start), "end": str(end), "rows": [], "total": 0,
                "by_currency": {}, "unclassified": 0,
                "note": "Нет подключённых аккаунтов Wildberries"}
    records = db.execute(select(MarketplaceRecord.external_id, MarketplaceRecord.payload)
                         .where(MarketplaceRecord.account_id.in_(ids),
                                MarketplaceRecord.source == "wildberries")).all()
    rows = []
    for external_id, payload in records:
        p = payload or {}
        if str(p.get("record_type") or "").lower() != "realization":
            continue
        dt = p.get("rrDate") or p.get("rr_dt") or p.get("saleDt") or p.get("sale_dt")
        try:
            day = marketplace_pipeline._day(dt)
        except (ValueError, TypeError, KeyError):
            continue
        if not start <= day <= end:
            continue
        operation = str(p.get("sellerOperName") or p.get("supplier_oper_name") or p.get("doc_type_name") or "").strip()
        low = operation.lower()
        if "возврат" in low or "return" in low: category = "return"
        elif "штраф" in low or "penalty" in low: category = "penalty"
        elif "логист" in low or "достав" in low or "delivery" in low: category = "logistics"
        elif "удержан" in low or "deduction" in low: category = "deduction"
        elif "хранен" in low or "storage" in low: category = "storage"
        elif "продаж" in low or low == "sale": category = "sale"
        else: category = "unclassified"
        currency = str(p.get("currency") or p.get("currency_name") or "RUB").upper()
        currency = {"РУБ": "RUB", "РУБ.": "RUB", "ТЕНГЕ": "KZT"}.get(currency, currency)
        # Monetary values are NOT combined into one amount: fields may overlap.
        fields = {"retail_amount": ("retailAmount", "retail_amount"),
                  "seller_for_pay": ("ppvzForPay", "ppvz_for_pay"),
                  "sales_commission": ("ppvzSalesCommission", "ppvz_sales_commission"),
                  "delivery": ("deliveryService", "delivery_rub"),
                  "penalty": ("penalty",), "deduction": ("deduction",),
                  "acquiring": ("acquiringFee", "acquiring_fee")}
        amounts = {}
        for label, keys in fields.items():
            value = next((p[k] for k in keys if p.get(k) is not None), None)
            if value is not None:
                try: amounts[label] = str(marketplace_pipeline._number(value))
                except (TypeError, ValueError): pass
        rows.append({"external_id": external_id, "date": str(day), "operation": operation,
                     "category": category, "currency": currency, "amounts": amounts,
                     "spp": p.get("spp") or p.get("spp_percent") or p.get("sppPercent"),
                     "commission_rate": p.get("commissionPercent", p.get("commission_percent"))})
    rows.sort(key=lambda x: (x["date"], x["external_id"]))
    totals = {}
    for row in rows:
        cur = row["currency"]
        t = totals.setdefault(cur, {"records": 0, "categories": {}})
        t["records"] += 1
        t["categories"][row["category"]] = t["categories"].get(row["category"], 0) + 1
    # Financial fields are independent measures. Aggregate each one separately,
    # preserving native currency and operation category; never infer commission
    # or profit by subtracting overlapping WB amounts.
    from decimal import Decimal
    weekly = {}
    for row in rows:
        day = date.fromisoformat(row["date"])
        monday = day.fromordinal(day.toordinal() - day.weekday())
        key = (str(monday), row["currency"])
        bucket = weekly.setdefault(key, {"week_start": str(monday),
            "currency": row["currency"], "records": 0, "categories": {},
            "amounts": {}})
        bucket["records"] += 1
        cat = row["category"]
        bucket["categories"][cat] = bucket["categories"].get(cat, 0) + 1
        for field, raw in row["amounts"].items():
            # Store category-qualified amounts to prevent unrelated events from
            # being combined into a misleading payout or sales total.
            group = bucket["amounts"].setdefault(cat, {})
            group[field] = str(Decimal(group.get(field, "0")) + Decimal(raw))
    weekly_reports = sorted(weekly.values(), key=lambda v: (v["week_start"], v["currency"]))
    return {"start": str(start), "end": str(end), "total": len(rows),
            "weekly_reports": weekly_reports,
            "unclassified": sum(r["category"] == "unclassified" for r in rows),
            "by_currency": totals, "rows": rows[offset:offset + limit],
            "note": "Исходные поля WB без смешивания оборота, СПП, начислений и банковских выплат. Для окончательной сверки нужен официальный отчёт за тот же период."}


def _wb_week_args(week_start: date):
    from .wb_weekly import finance_today
    if week_start.weekday() != 0:
        raise HTTPException(422, "Неделя должна начинаться в понедельник")
    end = week_start + timedelta(days=6)
    if end >= finance_today():
        raise HTTPException(422, "Отчёт доступен только за завершённую неделю")
    return end


def _wb_week_account(db, u, account_id):
    q = select(Account).where(Account.channel == "Wildberries", Account.mode == "api", Account.active.is_(True))
    if u.role != "Owner":
        q = q.where(Account.company_id.in_(u.company_ids or []))
    if account_id is not None:
        q = q.where(Account.id == account_id)
    account = db.scalar(q.order_by(Account.id))
    if not account:
        raise HTTPException(404, "Подключение Wildberries не найдено")
    scope(u, account.company_id)
    return account


def _wb_finance_period_args(period: str, start: date, end: date):
    from .wb_weekly import _validate_report_period
    try:
        _validate_report_period(period, start, end)
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc
    return period


@app.get("/api/wildberries/finance")
def wildberries_finance_report(period: str, start: date, end: date, account_id: int | None = None,
                               u=Depends(current_user), db=Depends(get_db)):
    """Read a daily, weekly, or month-aggregated official WB Finance report."""
    require(u, "analytics")
    period = _wb_finance_period_args(period, start, end)
    account = _wb_week_account(db, u, account_id)
    from .wb_weekly import build_report
    return build_report(db, account, start, end, period=period)


@app.get("/api/wildberries/costs")
def wildberries_cost_catalog(account_id: int, u=Depends(current_user), db=Depends(get_db)):
    require(u, "finance")
    from .wb_costs import account_for_cost, catalogue_products, current_cost_map
    account = account_for_cost(db, u, account_id)
    items, _ = catalogue_products(db, account)
    today = business_today(); costs = current_cost_map(db, account.id, today)
    for item in items:
        row = costs.get(item["external_key"])
        item["unit_cost"] = str(Decimal(row.unit_cost) / 100) if row else None
        item["effective_from"] = str(row.effective_from) if row else None
    return {"account_id": account.id, "items": items}


@app.put("/api/wildberries/costs")
def wildberries_cost_save(p: s.WBProductCostIn, u=Depends(current_user), db=Depends(get_db)):
    require(u, "finance", True)
    from .wb_costs import account_for_cost, catalogue_products, cost_external_key, money_to_tiyn
    account = account_for_cost(db, u, p.account_id)
    items, _ = catalogue_products(db, account)
    stable = cost_external_key(p.nm_id, p.seller_article, p.barcode)
    product = next((x for x in items if x["external_key"] == p.external_key or (stable and x["external_key"] == stable)), None)
    if not product:
        raise HTTPException(422, "Товар не найден в сопоставленном каталоге этого кабинета WB")
    amount = money_to_tiyn(p.unit_cost)
    row = db.scalar(select(MarketplaceProductCost).where(MarketplaceProductCost.account_id == account.id,
        MarketplaceProductCost.external_key == product["external_key"], MarketplaceProductCost.effective_from == p.effective_from))
    if row:
        row.unit_cost, row.updated_by = amount, u.id
    else:
        db.add(MarketplaceProductCost(company_id=account.company_id, account_id=account.id,
            external_key=product["external_key"], nm_id=product["nm_id"], seller_article=product["seller_article"],
            barcode=product["barcode"], product_name=product["name"], unit_cost=amount,
            currency="KZT", effective_from=p.effective_from, updated_by=u.id))
    db.commit()
    return {"saved": True, "external_key": product["external_key"], "unit_cost": str(p.unit_cost), "effective_from": str(p.effective_from)}


@app.post("/api/wildberries/costs/bulk")
def wildberries_cost_bulk(p: s.WBProductCostBulkIn, u=Depends(current_user), db=Depends(get_db)):
    require(u, "finance", True)
    from .wb_costs import account_for_cost, catalogue_products, money_to_tiyn
    account = account_for_cost(db, u, p.account_id)
    amount = money_to_tiyn(p.unit_cost)
    items, _ = catalogue_products(db, account)
    known = {x["external_key"]: x for x in items}
    unknown = [key for key in set(p.external_keys) if key not in known]
    if unknown:
        raise HTTPException(422, {"message": "Есть товары вне каталога кабинета", "external_keys": unknown[:20]})
    for key in set(p.external_keys):
        product = known[key]
        row = db.scalar(select(MarketplaceProductCost).where(MarketplaceProductCost.account_id == account.id,
            MarketplaceProductCost.external_key == key, MarketplaceProductCost.effective_from == p.effective_from))
        if row:
            row.unit_cost, row.updated_by = amount, u.id
        else:
            db.add(MarketplaceProductCost(company_id=account.company_id, account_id=account.id, external_key=key,
                nm_id=product["nm_id"], seller_article=product["seller_article"], barcode=product["barcode"],
                product_name=product["name"], unit_cost=amount, currency="KZT", effective_from=p.effective_from, updated_by=u.id))
    db.commit()
    return {"saved": len(set(p.external_keys)), "effective_from": str(p.effective_from)}


@app.get("/api/wildberries/costs/template.xlsx")
def wildberries_cost_template(account_id: int, u=Depends(current_user), db=Depends(get_db)):
    require(u, "finance")
    from .wb_costs import account_for_cost, catalogue_products, current_cost_map
    account = account_for_cost(db, u, account_id)
    from openpyxl import Workbook
    from io import BytesIO
    book = Workbook(); ws = book.active; ws.title = "Себестоимость"
    ws.append(["Артикул WB", "Название", "Себестоимость, ₸", "Ключ товара (не менять)", "Дата начала действия"])
    today = business_today(); costs = current_cost_map(db, account.id, today)
    for item in catalogue_products(db, account)[0]:
        cost = costs.get(item["external_key"])
        ws.append([item["nm_id"] or item["seller_article"], item["name"], float(Decimal(cost.unit_cost)/100) if cost else None, item["external_key"], today.isoformat()])
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    out = BytesIO(); book.save(out)
    return Response(out.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="wb-cost-template.xlsx"'})


@app.post("/api/wildberries/costs/import")
def wildberries_cost_import(account_id: int = Form(...), file: UploadFile = File(...),
                            u=Depends(current_user), db=Depends(get_db)):
    require(u, "finance", True)
    from .wb_costs import account_for_cost, catalogue_products, money_to_tiyn
    from openpyxl import load_workbook
    account = account_for_cost(db, u, account_id)
    content = file.file.read(5 * 1024 * 1024 + 1)
    if len(content) > 5 * 1024 * 1024: raise HTTPException(413, "Файл превышает 5 МБ")
    try:
        ws = load_workbook(io.BytesIO(content), read_only=True, data_only=True).active
        headers = [str(c or "").strip().lower() for c in next(ws.iter_rows(values_only=True))]
        key_col = headers.index("ключ товара (не менять)") if "ключ товара (не менять)" in headers else -1
        nm_col = headers.index("артикул wb") if "артикул wb" in headers else -1
        cost_col = headers.index("себестоимость, ₸") if "себестоимость, ₸" in headers else -1
        date_col = headers.index("дата начала действия") if "дата начала действия" in headers else -1
        if cost_col < 0 or (key_col < 0 and nm_col < 0): raise ValueError("Не найдены колонки шаблона")
        items, _ = catalogue_products(db, account)
        by_nm = {x["nm_id"]: x for x in items if x["nm_id"]}
        by_key = {x["external_key"]: x for x in items}
        rows, errors, seen = [], [], set()
        for line, values in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
            if line > 1001:
                errors.append({"row": line, "error": "лимит импорта — 1000 товаров"})
                break
            try:
                key = str(values[key_col]).strip() if key_col >= 0 and values[key_col] else ""
                nm = str(values[nm_col]).strip() if nm_col >= 0 and values[nm_col] else ""
                product = by_key.get(key) if key else by_nm.get(nm)
                if not product: raise ValueError("товар не найден по ключу/nmID")
                amount = money_to_tiyn(values[cost_col])
                raw_date = values[date_col] if date_col >= 0 else date.today()
                effective = raw_date.date() if isinstance(raw_date, datetime) else (raw_date if isinstance(raw_date, date) else date.fromisoformat(str(raw_date)))
                ident = (product["external_key"], effective)
                if ident in seen: raise ValueError("повтор товара и даты в файле")
                seen.add(ident); rows.append((product, amount, effective))
            except Exception as exc: errors.append({"row": line, "error": str(exc)[:180]})
        if errors: return {"saved": 0, "errors": errors[:100], "error_count": len(errors)}
        for product, amount, effective in rows:
            row = db.scalar(select(MarketplaceProductCost).where(MarketplaceProductCost.account_id == account.id,
                MarketplaceProductCost.external_key == product["external_key"], MarketplaceProductCost.effective_from == effective))
            if row: row.unit_cost, row.updated_by = amount, u.id
            else: db.add(MarketplaceProductCost(company_id=account.company_id, account_id=account.id,
                external_key=product["external_key"], nm_id=product["nm_id"], seller_article=product["seller_article"],
                barcode=product["barcode"], product_name=product["name"], unit_cost=amount,
                currency="KZT", effective_from=effective, updated_by=u.id))
        db.commit(); return {"saved": len(rows), "errors": []}
    except HTTPException: raise
    except Exception as exc: raise HTTPException(422, f"Не удалось прочитать Excel: {str(exc)[:180]}") from exc


def _sync_wb_finance_background(account_id: int, user_id: int, period: str, start: date, end: date):
    from . import wb_weekly
    with SessionLocal() as db:
        account = db.get(Account, account_id)
        user = db.get(User, user_id)
        if not account or not user or not user.active or account.channel != "Wildberries" or account.mode != "api":
            return
        try:
            credentials = load_saved_credentials(db, account.id)
            token = credentials.get("token")
            if not token:
                raise AccessRequired("Для выбранного аккаунта WB не найден сохранённый токен. Переподключите интеграцию.")
            wb_weekly.sync_period(db, account, token, period, start, end, force=True, project_user=user)
        except Exception as exc:
            db.rollback()
            meta = wb_weekly._meta(db, account.id, start, create=True, period=period)
            if meta is not None:
                state = dict(meta.payload or {})
                failed_at = now().isoformat()
                safe_error = str(exc).strip()[:400] or "Не удалось подготовить синхронизацию WB. Проверьте токен и настройки."
                state.update(status="ошибка", error=safe_error,
                             start=str(start), end=str(end), period=period,
                             attempt_at=failed_at, updated_at=failed_at)
                meta.payload = state
                meta.stage_status = "ошибка"
                db.commit()
            logger.warning("WB Finance report sync failed: trigger=manual account_id=%s period=%s error=%s",
                           account_id, period, str(exc)[:400])


@app.post("/api/wildberries/finance/sync", status_code=202)
def wildberries_finance_sync(tasks: BackgroundTasks, period: str, start: date, end: date,
                             account_id: int | None = None, u=Depends(current_user), db=Depends(get_db)):
    require(u, "analytics")
    period = _wb_finance_period_args(period, start, end)
    account = _wb_week_account(db, u, account_id)
    tasks.add_task(_sync_wb_finance_background, account.id, u.id, period, start, end)
    return {"status": "загружается", "period": period, "start": str(start), "end": str(end)}


@app.get("/api/wildberries/finance.xlsx")
def wildberries_finance_xlsx(period: str, start: date, end: date, account_id: int | None = None,
                             u=Depends(current_user), db=Depends(get_db)):
    require(u, "analytics")
    period = _wb_finance_period_args(period, start, end)
    account = _wb_week_account(db, u, account_id)
    from .wb_weekly import build_report, create_xlsx
    report = build_report(db, account, start, end, period=period)
    if report.get("status") != "готов" or not report.get("row_count"):
        raise HTTPException(409, "Финансовый отчёт ещё не получен; скачивание недоступно")
    filename = f"WB-{period}-{start}-{end}.xlsx"
    return Response(content=create_xlsx(report), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'})


@app.get("/api/wildberries/weekly")
def wildberries_weekly_report(week_start: date, account_id: int | None = None,
                              u=Depends(current_user), db=Depends(get_db)):
    require(u, "analytics")
    end = _wb_week_args(week_start)
    account = _wb_week_account(db, u, account_id)
    from .wb_weekly import build_report
    return build_report(db, account, week_start, end)


def _sync_wb_week_background(account_id: int, user_id: int, week_start: date, week_end: date):
    """Run a user-requested report fetch in a fresh DB session."""
    from . import wb_weekly
    with SessionLocal() as db:
        account = db.get(Account, account_id)
        user = db.get(User, user_id)
        if not account or not user or not user.active or account.channel != "Wildberries" or account.mode != "api":
            return
        try:
            credentials = load_saved_credentials(db, account.id)
            token = credentials.get("token")
            wb_weekly.sync_week(db, account, token, week_start, week_end, force=True, project_user=user)
        except Exception:
            db.rollback()
            meta = wb_weekly._meta(db, account.id, week_start, create=True)
            if meta is not None:
                state = dict(meta.payload or {})
                state.update(status="ошибка", error="Не удалось подготовить синхронизацию WB. Проверьте токен и настройки.",
                             week_start=str(week_start), week_end=str(week_end), updated_at=now().isoformat())
                meta.payload = state
                meta.stage_status = "ошибка"
                db.commit()
            logger.exception("WB weekly on-demand sync failed for account %s", account_id)


@app.post("/api/wildberries/weekly/sync", status_code=202)
def wildberries_weekly_sync(tasks: BackgroundTasks, week_start: date, account_id: int | None = None,
                            u=Depends(current_user), db=Depends(get_db)):
    """Queue a fetch for one completed week; report data remains read-only until WB returns it."""
    require(u, "analytics")
    end = _wb_week_args(week_start)
    account = _wb_week_account(db, u, account_id)
    tasks.add_task(_sync_wb_week_background, account.id, u.id, week_start, end)
    return {"status": "загружается", "week_start": str(week_start), "week_end": str(end)}


@app.get("/api/wildberries/weekly.xlsx")
def wildberries_weekly_xlsx(week_start: date, account_id: int | None = None,
                            u=Depends(current_user), db=Depends(get_db)):
    require(u, "analytics")
    end = _wb_week_args(week_start)
    account = _wb_week_account(db, u, account_id)
    from .wb_weekly import build_report, create_xlsx
    report = build_report(db, account, week_start, end)
    if report.get("status") != "готов" or not report.get("row_count"):
        raise HTTPException(409, "Недельный финансовый отчёт ещё не получен; скачивание недоступно")
    filename = f"WB-weekly-{week_start}-{end}.xlsx"
    return Response(content=create_xlsx(report), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'})


class MarketplaceMappingIn(BaseModel):
    product_id: int
    external_keys: list[str] = Field(min_length=1, max_length=30)


class ProductMarketplaceIdIn(BaseModel):
    account_id: int
    id_type: Literal["product", "sku", "barcode", "size"]
    external_id: str = Field(min_length=1, max_length=180)
    confirm_reassign: bool = False


class ProductMarketplaceMappingDeleteIn(BaseModel):
    account_id: int
    external_key: str = Field(min_length=1, max_length=200)


class MarketplaceAutoMatchIn(BaseModel):
    company_id: int | None = Field(default=None, ge=1)
    product_ids: list[int] | None = Field(default=None, max_length=10000)


@app.post("/api/products/marketplace-matching/auto")
def auto_match_marketplace_products(p: MarketplaceAutoMatchIn, tasks: BackgroundTasks,
                                    u=Depends(current_user), db=Depends(get_db)):
    """Link exact, unambiguous seller SKUs to QOR products across accessible shops.

    This is additive and catalogue-only: it never replays orders, changes stock,
    or writes into the existing financial ledgers.
    """
    require(u, "catalog", True)
    account_q = scope_query(select(Account), Account, u).where(
        Account.active.is_(True), Account.channel.in_(["Kaspi", "Wildberries"]))
    if p.company_id is not None:
        scope(u, p.company_id)
        account_q = account_q.where(Account.company_id == p.company_id)
    accounts = list(db.scalars(account_q.order_by(Account.company_id, Account.channel, Account.id)))
    product_q = select(Product).where(Product.active.is_(True), Product.catalog_kind == "master")
    if p.product_ids is not None:
        selected_ids = list(dict.fromkeys(p.product_ids))
        product_q = product_q.where(Product.id.in_(selected_ids)) if selected_ids else product_q.where(Product.id == -1)
    products = list(db.scalars(product_q.order_by(Product.id)))
    result = marketplace_matching.match_staged_catalog(db, accounts, products)
    new_wb_pairs = result.pop("_new_wb_pairs", [])
    audit(db, u, "marketplace_auto_match", "products", None, {
        "company_id": p.company_id, "products_scanned": result["products_scanned"],
        "accounts_scanned": len(accounts), "links_created": result["links_created"],
        "review_count": result["review_count"],
    })
    db.commit()
    # A single newly added SKU can use the existing opt-in auto-sync path. The
    # bulk matcher only creates links; it must not enqueue thousands of outbound
    # stock writes from one button press.
    if p.product_ids and len(set(p.product_ids)) == 1:
        for account_id, product_id in new_wb_pairs:
            account = next((a for a in accounts if a.id == account_id), None)
            if account:
                _schedule_wb_sync_for_product(tasks, product_id, account.company_id)
    return {"ok": True, "accounts_scanned": len(accounts), **result}


@app.post("/api/products/marketplace-catalogs/wb/refresh")
def refresh_wb_product_catalogs(p: MarketplaceAutoMatchIn, tasks: BackgroundTasks,
                                u=Depends(current_user), db=Depends(get_db)):
    """Queue a catalog-only WB refresh; do not run financial projection."""
    require(u, "catalog", True)
    account_q = scope_query(select(Account), Account, u).where(
        Account.active.is_(True), Account.channel == "Wildberries", Account.mode == "api")
    if p.company_id is not None:
        scope(u, p.company_id)
        account_q = account_q.where(Account.company_id == p.company_id)
    accounts = list(db.scalars(account_q.order_by(Account.company_id, Account.id)))
    if not accounts:
        raise HTTPException(409, "Нет активных API-магазинов Wildberries для выбранной компании")
    account_ids = [a.id for a in accounts]
    for active_job in db.scalars(select(Job).where(Job.kind == "marketplace_catalog_refresh",
                                                  Job.status.in_(["queued", "running"]))):
        if set((active_job.payload or {}).get("account_ids") or []) & set(account_ids):
            if active_job.user_id != u.id:
                raise HTTPException(409, "Загрузка каталога этих магазинов уже выполняется")
            return {"ok": True, "job": dump(active_job), "already_running": True}
    job = Job(company_id=p.company_id if p.company_id else (accounts[0].company_id if len({a.company_id for a in accounts}) == 1 else None),
              user_id=u.id, kind="marketplace_catalog_refresh", status="queued",
              payload={"account_ids": account_ids, "accounts": {str(a.id): {"account_id": a.id, "account_name": a.name, "status": "queued", "pages": 0, "records": 0} for a in accounts}},
              result={"links_created": 0, "review_count": 0, "review": [], "per_account": []})
    db.add(job)
    db.flush()
    audit(db, u, "enqueue_marketplace_catalog_refresh", "jobs", job.id, {"account_ids": account_ids})
    db.commit()
    tasks.add_task(_run_wb_product_catalog_refresh, job.id)
    return {"ok": True, "job": dump(job), "already_running": False}


def _run_wb_product_catalog_refresh(job_id):
    """Fetch WB product cards in resumable pages and auto-link exact articles.

    Catalog pages are staged and committed independently. No sales are projected,
    and Account.sync_status / sync_cursor remain owned by the financial sync path.
    """
    with SessionLocal() as db:
        job = db.get(Job, job_id)
        if not job or job.kind != "marketplace_catalog_refresh":
            return
        job.status = "running"
        job.started_at = now()
        job.attempts = int(job.attempts or 0) + 1
        db.commit()
        user = db.get(User, job.user_id)
        payload = dict(job.payload or {})
        account_ids = list(payload.get("account_ids") or [])
        progress = dict(payload.get("accounts") or {})
        aggregate = dict(job.result or {})
        aggregate.setdefault("per_account", [])
        aggregate.setdefault("review", [])
        aggregate.setdefault("links_created", 0)
        aggregate.setdefault("products_linked", 0)
        aggregate.setdefault("review_count", 0)
        any_success = False
        for account_id in account_ids:
            account = db.get(Account, account_id)
            state = dict(progress.get(str(account_id)) or {})
            if state.get("status") == "done":
                any_success = True
                continue
            try:
                if not user or not user.active:
                    raise AccessRequired("Пользователь задания отключён")
                if not account or not account.active or account.channel != "Wildberries" or account.mode != "api":
                    raise AccessRequired("Магазин WB отключён или не использует API")
                if user.role != "Owner" and account.company_id not in (user.company_ids or []):
                    raise AccessRequired("Нет доступа к компании магазина WB")
                worker = connector(account, load_saved_credentials(db, account.id))
                cursor = state.get("cursor") or "{}"
                state.update({"status": "running", "error": ""})
                progress[str(account_id)] = state
                payload["accounts"] = progress
                job.payload = payload
                db.commit()
                pages = int(state.get("pages") or 0)
                records_count = int(state.get("records") or 0)
                inserted_count = int(state.get("inserted") or 0)
                updated_count = int(state.get("updated") or 0)
                for piece in worker.fetch_section_pages("products", cursor):
                    inserted_page, updated_page = _stage_marketplace_catalog_records(db, account, piece)
                    pages += 1
                    records_count += len(piece.records)
                    inserted_count += inserted_page
                    updated_count += updated_page
                    cursor = piece.next_cursor or cursor
                    state.update({"status": "running", "pages": pages, "records": records_count,
                                  "inserted": inserted_count, "updated": updated_count,
                                  "cursor": cursor, "updated_at": now().isoformat()})
                    progress[str(account_id)] = state
                    payload["accounts"] = progress
                    job.payload = payload
                    db.commit()
                products = list(db.scalars(select(Product).where(Product.active.is_(True)).order_by(Product.id)))
                matched = marketplace_matching.match_staged_catalog(db, [account], products)
                matched.pop("_new_wb_pairs", None)
                aggregate["links_created"] += matched["links_created"]
                aggregate["products_linked"] += matched["products_linked"]
                aggregate["review_count"] += matched["review_count"]
                aggregate["review"] = (aggregate["review"] + matched["review"])[:200]
                account_summary = matched["per_account"][0] if matched["per_account"] else {}
                aggregate["per_account"].append(account_summary)
                state.update({"status": "done", "pages": pages, "records": records_count,
                              "inserted": inserted_count, "updated": updated_count,
                              "links_created": matched["links_created"], "review_count": matched["review_count"],
                              "cursor": "", "updated_at": now().isoformat()})
                progress[str(account_id)] = state
                payload["accounts"] = progress
                job.payload = payload
                job.result = aggregate
                audit(db, user, "marketplace_catalog_refresh", "accounts", account.id,
                      {"pages": pages, "records": records_count, "links_created": matched["links_created"],
                       "review_count": matched["review_count"]})
                db.commit()
                any_success = True
            except Exception as exc:
                db.rollback()
                job = db.get(Job, job_id)
                payload = dict(job.payload or {})
                progress = dict(payload.get("accounts") or {})
                state = dict(progress.get(str(account_id)) or {})
                state.update({"status": "failed", "error": str(exc)[:400], "updated_at": now().isoformat()})
                progress[str(account_id)] = state
                payload["accounts"] = progress
                job.payload = payload
                db.commit()
                logger.exception("WB product catalog refresh failed for account %s", account_id)
        job = db.get(Job, job_id)
        statuses = [str((value or {}).get("status")) for value in (job.payload or {}).get("accounts", {}).values()]
        job.status = "partial" if any(status == "failed" for status in statuses) and any_success else (
            "failed" if all(status == "failed" for status in statuses) else "done")
        if job.status in {"partial", "failed"}:
            job.error = "; ".join(f"магазин {aid}: {row.get('error','ошибка')}" for aid, row in (job.payload or {}).get("accounts", {}).items() if row.get("status") == "failed")[:1000]
        db.commit()


@app.post("/api/products/{product_id}/marketplace-mappings")
def product_marketplace_mapping(product_id: int, p: ProductMarketplaceIdIn,
                                tasks: BackgroundTasks, u=Depends(current_user), db=Depends(get_db)):
    """Attach an external marketplace identifier and replay matching staged rows."""
    require(u, "inventory", True)
    product = db.get(Product, product_id)
    if not product or not product.active:
        raise HTTPException(404, "Активный товар QOR не найден")
    account = db.scalar(select(Account).where(Account.id == p.account_id).with_for_update())
    if not account or account.channel not in {"Kaspi", "Wildberries", "Ozon"}:
        raise HTTPException(404, "Магазин маркетплейса не найден")
    sync_state, _sync_error = sync_status_for_display(account)
    sync_heartbeat = account.last_sync
    if sync_heartbeat is not None and sync_heartbeat.tzinfo is None:
        sync_heartbeat = sync_heartbeat.replace(tzinfo=timezone.utc)
    abandoned_kaspi_period_sync = (
        sync_state == "running"
        and account.channel == "Kaspi"
        and str(account.sync_error or "").startswith("KASPI_PERIOD_SYNC")
        and (sync_heartbeat is None or sync_heartbeat < now() - timedelta(minutes=5))
    )
    if sync_state == "running" and not abandoned_kaspi_period_sync:
        raise HTTPException(409, "Синхронизация магазина ещё выполняется. Дождитесь завершения и повторите связь.")
    scope(u, account.company_id)
    product_company_ids = {int(value) for value in ((product.specifications or {}).get("_qor_company_ids") or []) if str(value).isdigit()}
    if product_company_ids and account.company_id not in product_company_ids:
        raise HTTPException(422, "Товар и магазин должны принадлежать одной компании")
    raw_id = p.external_id.strip()
    if not raw_id or any(ch.isspace() for ch in raw_id):
        raise HTTPException(422, "Укажите ID без пробелов")
    external_key = f"{p.id_type}:{raw_id}"
    existing = db.scalar(select(MarketplaceProductMapping).where(
        MarketplaceProductMapping.account_id == account.id,
        MarketplaceProductMapping.external_key == external_key,
    ))
    if existing and existing.product_id != product.id:
        if not p.confirm_reassign:
            raise HTTPException(409, "Этот ID уже привязан к другому SKU. Подтвердите перенос связи.")
        old_product_id = existing.product_id
        old_product = db.get(Product, old_product_id)
        if old_product and (old_product.specifications or {}).get("_qor_company_ids"):
            old_company_ids = {int(value) for value in ((old_product.specifications or {}).get("_qor_company_ids") or []) if str(value).isdigit()}
            if account.company_id not in old_company_ids:
                raise HTTPException(422, "Карточка маркетплейса и новый SKU должны принадлежать одной компании")
        existing.product_id = product.id
        audit(db, u, "mapping_reassign", "products", product.id, {
            "account_id": account.id, "channel": account.channel,
            "external_key": external_key, "from_product_id": old_product_id,
            "to_product_id": product.id,
        })
    if not existing:
        try:
            with db.begin_nested():
                db.add(MarketplaceProductMapping(company_id=account.company_id,
                    account_id=account.id, channel=account.channel,
                    external_key=external_key, product_id=product.id))
                db.flush()
                audit(db, u, "mapping", "products", product.id, {
                    "account_id": account.id, "channel": account.channel,
                    "external_key": external_key,
                })
        except IntegrityError:
            existing = db.scalar(select(MarketplaceProductMapping).where(
                MarketplaceProductMapping.account_id == account.id,
                MarketplaceProductMapping.external_key == external_key,
            ))
            if not existing or existing.product_id != product.id:
                raise HTTPException(409, "Этот ID уже привязан к другому товару QOR")
    # A saved link must immediately resolve matching staged sales. Otherwise
    # users see a successful save while the card and sale remain "needs link"
    # until an unrelated manual reprocess.
    matching_rows = []
    candidates = db.scalars(select(MarketplaceRecord).where(
        MarketplaceRecord.account_id == account.id,
        MarketplaceRecord.company_id == account.company_id,
        MarketplaceRecord.stage_status == "pending_mapping",
    )).all()
    for candidate in candidates:
        if external_key in marketplace_pipeline.identifiers(candidate.payload or {}):
            matching_rows.append(candidate)
    projected_now = 0
    pending_now = 0
    for candidate in matching_rows:
        ok, _reason = marketplace_pipeline.project_record(db, u, account, candidate)
        projected_now += int(ok)
        pending_now += int(candidate.stage_status in {"pending_fifo", "pending_fx", "pending_mapping"})
    db.commit()
    if account.channel == "Wildberries":
        # If the seller explicitly enabled WB auto-sync, a newly linked card or
        # size becomes eligible for the same guarded outbound sync as later edits.
        _schedule_wb_sync_for_product(tasks, product.id, account.company_id)
    return {"ok": True, "product_id": product.id, "account_id": account.id,
            "channel": account.channel, "external_key": external_key,
            "reprocessed_records": len(matching_rows), "projected_records": projected_now,
            "pending_records": pending_now}


@app.delete("/api/products/{product_id}/marketplace-mappings")
def delete_product_marketplace_mapping(product_id: int, p: ProductMarketplaceMappingDeleteIn,
                                       u=Depends(current_user), db=Depends(get_db)):
    require(u, "inventory", True)
    product = db.get(Product, product_id)
    if not product:
        raise HTTPException(404, "Товар QOR не найден")
    mapping = db.scalar(select(MarketplaceProductMapping).where(
        MarketplaceProductMapping.product_id == product_id,
        MarketplaceProductMapping.account_id == p.account_id,
        MarketplaceProductMapping.external_key == p.external_key,
    ))
    if not mapping:
        raise HTTPException(404, "Связь ID с этим товаром уже удалена")
    account = db.get(Account, mapping.account_id)
    if not account:
        raise HTTPException(404, "Магазин маркетплейса не найден")
    scope(u, account.company_id)
    db.delete(mapping)
    audit(db, u, "mapping_remove", "products", product_id, {
        "account_id": account.id, "channel": account.channel,
        "external_key": p.external_key,
    })
    db.commit()
    return {"ok": True, "product_id": product_id, "account_id": account.id,
            "external_key": p.external_key, "history_preserved": True}


@app.post("/api/marketplace-records/{id}/mapping")
def marketplace_mapping(id: int, p: MarketplaceMappingIn, u=Depends(current_user), db=Depends(get_db)):
    require(u, "accounts", True)
    row = db.get(MarketplaceRecord, id)
    if not row: raise HTTPException(404, "Запись не найдена")
    scope(u, row.company_id)
    acc = db.scalar(select(Account).where(Account.id == row.account_id).with_for_update())
    if acc.sync_status == 'running': raise HTTPException(409, 'Синхронизация ещё выполняется')
    product = db.get(Product, p.product_id)
    if not product or not product.active: raise HTTPException(422, "Товар не найден")
    payload = row.payload or {}
    permitted = set(marketplace_pipeline.identifiers(payload))
    for item in payload.get('products') or []:
        permitted.update(marketplace_pipeline.identifiers(item))
    if not set(p.external_keys).issubset(permitted):
        raise HTTPException(422, "Идентификатор отсутствует в исходной записи")
    for key in set(p.external_keys):
        mapping = db.scalar(select(MarketplaceProductMapping).where(
            MarketplaceProductMapping.account_id == acc.id, MarketplaceProductMapping.external_key == key))
        if mapping and mapping.product_id != product.id:
            raise HTTPException(409, "Идентификатор уже связан с другим товаром")
        if not mapping:
            db.add(MarketplaceProductMapping(company_id=acc.company_id, account_id=acc.id,
                channel=acc.channel, external_key=key, product_id=product.id))
    db.flush()
    audit(db, u, 'mapping', 'accounts', acc.id, {'product_id':product.id, 'keys':p.external_keys})
    result = marketplace_pipeline.project_account(db, u, acc)
    db.commit()
    return {'ok':True, **result}


@app.post("/api/accounts/{id}/reprocess")
def reprocess_marketplace(id: int, u=Depends(current_user), db=Depends(get_db)):
    require(u, 'accounts', True)
    acc = db.scalar(select(Account).where(Account.id == id).with_for_update())
    if not acc: raise HTTPException(404, 'Интеграция не найдена')
    scope(u, acc.company_id)
    if acc.sync_status == 'running': raise HTTPException(409, 'Синхронизация ещё выполняется')
    result = marketplace_pipeline.project_account(db, u, acc)
    db.commit()
    return result


class SettlementIn(BaseModel):
    effective_date: date
    currency: str = Field(pattern=r'^[A-Z]{3}$')
    fx_rate: int = Field(gt=0, le=10**10)


@app.put('/api/accounts/{id}/settlement')
def account_settlement(id: int, p: SettlementIn, u=Depends(current_user), db=Depends(get_db)):
    require(u, 'accounts', True)
    acc = db.scalar(select(Account).where(Account.id == id).with_for_update())
    if not acc: raise HTTPException(404, 'Интеграция не найдена')
    scope(u, acc.company_id)
    if acc.sync_status == 'running': raise HTTPException(409, 'Синхронизация ещё выполняется')
    if p.currency == 'KZT' and p.fx_rate != 10000:
        raise HTTPException(422, 'Курс KZT должен быть 10000')
    existing=db.scalar(select(AccountExchangeRate).where(AccountExchangeRate.account_id==id,
        AccountExchangeRate.currency==p.currency,AccountExchangeRate.rate_date==p.effective_date))
    if existing and existing.rate_scaled!=p.fx_rate: raise HTTPException(409,'Курс на эту дату уже зафиксирован')
    if not existing: db.add(AccountExchangeRate(account_id=id,currency=p.currency,rate_date=p.effective_date,rate_scaled=p.fx_rate,source='manual_confirmed'))
    acc.settlement_currency, acc.settlement_fx_rate = p.currency, p.fx_rate
    audit(db, u, 'settlement', 'accounts', id, p.model_dump())
    db.commit()
    return {'ok':True}


@app.post('/api/post/opening')
def post_opening(p:s.OpeningInventoryIn,u=Depends(current_user),db=Depends(get_db)):
    require(u,'inventory',True)
    obj=canonical.post_event(db,u,'opening',p);db.commit();return dump(obj)


@app.post('/api/admin/reconcile-marketplaces')
def reconcile_existing(u=Depends(current_user),db=Depends(get_db)):
    require(u,'accounts',True)
    results=marketplace_pipeline.reconcile_all_staged_marketplaces(db,u)
    audit(db,u,'reconcile','marketplace_records',details={'results':results});db.commit()
    return {'ok':True,'results':results}


@app.get("/api/data/{entity}")
def listing(
    entity: str,
    search: str = "",
    company_id: int | None = None,
    account_id: int | None = None,
    catalog_kind: str | None = None,
    channel: str | None = None,
    start: date | None = None,
    end: date | None = None,
    offset: int = Query(0, ge=0),
    limit: int = Query(50, ge=1, le=500),
    u=Depends(current_user),
    db=Depends(get_db),
):
    q = access_query(db, u, entity, search, company_id, account_id)
    model = MODELS[entity][0]
    if entity == "products" and catalog_kind is not None:
        if catalog_kind not in {"master", "marketplace"}:
            raise HTTPException(422, "Неизвестный тип каталога")
        q = q.where(Product.catalog_kind == catalog_kind)
    if entity in {"sales", "expenses", "cash", "purchases"} and (start is not None or end is not None):
        if start is None or end is None or end < start:
            raise HTTPException(422, "Укажите корректный диапазон дат")
        q = q.where(model.date >= start, model.date <= end)
    if channel and entity == "sales":
        q = q.where(
            Sale.account_id.in_(select(Account.id).where(Account.channel == channel))
        )
    count = db.scalar(select(func.count()).select_from(q.subquery()))
    rows = []
    for obj in db.scalars(q.order_by(model.id.desc()).offset(offset).limit(limit)):
        row = dump(obj)
        if entity == "accounts":
            row["sync_status"], row["sync_error"] = sync_status_for_display(obj)
        rows.append(row)
    if entity == "niches":
        for row in rows:
            row["analysis"] = ai.niche_score(row["data"])
    if entity == "batches" and rows:
        batch_ids = [r["id"] for r in rows]
        allocation_rows = db.execute(
            select(Allocation.batch_id, func.coalesce(func.sum(Allocation.quantity), 0),
                   func.coalesce(func.sum(Allocation.returned), 0))
            .where(Allocation.batch_id.in_(batch_ids))
            .group_by(Allocation.batch_id)
        ).all()
        allocations = {int(batch_id): (int(sold or 0), int(returned or 0))
                       for batch_id, sold, returned in allocation_rows}
        for row in rows:
            sold, returned = allocations.get(row["id"], (0, 0))
            row["sold_quantity"] = sold
            row["returned_quantity"] = returned
            row["net_sold_quantity"] = max(0, sold - returned)
    if entity == "products" and rows:
        product_ids = [r["id"] for r in rows]
        from_date = start or business_today().replace(day=1)
        to_date = end or business_today()
        if to_date < from_date:
            raise HTTPException(422, "Конец периода раньше начала")
        batch_q = scope_query(select(Batch.product_id, func.coalesce(func.sum(Batch.remaining), 0)), Batch, u)
        batch_q = batch_q.where(Batch.product_id.in_(product_ids)).group_by(Batch.product_id)
        if company_id:
            batch_q = batch_q.where(Batch.company_id == company_id)
        physical = dict(db.execute(batch_q).all())
        sold_q = scope_query(select(Sale.product_id, func.coalesce(func.sum(Sale.quantity), 0)), Sale, u)
        sold_q = sold_q.where(Sale.product_id.in_(product_ids), Sale.date >= from_date, Sale.date <= to_date).group_by(Sale.product_id)
        if company_id:
            sold_q = sold_q.where(Sale.company_id == company_id)
        sold = dict(db.execute(sold_q).all())
        returned_q = select(Sale.product_id, func.coalesce(func.sum(Return.quantity), 0)).join(Return, Return.sale_id == Sale.id)
        returned_q = scope_query(returned_q, Sale, u).where(Sale.product_id.in_(product_ids), Return.date >= from_date, Return.date <= to_date).group_by(Sale.product_id)
        if company_id:
            returned_q = returned_q.where(Sale.company_id == company_id)
        returned = dict(db.execute(returned_q).all())
        channel_sales_q = scope_query(select(
            Sale.product_id, Account.id, Account.channel, Account.name,
            func.coalesce(func.sum(Sale.quantity), 0),
        ).join(Account, Account.id == Sale.account_id), Sale, u)
        channel_sales_q = channel_sales_q.where(
            Sale.product_id.in_(product_ids), Sale.date >= from_date,
            Sale.date <= to_date, Account.active.is_(True),
        ).group_by(Sale.product_id, Account.id, Account.channel, Account.name)
        if company_id:
            channel_sales_q = channel_sales_q.where(Sale.company_id == company_id)
        channel_sales = {(pid, aid): {"channel": channel, "account_name": name, "sold": int(qty or 0), "returned": 0}
                         for pid, aid, channel, name, qty in db.execute(channel_sales_q).all()}
        channel_returns_q = scope_query(select(
            Sale.product_id, Account.id, Account.channel, Account.name,
            func.coalesce(func.sum(Return.quantity), 0),
        ).join(Return, Return.sale_id == Sale.id).join(Account, Account.id == Sale.account_id), Sale, u)
        channel_returns_q = channel_returns_q.where(
            Sale.product_id.in_(product_ids), Return.date >= from_date,
            Return.date <= to_date, Account.active.is_(True),
        ).group_by(Sale.product_id, Account.id, Account.channel, Account.name)
        if company_id:
            channel_returns_q = channel_returns_q.where(Sale.company_id == company_id)
        for pid, aid, channel, name, qty in db.execute(channel_returns_q).all():
            channel_sales.setdefault((pid, aid), {"channel": channel, "account_name": name, "sold": 0, "returned": 0})["returned"] = int(qty or 0)
        defective_q = select(Sale.product_id, func.coalesce(func.sum(Return.quantity), 0)).join(Return, Return.sale_id == Sale.id)
        defective_q = scope_query(defective_q, Sale, u).where(
            Sale.product_id.in_(product_ids), Return.inventory_status == "defective"
        ).group_by(Sale.product_id)
        if company_id:
            defective_q = defective_q.where(Sale.company_id == company_id)
        defective = dict(db.execute(defective_q).all())
        accounts_q = scope_query(select(Account), Account, u)
        if company_id:
            accounts_q = accounts_q.where(Account.company_id == company_id)
        accessible_accounts = db.scalars(accounts_q).all()
        accessible_accounts_by_id = {a.id: a for a in accessible_accounts}
        accessible_account_ids = {a.id for a in accessible_accounts}
        active_marketplace_accounts = [a for a in accessible_accounts if a.active and a.channel in {"Kaspi", "Wildberries", "Ozon"}]
        active_account_ids = {a.id for a in accessible_accounts if a.active}
        sync_times = [a.last_sync for a in accessible_accounts if a.last_sync]
        latest_sync = max(sync_times) if sync_times else None
        map_q = scope_query(select(MarketplaceProductMapping), MarketplaceProductMapping, u)
        if company_id:
            map_q = map_q.where(MarketplaceProductMapping.company_id == company_id)
        mapping_rows = db.scalars(map_q.where(MarketplaceProductMapping.product_id.in_(product_ids))).all()
        for row in rows:
            specs = row.get("specifications") or {}
            market_snapshot = specs.get("marketplace_stock") or {}
            visible_snapshot = {}
            for key, value in market_snapshot.items():
                try: snapshot_account_id = int(str(key).rsplit(":", 1)[1])
                except (IndexError, TypeError, ValueError): continue
                if snapshot_account_id in active_account_ids and isinstance(value, dict):
                    visible_snapshot[key] = value
            marketplace_units = sum(max(0, int(v.get("quantity") or 0)) for v in visible_snapshot.values())
            physical_units = int(physical.get(row["id"], 0) or 0)
            defective_units = int(defective.get(row["id"], 0) or 0)
            in_transit = max(0, int(specs.get("in_transit_quantity") or 0))
            reserved = max(0, int(specs.get("reserved_quantity") or 0))
            allocations = specs.get("marketplace_stock_allocations") or {}
            allocated_units = 0
            for account_key, allocation in allocations.items():
                try: allocation_account_id = int(account_key)
                except (TypeError, ValueError): continue
                if allocation_account_id not in accessible_account_ids: continue
                if not isinstance(allocation, dict):
                    try: allocated_units += max(0, int(allocation or 0))
                    except (TypeError, ValueError): pass
                    continue
                try: allocated_units += max(0, int(allocation.get("quantity") or 0))
                except (TypeError, ValueError): pass
                for amount in (allocation.get("sizes") or {}).values():
                    try: allocated_units += max(0, int(amount or 0))
                    except (TypeError, ValueError): pass
            available_units = max(0, physical_units - reserved - allocated_units)
            # Marketplace snapshots mirror listing/account availability; absent a
            # transfer ledger they cannot be safely added to QOR batches without
            # risking that the same units are counted twice.
            total_units = physical_units + in_transit + defective_units
            product_mappings = [m for m in mapping_rows if m.product_id == row["id"]]
            mapped_account_ids = {m.account_id for m in product_mappings}
            linked_count = sum(1 for account in active_marketplace_accounts if account.id in mapped_account_ids)
            marketplace_count = len(active_marketplace_accounts)
            coverage_status = "no_marketplaces" if marketplace_count == 0 else (
                "complete" if linked_count == marketplace_count else ("partial" if linked_count else "unmapped")
            )
            relevant_sales = {aid: data for (pid, aid), data in channel_sales.items() if pid == row["id"]}
            for account in active_marketplace_accounts:
                if account.id in mapped_account_ids:
                    relevant_sales.setdefault(account.id, {"channel": account.channel, "account_name": account.name, "sold": 0, "returned": 0})
            snapshot_review = bool(visible_snapshot)
            allocation_overdrawn = allocated_units > max(0, physical_units - reserved)
            if coverage_status == "no_marketplaces":
                reconciliation_status = "no_marketplaces"
            else:
                reconciliation_status = "partial" if coverage_status in {"partial", "unmapped"} or snapshot_review or allocation_overdrawn else "confirmed"
            row["inventory"] = {
                "total": total_units,
                "physical": physical_units,
                "available": available_units,
                "reserved": reserved,
                "allocated": allocated_units,
                "marketplace": marketplace_units,
                "in_transit": in_transit,
                "defective": defective_units,
                "sold": int(sold.get(row["id"], 0) or 0),
                "returned": int(returned.get(row["id"], 0) or 0),
                "sales_by_marketplace": [
                    {"account_id": aid, **data} for aid, data in sorted(relevant_sales.items())
                ],
                "mapping_coverage": {"linked": linked_count, "active_accounts": marketplace_count, "status": coverage_status},
                "reconciliation_status": reconciliation_status,
                "snapshot_review_required": snapshot_review,
                "allocation_overdrawn": allocation_overdrawn,
                "last_sync": latest_sync,
                "low_stock": available_units <= int(row.get("min_stock") or 0),
            }
        for row in rows:
            row["marketplace_mappings"] = [
                {"account_id": m.account_id, "channel": m.channel, "external_key": m.external_key}
                for m in mapping_rows if m.product_id == row["id"]
            ]
            product_mapping_rows = [m for m in mapping_rows if m.product_id == row["id"]]
            mapping_by_account = {}
            for mapping in product_mapping_rows:
                mapping_by_account.setdefault(mapping.account_id, []).append(mapping.external_key)
            row["marketplace_links"] = []
            for account in active_marketplace_accounts:
                keys = sorted(set(mapping_by_account.get(account.id, [])))
                status, _sync_error = sync_status_for_display(account)
                row["marketplace_links"].append({
                    "account_id": account.id,
                    "channel": account.channel,
                    "account_name": account.name,
                    "active": bool(account.active),
                    "mapping_status": "linked" if keys else "needs_link",
                    "mapping_count": len(keys),
                    "external_keys": keys,
                    "sync_status": status,
                    "sync_attention": status in {"failed", "partial"},
                    "sync_status_label": {
                        "failed": "Ошибка синхронизации",
                        "partial": "Синхронизация частичная",
                        "running": "Синхронизация выполняется",
                        "done": "Синхронизировано",
                        "never": "Ожидает первой синхронизации",
                    }.get(status, "Статус синхронизации не подтверждён"),
                    "sync_attention_message": (
                        "Проверьте магазин в разделе «Интеграции»." if status in {"failed", "partial"} else ""
                    ),
                })
            for account_id, keys in mapping_by_account.items():
                if any(link["account_id"] == account_id for link in row["marketplace_links"]):
                    continue
                mapping = next(m for m in product_mapping_rows if m.account_id == account_id)
                mapped_account = accessible_accounts_by_id.get(account_id)
                row["marketplace_links"].append({
                    "account_id": account_id,
                    "channel": mapping.channel,
                    "account_name": mapped_account.name if mapped_account else f"магазин {account_id}",
                    "active": False,
                    "mapping_status": "inactive",
                    "mapping_count": len(set(keys)),
                    "external_keys": sorted(set(keys)),
                    "sync_status": "inactive",
                    "sync_attention": False,
                    "sync_status_label": "Интеграция отключена",
                    "sync_attention_message": "",
                })
            row["marketplace_links"].sort(key=lambda link: (link["channel"], link["account_name"], link["account_id"]))
    return {"items": rows, "total": count}


def validate_master(db, u, entity, p, existing=None):
    if hasattr(p, "company_id"):
        f.get_company(db, u, p.company_id)
    if entity == "products":
        scoped_companies = (p.specifications or {}).get("_qor_company_ids", [])
        if not isinstance(scoped_companies, list) or len(scoped_companies) > 50:
            raise HTTPException(422, "Некорректная привязка товара к компаниям")
        for company_id in scoped_companies:
            try:
                company_id = int(company_id)
            except (TypeError, ValueError):
                raise HTTPException(422, "Некорректная компания товара")
            if not db.get(Company, company_id):
                raise HTTPException(422, "Компания товара не найдена")
            scope(u, company_id)
    if entity == "companies" and u.role != "Owner":
        raise HTTPException(403, "Компании изменяет Owner")
    if entity == "users":
        if u.role != "Owner":
            raise HTTPException(403, "Назначение ролей доступно Owner")
        if existing and existing.id == u.id and (not p.active or p.role != "Owner"):
            raise HTTPException(422, "Нельзя отключить себя или снять свою роль Owner")
        if not existing and len(p.password) < 12:
            raise HTTPException(422, "Пароль от 12 символов")
        if p.password and len(p.password) < 12:
            raise HTTPException(422, "Пароль от 12 символов")
        for cid in p.company_ids:
            if not db.get(Company, cid):
                raise HTTPException(422, "Неизвестная компания")
    if entity == "accounts" and existing:
        linked = db.scalar(
            select(Sale.id).where(Sale.account_id == existing.id).limit(1)
        )
        if linked and (
            p.company_id != existing.company_id or p.channel != existing.channel
        ):
            raise HTTPException(
                422, "Нельзя менять компанию или канал аккаунта с продажами"
            )
    if entity == "niches":
        ai.niche_score(p.data)


def master_values(p, entity, existing=None):
    values = p.model_dump()
    # Product version is an optimistic concurrency token, not an editable field.
    if entity == "products":
        values.pop("version", None)
        submitted_catalog_kind = values.pop("catalog_kind", None)
        values["catalog_kind"] = submitted_catalog_kind or (existing.catalog_kind if existing is not None else "master")
        description = values.pop("description", "")
        barcode = values.pop("barcode", "")
        purchase_cost = values.pop("purchase_cost", 0)
        custom_fields = values.pop("custom_fields", {})
        marketplace_stock_allocations = values.pop("marketplace_stock_allocations", None)
        specs = dict(values.get("specifications") or {})
        old_specs = dict((existing.specifications or {}) if existing is not None else {})
        history = list(old_specs.get("purchase_cost_history") or [])
        old_cost = int(old_specs.get("purchase_cost") or 0)
        today = business_today()
        if existing is None and purchase_cost > 0:
            history.append({"effective_from": str(today), "unit_cost": int(purchase_cost)})
        elif existing is not None and int(purchase_cost or 0) == old_cost and old_cost > 0 and not history:
            # Start tracking an existing SKU from today without rewriting its
            # untracked historical sales.
            history.append({"effective_from": str(today), "unit_cost": old_cost})
        elif existing is not None and int(purchase_cost or 0) != old_cost:
            # Keep today's already sold units on the previous value. A changed
            # SKU cost starts on the next Almaty business day; historical batches
            # and posted sale COGS remain untouched.
            if old_cost > 0 and not any(
                isinstance(item, dict) and item.get("effective_from") == str(today)
                for item in history
            ):
                history.append({"effective_from": str(today), "unit_cost": old_cost})
            if int(purchase_cost or 0) >= 0:
                # If there was no known previous SKU cost, the first confirmed
                # value applies today. Replacements take effect tomorrow.
                effective = str(today + timedelta(days=1) if old_cost > 0 else today)
                history = [item for item in history if not (
                    isinstance(item, dict) and item.get("effective_from") == effective
                )]
                history.append({"effective_from": effective, "unit_cost": int(purchase_cost)})
        history = sorted(
            [item for item in history if isinstance(item, dict) and item.get("effective_from") and item.get("unit_cost") is not None],
            key=lambda item: str(item["effective_from"]),
        )
        specs.update({"description": description, "barcode": barcode,
                      "purchase_cost": purchase_cost, "purchase_cost_history": history,
                      "custom_fields": custom_fields})
        if marketplace_stock_allocations is not None:
            specs["marketplace_stock_allocations"] = marketplace_stock_allocations
        values["specifications"] = specs
    if entity == "users":
        pw = values.pop("password")
        values["email"] = values["email"].lower()
        if pw:
            values["password_hash"] = passwords.hash(pw)
    return values


@app.post("/api/products/ai-draft")
def product_ai_draft(p: dict, u=Depends(current_user)):
    """Generate an editable product draft; the caller must review and save it."""
    require(u, "catalog", True)
    return ai.product_draft(p)


@app.post("/api/data/{entity}")
def create_master(
    entity: str, payload: dict, tasks: BackgroundTasks, u=Depends(current_user), db=Depends(get_db)
):
    model, cap, schema = model_info(entity)
    require(u, cap, True)
    if not schema:
        raise HTTPException(405, "Используйте проведение документа")
    p = schema.model_validate(payload)
    validate_master(db, u, entity, p)
    obj = model(**master_values(p, entity))
    db.add(obj)
    db.flush()
    audit(db, u, "create", entity, obj.id, {"after": dump(obj)})
    db.commit()
    if entity == "products": _schedule_wb_sync_for_product(tasks, obj.id)
    return dump(obj)


@app.put("/api/data/{entity}/{id}")
def edit_master(
    entity: str, id: int, payload: dict, tasks: BackgroundTasks, u=Depends(current_user), db=Depends(get_db)
):
    model, cap, schema = model_info(entity)
    require(u, cap, True)
    if not schema:
        raise HTTPException(405, "Проведённые финансовые записи неизменяемы")
    obj = db.get(model, id)
    if not obj:
        raise HTTPException(404, "Не найдено")
    if entity == "products":
        obj = db.scalar(select(Product).where(Product.id == id).with_for_update().execution_options(populate_existing=True))
    if hasattr(obj, "company_id"):
        scope(u, obj.company_id)
    p = schema.model_validate(payload)
    if entity == "products" and p.version is None:
        raise HTTPException(428, "Для сохранения товара нужна версия карточки. Обновите список товаров.")
    if entity == "products" and p.version != obj.version:
        raise HTTPException(409, "Товар уже изменён другим сотрудником. Обновите карточку и повторите правку.")
    validate_master(db, u, entity, p, obj)
    if entity == "customers" and p.company_id != obj.company_id:
        raise HTTPException(422, "Компания покупателя неизменяема")
    before = dump(obj)
    for k, v in master_values(p, entity, obj).items():
        setattr(obj, k, v)
    if entity == "products":
        obj.version += 1
    audit(db, u, "update", entity, id, {"before": before, "after": dump(obj)})
    db.commit()
    if entity == "products": _schedule_wb_sync_for_product(tasks, obj.id)
    return dump(obj)


@app.delete("/api/data/{entity}/{id}")
def delete_master(entity: str, id: int, u=Depends(current_user), db=Depends(get_db)):
    model, cap, schema = model_info(entity)
    require(u, cap, True)
    if not schema or entity == "users":
        raise HTTPException(
            405,
            "Удаление запрещено; отключите пользователя или используйте корректирующий документ",
        )
    if entity == "companies" and u.role != "Owner":
        raise HTTPException(403, "Только Owner")
    obj = db.get(model, id)
    if not obj:
        raise HTTPException(404, "Не найдено")
    if hasattr(obj, "company_id"):
        scope(u, obj.company_id)
    if entity == "accounts":
        secret = db.scalar(select(IntegrationSecret).where(IntegrationSecret.account_id == id))
        if secret:
            db.delete(secret)
        obj.active = False
        obj.sync_status = "never"
        obj.sync_error = ""
        audit(db, u, "archive", entity, id, {"name": obj.name, "channel": obj.channel})
        db.commit()
        return {"ok": True, "archived": True}
    before = dump(obj)
    db.delete(obj)
    audit(db, u, "delete", entity, id, {"before": before})
    db.commit()
    return {"ok": True}


OPERATIONS = {
    "purchases": (s.PurchaseIn, f.purchase, "inventory"),
    "sales": (s.SaleIn, f.sale, "sales"),
    "returns": (s.ReturnIn, f.return_sale, "sales"),
    "expenses": (s.ExpenseIn, f.expense, "finance"),
    "cash": (s.CashIn, f.payment, "finance"),
    "adjustments": (s.AdjustIn, f.adjust, "inventory"),
}


@app.post("/api/post/{kind}")
def post(kind: str, payload: dict, tasks: BackgroundTasks, u=Depends(current_user), db=Depends(get_db)):
    if kind not in OPERATIONS:
        raise HTTPException(404, "Операция не найдена")
    schema, fn, cap = OPERATIONS[kind]
    require(u, cap, True)
    parsed = schema.model_validate(payload)
    affected = []
    affected_company = getattr(parsed, "company_id", None)
    if kind == "purchases": affected = [x.product_id for x in parsed.lines]
    elif kind in {"sales", "adjustments"}: affected = [parsed.product_id]
    elif kind == "returns":
        sale_row = db.get(Sale, parsed.sale_id)
        if sale_row:
            affected = [sale_row.product_id]
            affected_company = sale_row.company_id
    obj = canonical.post_event(db, u, {"purchases":"purchase", "sales":"sale", "returns":"return", "expenses":"expense", "cash":"payment", "adjustments":"stock"}[kind], parsed, reject_duplicates=True)
    db.commit()
    if kind in {"purchases", "sales", "returns", "adjustments"}:
        for product_id in set(affected): _schedule_wb_sync_for_product(tasks, product_id, affected_company)
    return dump(obj)


@app.post("/api/returns/{return_id}/status")
def update_return_status(return_id: int, payload: dict, tasks: BackgroundTasks, u=Depends(current_user), db=Depends(get_db)):
    require(u, "inventory", True)
    ret = db.get(Return, return_id)
    if not ret:
        raise HTTPException(404, "Возврат не найден")
    scope(u, ret.company_id)
    sale = db.get(Sale, ret.sale_id)
    if not sale:
        raise HTTPException(409, "Исходная продажа не найдена")
    f.lock_product(db, sale.product_id)
    ret = db.scalar(select(Return).where(Return.id == return_id).with_for_update())
    change = s.ReturnStatusIn.model_validate(payload)
    transitions = {
        "requested": {"approved"},
        "approved": {"in_transit"},
        "in_transit": {"received"},
        "received": {"inspection"},
        "inspection": {"restocked", "defective"},
        "defective": {"written_off"},
        "restocked": set(),
        "written_off": set(),
    }
    if change.status != ret.inventory_status:
        if change.status not in transitions.get(ret.inventory_status, set()):
            raise HTTPException(409, "Недопустимый переход статуса возврата")
        if change.status in {"defective", "written_off"} and len(change.reason.strip()) < 3:
            raise HTTPException(422, "Для брака и списания укажите причину")
        before = ret.inventory_status
        ret.inventory_status = change.status
        audit(db, u, "inventory_status", "returns", ret.id, {
            "from": before, "to": change.status, "reason": change.reason.strip(),
        })
    if change.status == "restocked":
        f.restock_return_inventory(db, u, ret)
    db.commit()
    if change.status == "restocked": _schedule_wb_sync_for_product(tasks, sale.product_id, sale.company_id)
    return dump(ret)


@app.get("/api/accounting/summary")
def accounting_summary(company_id: int, start: date, end: date, u=Depends(current_user), db=Depends(get_db)):
    require(u, "finance")
    scope(u, company_id)
    if end < start:
        raise HTTPException(422, "Некорректный период")
    bank = list(db.scalars(select(BankTransaction).where(BankTransaction.company_id == company_id, BankTransaction.date >= start, BankTransaction.date <= end)))
    incoming = sum(x.amount for x in bank if x.direction == "in")
    outgoing = sum(x.amount for x in bank if x.direction == "out")
    # Receivables are authoritative only in the central finance report: it also
    # subtracts confirmed returns and therefore cannot drift from Finance/AI.
    receivable = f.report(db, u, date.min, end, company_id)["b2b_debt"]
    purchase_paid = {pid: -amt for pid, amt in db.execute(select(CashEntry.purchase_id, func.coalesce(func.sum(CashEntry.amount), 0)).where(CashEntry.purchase_id.is_not(None), CashEntry.date <= end).group_by(CashEntry.purchase_id))}
    purchases = list(db.scalars(select(Purchase).where(Purchase.company_id == company_id, Purchase.date <= end)))
    payable = sum(max(0, x.total - max(0, purchase_paid.get(x.id, 0))) for x in purchases)
    rules = [dump(x) for x in db.scalars(select(TaxRule).where(TaxRule.company_id == company_id, TaxRule.active == True, TaxRule.valid_from <= end))]
    return {"period": {"start": start, "end": end}, "bank_in": incoming, "bank_out": outgoing, "bank_net": incoming-outgoing, "receivable": receivable, "payable": payable, "tax_rules": rules, "tax_notice": "Ставки являются конфигурацией. Налоговый расчет требует проверки бухгалтером."}

@app.get("/api/report")
def get_report(
    start: date,
    end: date,
    company_id: int | None = None,
    account_id: int | None = None,
    u=Depends(current_user),
    db=Depends(get_db),
    active_only: bool = False,
):
    require(u, "analytics")
    return f.report(db, u, start, end, company_id, account_id, active_only)


@app.get("/api/reports/executive.xlsx")
def executive_xlsx(
    start: date,
    end: date,
    company_id: int | None = None,
    account_id: int | None = None,
    u=Depends(current_user),
    db=Depends(get_db),
):
    require(u, "analytics")
    if end < start:
        raise HTTPException(422, "Конечная дата раньше начальной")
    report = f.report(db, u, start, end, company_id, account_id)
    raw = reporting_xlsx.executive_workbook(report, start, end)
    filename = f"QOR-report-{start}-{end}.xlsx"
    return Response(
        raw,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.get("/api/dashboard/summary")
def dashboard_summary(
    start: date,
    end: date,
    company_id: int | None = None,
    u=Depends(current_user),
    db=Depends(get_db),
):
    """Dashboard contract from the client specification.

    Returns selected-period metrics plus immutable today/month snapshots so the
    headline KPIs never change meaning when the user changes the period filter.
    All values come from the deterministic finance engine.
    """
    require(u, "analytics")
    today = business_today()
    # The selected report is loaded separately for charts and drilldowns. This
    # endpoint returns only date-scoped KPI data so the overview can paint quickly.
    selected = f.report_headline(db, u, start, end, company_id)
    month_start = today.replace(day=1)
    month_report = (
        selected if (start, end) == (month_start, today)
        else f.report_headline(db, u, month_start, today, company_id)
    )
    today_report = (
        selected if (start, end) == (today, today)
        else f.report_headline(db, u, today, today, company_id)
    )

    def channels(report):
        found = {str(x.get("channel", "")): x for x in report.get("channel_summary", [])}
        values = {
            name: int(found.get(name, {}).get("revenue", 0) or 0)
            for name in ("Kaspi", "Wildberries", "Ozon", "Offline", "B2B")
        }
        values["WB"] = values["Wildberries"]
        return values

    return {
        "as_of": str(today),
        "today": {
            "revenue": today_report["revenue"],
            "net_profit": today_report["net_profit"],
            "units": today_report["units"],
        },
        "month": {
            "revenue": month_report["revenue"],
            "gross_profit": month_report["gross_profit"],
            "net_profit": month_report["net_profit"],
            "margin": month_report["margin"],
        },
        "report": selected,
        "selected": {
            "revenue": selected["revenue"],
            "gross_profit": selected["gross_profit"],
            "net_profit": selected["net_profit"],
            "margin": selected["margin"],
            "units": selected["units"],
            "returns": selected.get("returns"),
            "advertising": selected.get("advertising"),
            "expenses": selected.get("expenses"),
            "stock_value": selected.get("stock_value"),
            "cash": selected.get("cash"),
            "b2b_debt": selected.get("b2b_debt"),
            "channels": channels(selected),
        },
    }


@app.get("/api/report/details")
def get_report_details(
    start: date,
    end: date,
    metric: str,
    company_id: int | None = None,
    account_id: int | None = None,
    active_only: bool = False,
    u=Depends(current_user),
    db=Depends(get_db),
):
    require(u, "analytics")
    return f.report_details(db, u, start, end, metric, company_id, account_id, active_only)


@app.get("/api/stock")
def stock(company_id: int | None = None, u=Depends(current_user), db=Depends(get_db)):
    require(u, "inventory")
    if company_id is not None:
        f.get_company(db, u, company_id)
    today = business_today()
    month_start = today.replace(day=1)
    movement_query = scope_query(select(
        Movement.product_id,
        func.coalesce(func.sum(Movement.quantity), 0),
        func.coalesce(func.sum(Movement.cost), 0),
    ), Movement, u).where(Movement.date <= today).group_by(Movement.product_id)
    if company_id is not None:
        movement_query = movement_query.where(Movement.company_id == company_id)
    movement_totals = {
        pid: (int(quantity or 0), int(value or 0))
        for pid, quantity, value in db.execute(movement_query)
    }
    sales_query = scope_query(select(
        Sale.product_id, func.coalesce(func.sum(Sale.quantity), 0)
    ), Sale, u).where(Sale.date >= month_start, Sale.date <= today).group_by(Sale.product_id)
    if company_id is not None:
        sales_query = sales_query.where(Sale.company_id == company_id)
    sold_totals = {pid: int(quantity or 0) for pid, quantity in db.execute(sales_query)}
    returns_query = scope_query(select(
        Sale.product_id, func.coalesce(func.sum(Return.quantity), 0)
    ).join(Sale, Return.sale_id == Sale.id), Sale, u).where(
        Return.date >= month_start, Return.date <= today
    ).group_by(Sale.product_id)
    if company_id is not None:
        returns_query = returns_query.where(Sale.company_id == company_id, Return.company_id == company_id)
    for pid, quantity in db.execute(returns_query):
        sold_totals[pid] = sold_totals.get(pid, 0) - int(quantity or 0)
    products = db.scalars(scope_query(select(Product).where(Product.active.is_(True)), Product, u)).all()
    rows = []
    for product in products:
        specs = product.specifications or {}
        company_ids = {int(value) for value in (specs.get("_qor_company_ids") or []) if str(value).isdigit()}
        if company_id is not None and company_ids and company_id not in company_ids:
            continue
        if u.role != "Owner" and company_ids and not company_ids.intersection(u.company_ids):
            continue
        if getattr(product, "catalog_kind", "master") != "master":
            continue
        quantity, value = movement_totals.get(product.id, (0, 0))
        sold = sold_totals.get(product.id, 0)
        days = max(1, (today - month_start).days + 1)
        row = {
            "product_id": product.id, "sku": product.sku, "name": product.name,
            "quantity": quantity, "value": value, "min_stock": int(product.min_stock or 0),
            "sold": sold, "unit_cost": f.rounded(value, quantity) if quantity > 0 else 0,
            "days_of_stock": round(quantity / (sold / days), 1) if sold > 0 else None,
            "turnover_days": round(days * quantity / sold, 1) if sold > 0 else None,
            "velocity_per_day": round(sold / days, 3),
        }
        rows.append(row)
    if not allowed(u, "finance") and not allowed(u, "analytics"):
        for x in rows:
            x.pop("value", None)
    return rows


@app.post("/api/inventory/count")
def inventory_count(p: s.StockCountIn, tasks: BackgroundTasks, u=Depends(current_user), db=Depends(get_db)):
    """Reconcile counted physical quantity as one auditable, replay-safe movement."""
    require(u, "inventory", True)
    f.get_company(db, u, p.company_id)
    identity = f"stock-count:{p.company_id}:{p.idempotency_key}"
    body = p.model_dump(mode="json")
    old = db.scalar(select(CanonicalEvent).where(CanonicalEvent.identity == identity))
    if old:
        if old.event_type != "stock_count" or old.payload.get("request") != body:
            raise HTTPException(409, "Ключ операции уже использован для другой корректировки")
        db.commit()
        return {"ok": True, "replayed": True, "movement_id": old.target_id, **old.payload["result"]}
    product = f.lock_product(db, p.product_id)
    current = db.scalar(select(func.coalesce(func.sum(Batch.remaining), 0)).where(
        Batch.company_id == p.company_id, Batch.product_id == product.id
    )) or 0
    delta = p.counted_quantity - current
    movement_id = None
    if delta:
        unit_cost = p.unit_cost or int((product.specifications or {}).get("purchase_cost") or 0)
        if delta > 0 and unit_cost <= 0:
            raise HTTPException(422, "При увеличении остатка укажите подтверждённую себестоимость единицы")
        adjustment = s.AdjustIn(
            company_id=p.company_id, product_id=product.id, date=business_today(),
            quantity=delta, unit_cost=unit_cost if delta > 0 else 0,
            kind="adjustment", reason=p.reason,
            reference=f"count:{p.company_id}:{p.idempotency_key}",
        )
        posted = f.adjust(db, u, adjustment)
        movement_id = posted["movement_ids"][0] if posted.get("movement_ids") else None
    result = {"previous_quantity": current, "counted_quantity": p.counted_quantity, "delta": delta}
    db.add(CanonicalEvent(company_id=p.company_id, account_id=None, event_type="stock_count",
        identity=identity, payload={"request": body, "result": result}, target_id=movement_id, status="posted"))
    audit(db, u, "stock_count", "inventory", product.id, {**body, **result, "movement_id": movement_id})
    db.commit()
    _schedule_wb_sync_for_product(tasks, product.id, p.company_id)
    return {"ok": True, "replayed": False, "movement_id": movement_id, **result}


@app.get("/api/forecast")
def forecast(
    start: date,
    end: date,
    company_id: int | None = None,
    u=Depends(current_user),
    db=Depends(get_db),
):
    require(u, "analytics")
    return ai.forecast(f.report(db, u, start, end, company_id))


@app.get("/api/alerts")
def alerts(
    start: date, end: date, company_id: int | None = None,
    u=Depends(current_user), db=Depends(get_db),
):
    require(u, "analytics")
    if end < start:
        raise HTTPException(422, "Конец периода раньше начала")

    # Alert comparisons need only financial headline fields. The full report also
    # materializes all historical cash, sales, returns, movements and products;
    # building it twice here blocked the single web worker for every alert refresh.
    current = f.report_headline(db, u, start, end, company_id)
    days = (end - start).days + 1
    prev_end = start - timedelta(days=1)
    prev_start = prev_end - timedelta(days=days - 1)
    previous = f.report_headline(db, u, prev_start, prev_end, company_id)

    account_query = scope_query(select(Account), Account, u)
    if company_id is not None:
        account_query = account_query.where(Account.company_id == company_id)
    account_ids = {a.id for a in db.scalars(account_query)}

    # Recreate the report's inventory slice with grouped SQL instead of loading
    # every historical row into Python. This does not alter marketplace reports.
    stock_query = scope_query(
        select(
            Movement.product_id,
            func.coalesce(func.sum(Movement.quantity), 0).label("quantity"),
            func.coalesce(func.sum(Movement.cost), 0).label("value"),
        ),
        Movement,
        u,
    ).where(Movement.date <= end).group_by(Movement.product_id)
    if company_id is not None:
        stock_query = stock_query.where(Movement.company_id == company_id)
    stock_by_product = {
        pid: {"quantity": int(quantity or 0), "value": int(value or 0)}
        for pid, quantity, value in db.execute(stock_query)
    }

    sold_by_product = {}
    if account_ids:
        sold_query = scope_query(
            select(Sale.product_id, func.coalesce(func.sum(Sale.quantity), 0)),
            Sale,
            u,
        ).where(
            Sale.account_id.in_(account_ids), Sale.date >= start, Sale.date <= end
        )
        if company_id is not None:
            sold_query = sold_query.where(Sale.company_id == company_id)
        for pid, quantity in db.execute(sold_query.group_by(Sale.product_id)):
            sold_by_product[pid] = int(quantity or 0)

        returned_query = scope_query(
            select(Sale.product_id, func.coalesce(func.sum(Return.quantity), 0))
            .join(Sale, Return.sale_id == Sale.id),
            Return,
            u,
        ).where(
            Sale.account_id.in_(account_ids),
            Sale.date <= end,
            Return.date >= start,
            Return.date <= end,
        )
        if company_id is not None:
            returned_query = returned_query.where(
                Return.company_id == company_id, Sale.company_id == company_id
            )
        for pid, quantity in db.execute(returned_query.group_by(Sale.product_id)):
            sold_by_product[pid] = sold_by_product.get(pid, 0) - int(quantity or 0)

    stockrows = []
    products_query = scope_query(select(Product).where(Product.active.is_(True)), Product, u)
    for product in db.scalars(products_query):
        specs = product.specifications or {}
        company_ids = {int(value) for value in (specs.get("_qor_company_ids") or []) if str(value).isdigit()}
        if company_id is not None and company_ids and company_id not in company_ids:
            continue
        if u.role != "Owner" and company_ids and not company_ids.intersection(u.company_ids):
            continue
        if getattr(product, "catalog_kind", "master") != "master":
            continue
        minimum = int(product.min_stock or 0)
        # A zero threshold means the owner has not configured a low-stock alert.
        # Avoid treating every imported marketplace product with stock 0 as a warning.
        if minimum <= 0:
            continue
        stock = stock_by_product.get(product.id, {"quantity": 0, "value": 0})
        sold = sold_by_product.get(product.id, 0)
        quantity = stock["quantity"]
        stockrows.append({
            "product_id": product.id,
            "sku": product.sku,
            "name": product.name,
            "quantity": quantity,
            "value": stock["value"],
            "min_stock": minimum,
            "sold": sold,
            "unit_cost": f.rounded(stock["value"], quantity) if quantity > 0 else 0,
            "days_of_stock": round(quantity / (sold / max(1, days)), 1) if sold > 0 else None,
            "turnover_days": round(max(1, days) * quantity / sold, 1) if sold > 0 else None,
            "velocity_per_day": round(sold / max(1, days), 3),
        })
    current["stock"] = stockrows
    current["stock_alerts"] = [row for row in stockrows if row["quantity"] <= row["min_stock"]]

    # Preserve the receivables rows previously included in the alert response,
    # but aggregate returns and payments in SQL rather than materializing history.
    receivables = []
    if account_ids:
        sales_query = scope_query(
            select(
                Sale.id, Sale.reference, Sale.customer_id, Sale.account_id,
                Sale.revenue, Sale.due_date, Account.channel,
            ).join(Account, Account.id == Sale.account_id),
            Sale,
            u,
        ).where(Sale.account_id.in_(account_ids), Sale.date <= end)
        if company_id is not None:
            sales_query = sales_query.where(Sale.company_id == company_id)
        all_sales = sales_query.subquery()

        returns_query = scope_query(
            select(Return.sale_id, func.coalesce(func.sum(Return.revenue), 0).label("returned"))
            .join(Sale, Return.sale_id == Sale.id),
            Return,
            u,
        ).where(Sale.account_id.in_(account_ids), Return.date <= end)
        if company_id is not None:
            returns_query = returns_query.where(
                Return.company_id == company_id, Sale.company_id == company_id
            )
        return_totals = returns_query.group_by(Return.sale_id).subquery()

        paid_query = scope_query(
            select(CashEntry.sale_id, func.coalesce(func.sum(CashEntry.amount), 0).label("paid")),
            CashEntry,
            u,
        ).where(CashEntry.date <= end, CashEntry.sale_id.is_not(None))
        if company_id is not None:
            paid_query = paid_query.where(CashEntry.company_id == company_id)
        paid_totals = paid_query.group_by(CashEntry.sale_id).subquery()

        receivable_rows = db.execute(
            select(
                all_sales.c.id, all_sales.c.reference, all_sales.c.customer_id,
                all_sales.c.channel, all_sales.c.revenue, all_sales.c.due_date,
                func.coalesce(return_totals.c.returned, 0),
                func.coalesce(paid_totals.c.paid, 0),
            )
            .select_from(all_sales)
            .outerjoin(return_totals, return_totals.c.sale_id == all_sales.c.id)
            .outerjoin(paid_totals, paid_totals.c.sale_id == all_sales.c.id)
        )
        for sale_id, reference, customer_id, channel, revenue, due_date, returned, paid in receivable_rows:
            due = max(0, int(revenue) - int(returned or 0) - int(paid or 0))
            if due:
                receivables.append({
                    "sale_id": sale_id, "reference": reference, "customer_id": customer_id,
                    "channel": channel, "due": due, "due_date": due_date,
                    "overdue": bool(due_date and due_date < end),
                })

    items = []
    for x in current["stock"]:
        dl = x.get("days_of_stock")
        if dl is not None and dl <= 7:
            items.append({"severity":"critical", "type":"stockout", "title":f"{x['name']}: запас примерно на {dl} дн.", "product_id":x.get("product_id")})
        if x.get("stale"):
            items.append({"severity":"warning", "type":"stale_stock", "title":f"{x['name']}: товар залежался на складе", "product_id":x.get("product_id")})
    if current["revenue"] and current["margin"] is not None and current["margin"] < 0:
        items.append({"severity":"critical", "type":"negative_margin", "title":"За период чистая маржа отрицательная"})
    def pct(cur, prev):
        return ((cur-prev)/abs(prev)*100) if prev else None
    ad = pct(current["advertising"], previous["advertising"]); mg = pct(current["margin"] or 0, previous["margin"] or 0)
    if ad is not None and ad >= 30:
        items.append({"severity":"warning", "type":"advertising_growth", "title":f"Расходы на рекламу выросли на {ad:.1f}%", "change_percent":round(ad,1)})
    if mg is not None and mg <= -10:
        items.append({"severity":"warning", "type":"margin_drop", "title":f"Маржа снизилась на {abs(mg):.1f}%", "change_percent":round(mg,1)})
    prev_channels={x["channel"]:x for x in previous["channel_summary"]}
    for x in current["channel_summary"]:
        old=prev_channels.get(x["channel"],{}).get("revenue",0); ch=pct(x["revenue"],old)
        if ch is not None and ch <= -20:
            items.append({"severity":"warning", "type":"channel_drop", "title":f"{x['channel']}: оборот снизился на {abs(ch):.1f}%", "change_percent":round(ch,1)})
    if current["stock_value"] >= 50_000_000 * 100:
        items.append({"severity":"info", "type":"stock_value", "title":"На складе находится товар на сумму 50 млн ₸ или более"})
    if current.get("pending_cogs_sales", 0):
        count = int(current["pending_cogs_sales"])
        items.append({"severity":"warning", "type":"missing_cost", "title":f"Для {count} продаж не задана себестоимость",
            "count":count,"action":"Заполните себестоимость или импортируйте закупки; прибыль за период пока неполная."})
    account_query=select(Account).where(Account.active.is_(True),Account.channel.in_(["Kaspi","Wildberries","Ozon"]))
    if u.role != "Owner": account_query=account_query.where(Account.company_id.in_(u.company_ids or []))
    if company_id is not None: account_query=account_query.where(Account.company_id==company_id)
    active_accounts=db.scalars(account_query).all()
    active_ids=[a.id for a in active_accounts]
    account_by_id={a.id:a for a in active_accounts}
    if active_ids:
        pending_mapping_rows=db.execute(select(MarketplaceRecord.account_id,func.count(MarketplaceRecord.id))
            .where(MarketplaceRecord.account_id.in_(active_ids),MarketplaceRecord.stage_status=="pending_mapping")
            .group_by(MarketplaceRecord.account_id)).all()
        for account_id,count in pending_mapping_rows:
            account=account_by_id[account_id]
            items.append({"severity":"warning","type":"unmapped_products",
                "title":f"{account.name}: {count} операций не сопоставлены с товарами ERP",
                "account_id":account.id,"channel":account.channel,"count":int(count),
                "action":"Сопоставьте SKU / артикул маркетплейса с товаром ERP; склад и прибыль не должны угадывать товар."})
        for account in active_accounts:
            if account.sync_status in {"failed","partial"} or account.sync_error:
                items.append({"severity":"critical" if account.sync_status=="failed" else "warning",
                    "type":"marketplace_sync","title":f"{account.name}: синхронизация {account.sync_status}",
                    "account_id":account.id,"channel":account.channel,
                    "action":(account.sync_error or "Откройте интеграцию и проверьте состояние подключения.")[:300]})
    return {"items":items, "count":len(items), "stock_alerts":current.get("stock_alerts", []),
            "receivables":receivables, "period":[str(start),str(end)], "comparison_period":[str(prev_start),str(prev_end)],
            "method":"Детерминированные правила ERP; без LLM. Рыночные сигналы добавляются только при наличии подтверждённого источника."}


@app.post("/api/calculator")
def calculator(p: dict, u=Depends(current_user)):
    require(u, "marketing")
    return ai.calculator(p)


@app.get("/api/roles")
def roles(u=Depends(current_user)):
    require(u, "users")
    return {k: sorted(v) for k, v in PERMISSIONS.items()}


@app.get("/api/ai/agents")
def ai_agents(u=Depends(current_user)):
    require(u, "ai")
    return {"items": ai.agent_catalog()}


@app.get("/api/settings")
def settings(u=Depends(current_user), db=Depends(get_db)):
    obj = db.get(Setting, 1)
    return {
        **dump(obj),
        "ai_enabled": bool(os.getenv("ANTHROPIC_API_KEY")),
        "claude_skills": ai.claude_skills_status(),
        "demo": os.getenv("DEMO_MODE", "false").strip().lower() == "true",
    }


@app.post("/api/ai/skills/install")
def install_claude_skills(u=Depends(current_user), db=Depends(get_db)):
    """Upload the bundled QOR Skills to the Anthropic workspace behind the server key."""
    require(u, "settings", True)
    if not CLAUDE_SKILL_INSTALL_LOCK.acquire(blocking=False):
        raise HTTPException(409, "Загрузка Skills уже выполняется")
    try:
        if ai.claude_skills_status()["enabled"]:
            raise HTTPException(409, "Все 11 Skills уже подключены")
        skills = ai.install_claude_skills()
        audit(db, u, "create", "claude_skills", 1)
        db.commit()
        return {"skills": skills, "configured": len(skills), "required": len(ai.CLAUDE_SKILL_KEYS)}
    finally:
        CLAUDE_SKILL_INSTALL_LOCK.release()


@app.put("/api/settings")
def save_settings(p: dict, u=Depends(current_user), db=Depends(get_db)):
    require(u, "settings", True)
    if set(p) - {"business_name"}:
        raise HTTPException(
            422, "В этой версии валюта KZT, зона Asia/Almaty фиксированы"
        )
    obj = db.get(Setting, 1)
    obj.business_name = str(p.get("business_name", "QOR Business"))[:160]
    audit(db, u, "update", "settings", 1)
    db.commit()
    return dump(obj)


def _credential_cipher():
    # Production credentials are encrypted with a dedicated system secret, never
    # with an account password or database URL. Development retains a local-only
    # fallback so a fresh demo can run without copying a secret into source code.
    material = os.getenv("CREDENTIAL_ENCRYPTION_KEY")
    if not material:
        if os.getenv("APP_ENV", "development").lower() == "production":
            raise HTTPException(503, "Сервер интеграций не настроен: отсутствует CREDENTIAL_ENCRYPTION_KEY")
        material = os.getenv("OWNER_PASSWORD") or os.getenv("DATABASE_URL", "qor-local-dev")
    key = base64.urlsafe_b64encode(hashlib.sha256(material.encode()).digest())
    return Fernet(key)


@app.put("/api/accounts/{id}/credentials")
def save_account_credentials(id: int, p: dict, u=Depends(current_user), db=Depends(get_db)):
    require(u, "accounts", True)
    acc = db.get(Account, id)
    if not acc: raise HTTPException(404, "Аккаунт не найден")
    scope(u, acc.company_id)
    allowed = {"Kaspi": {"token"}, "Wildberries": {"token"}, "Ozon": {"client_id", "api_key"}}.get(acc.channel, set())
    if not allowed: raise HTTPException(422, "Для этого канала API-ключи не требуются")
    values = {k: str(p.get(k, "")).strip() for k in allowed}
    if any(not values[k] for k in allowed): raise HTTPException(422, "Заполните все данные подключения")
    # Store canonical per-account credential fields. Marketplace connections created in QOR
    # never depend on Render ENV variable names or Account.credential_ref.
    names = values
    enc = _credential_cipher().encrypt(json.dumps(names).encode()).decode()
    row = db.scalar(select(IntegrationSecret).where(IntegrationSecret.account_id == id))
    if row: row.encrypted_payload, row.key_hint = enc, "••••" + list(values.values())[0][-4:]
    else: db.add(IntegrationSecret(account_id=id, encrypted_payload=enc, key_hint="••••" + list(values.values())[0][-4:]))
    acc.mode = "api"
    acc.active = True
    audit(db, u, "credentials_update", "accounts", id, {"channel": acc.channel})
    db.commit()
    return {"ok": True, "channel": acc.channel, "configured": True, "hint": "••••" + list(values.values())[0][-4:]}

@app.delete("/api/accounts/{id}/credentials")
def delete_account_credentials(id: int, u=Depends(current_user), db=Depends(get_db)):
    require(u, "accounts", True)
    acc=db.get(Account,id)
    if not acc: raise HTTPException(404,"Аккаунт не найден")
    scope(u,acc.company_id)
    row=db.scalar(select(IntegrationSecret).where(IntegrationSecret.account_id==id))
    if row: db.delete(row)
    acc.active=False; acc.sync_status="never"; acc.sync_error=""
    audit(db,u,"credentials_delete","accounts",id,{"channel":acc.channel}); db.commit()
    return {"ok":True,"archived":True}


@app.get("/api/accounts/{id}/marketplace-settings")
def marketplace_settings_get(id: int, u=Depends(current_user), db=Depends(get_db)):
    require(u, "accounts")
    acc = db.get(Account, id)
    if not acc or acc.channel not in {"Kaspi", "Wildberries"}:
        raise HTTPException(404, "Интеграция не найдена")
    scope(u, acc.company_id)
    settings = acc.integration_settings or {}
    return {"channel": acc.channel, "auto_sync": bool(settings.get("auto_sync", False)),
            "kaspi_store_id": settings.get("kaspi_store_id", ""),
            "kaspi_merchant_id": settings.get("kaspi_merchant_id", ""),
            "kaspi_full_catalog_confirmed": bool(settings.get("kaspi_full_catalog_confirmed", False)),
            "kaspi_feed_configured": bool(settings.get("kaspi_feed_token_hash")),
            "wb_warehouse_id": str(settings.get("wb_warehouse_id", ""))}


@app.put("/api/accounts/{id}/marketplace-settings")
def marketplace_settings_save(id: int, p: dict, u=Depends(current_user), db=Depends(get_db)):
    require(u, "accounts", True)
    acc = db.get(Account, id)
    if not acc or acc.channel not in {"Kaspi", "Wildberries"}:
        raise HTTPException(404, "Интеграция не найдена")
    scope(u, acc.company_id)
    settings = dict(acc.integration_settings or {})
    if acc.channel == "Kaspi":
        allowed = {"kaspi_store_id", "kaspi_merchant_id", "kaspi_full_catalog_confirmed", "auto_sync"}
        if set(p) - allowed:
            raise HTTPException(422, "Неизвестная настройка Kaspi")
        store = str(p.get("kaspi_store_id", settings.get("kaspi_store_id", ""))).strip()
        merchant = str(p.get("kaspi_merchant_id", settings.get("kaspi_merchant_id", ""))).strip()
        if len(store) > 120 or len(merchant) > 120:
            raise HTTPException(422, "Код магазина или продавца слишком длинный")
        settings.update({"kaspi_store_id": store, "kaspi_merchant_id": merchant,
                         "kaspi_full_catalog_confirmed": bool(p.get("kaspi_full_catalog_confirmed", settings.get("kaspi_full_catalog_confirmed", False)))})
    else:
        allowed = {"wb_warehouse_id", "auto_sync"}
        if set(p) - allowed:
            raise HTTPException(422, "Неизвестная настройка Wildberries")
        warehouse_id = str(p.get("wb_warehouse_id", settings.get("wb_warehouse_id", ""))).strip()[:40]
        if warehouse_id and not warehouse_id.isdigit():
            raise HTTPException(422, "ID склада WB должен быть числовым")
        settings["wb_warehouse_id"] = warehouse_id
    if "auto_sync" in p:
        settings["auto_sync"] = bool(p["auto_sync"])
    acc.integration_settings = settings
    audit(db, u, "marketplace_settings_update", "accounts", acc.id,
          {"channel": acc.channel, "keys": sorted(set(p) - {"auto_sync"})})
    db.commit()
    return {"ok": True, "channel": acc.channel, "auto_sync": bool(settings.get("auto_sync", False))}


@app.post("/api/accounts/{id}/kaspi/feed-url")
def kaspi_feed_url(id: int, request: Request, u=Depends(current_user), db=Depends(get_db)):
    require(u, "accounts", True)
    acc = db.get(Account, id)
    if not acc or acc.channel != "Kaspi" or not acc.active:
        raise HTTPException(404, "Активный магазин Kaspi не найден")
    scope(u, acc.company_id)
    settings = dict(acc.integration_settings or {})
    if not settings.get("kaspi_store_id") or not settings.get("kaspi_merchant_id"):
        raise HTTPException(422, "Сначала укажите код склада и merchant ID из кабинета Kaspi")
    if not settings.get("kaspi_full_catalog_confirmed"):
        raise HTTPException(409, "Подтвердите, что все товары действующего прайс-листа Kaspi будут связаны с QOR")
    token = secrets.token_urlsafe(32)
    settings["kaspi_feed_token_hash"] = hashlib.sha256(token.encode()).hexdigest()
    acc.integration_settings = settings
    audit(db, u, "kaspi_feed_token_rotated", "accounts", acc.id, {"channel": "Kaspi"})
    db.commit()
    base = str(request.base_url).rstrip("/")
    return {"feed_url": f"{base}/public/kaspi-feeds/{acc.id}/{token}.xml",
            "setup_note": "Добавьте эту ссылку в кабинете Kaspi → Товары → Прайс-лист → Автоматическая загрузка. Kaspi проверяет обновления примерно раз в час. Ссылка показывает только сопоставленные товары этого магазина."}


@app.get("/public/kaspi-feeds/{account_id}/{token}.xml")
def kaspi_public_feed(account_id: int, token: str, db=Depends(get_db)):
    acc = db.get(Account, account_id)
    if not acc or not acc.active or acc.channel != "Kaspi":
        raise HTTPException(404, "Прайс-лист не найден")
    settings = acc.integration_settings or {}
    expected = str(settings.get("kaspi_feed_token_hash") or "")
    actual = hashlib.sha256(token.encode()).hexdigest()
    if not expected or not hmac.compare_digest(expected, actual):
        raise HTTPException(404, "Прайс-лист не найден")
    if not settings.get("kaspi_full_catalog_confirmed") or not settings.get("kaspi_store_id") or not settings.get("kaspi_merchant_id"):
        raise HTTPException(409, "Настройки прайс-листа не завершены")
    mappings = db.scalars(select(MarketplaceProductMapping).where(
        MarketplaceProductMapping.account_id == acc.id,
        MarketplaceProductMapping.company_id == acc.company_id,
        MarketplaceProductMapping.channel == "Kaspi",
    )).all()
    product_ids = sorted({m.product_id for m in mappings})
    products = {p.id: p for p in db.scalars(select(Product).where(Product.id.in_(product_ids), Product.active.is_(True))).all()} if product_ids else {}
    skus_by_product = {}
    for mapping in mappings:
        key = mapping.external_key
        if key.startswith("sku:"):
            sku = key[4:]
            if re.fullmatch(r"[A-Za-z0-9._-]{1,80}", sku):
                skus_by_product.setdefault(mapping.product_id, set()).add(sku)
    root = ET.Element("kaspi_catalog", {"date": business_today().isoformat(), "xmlns": "kaspiShopping",
        "xmlns:xsi": "http://www.w3.org/2001/XMLSchema-instance",
        "xsi:schemaLocation": "kaspiShopping http://kaspi.kz/kaspishopping.xsd"})
    ET.SubElement(root, "company").text = acc.name
    ET.SubElement(root, "merchantid").text = str(settings["kaspi_merchant_id"])
    offers = ET.SubElement(root, "offers")
    physical_by_product = dict(db.execute(select(Batch.product_id, func.coalesce(func.sum(Batch.remaining), 0))
        .where(Batch.company_id == acc.company_id, Batch.product_id.in_(list(products)))
        .group_by(Batch.product_id)).all()) if products else {}
    allocated_elsewhere = {}
    for product_id, product in products.items():
        specs = product.specifications or {}
        allocation_map = specs.get("marketplace_stock_allocations") or {}
        for other_id, value in allocation_map.items():
            try:
                count = int(value.get("quantity", 0)) if isinstance(value, dict) else int(value or 0)
            except (TypeError, ValueError):
                count = 0
            if isinstance(value, dict):
                for size_amount in (value.get("sizes") or {}).values():
                    try: count += max(0, int(size_amount or 0))
                    except (TypeError, ValueError): pass
            allocated_elsewhere[product_id] = allocated_elsewhere.get(product_id, 0) + max(0, count)
    for product_id, product in products.items():
        skus = sorted(skus_by_product.get(product_id, set()))
        if not skus or product.price <= 0:
            continue
        specs = product.specifications or {}
        allocation_map = specs.get("marketplace_stock_allocations") or {}
        try:
            qty_value = allocation_map.get(str(acc.id), 0)
            qty = int(qty_value.get("quantity", 0)) if isinstance(qty_value, dict) else int(qty_value or 0)
        except (TypeError, ValueError):
            qty = 0
        qty = max(0, qty)
        physical = max(0, int(physical_by_product.get(product_id, 0) or 0) - int(specs.get("reserved_quantity") or 0))
        # Do not publish a platform allocation that exceeds the verified shared physical pool.
        if qty > physical or allocated_elsewhere.get(product_id, 0) > physical:
            qty = 0
        for sku in skus:
            offer = ET.SubElement(offers, "offer", {"sku": sku})
            ET.SubElement(offer, "model").text = product.name
            ET.SubElement(offer, "brand").text = product.brand or "Без бренда"
            availabilities = ET.SubElement(offer, "availabilities")
            ET.SubElement(availabilities, "availability", {
                "available": "yes" if qty > 0 else "no", "storeId": str(settings["kaspi_store_id"]), "stockCount": str(qty)})
            ET.SubElement(offer, "price").text = str((int(product.price) + 50) // 100)
    xml = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    return Response(content=xml, media_type="application/xml", headers={"Cache-Control": "no-store"})


def _wb_mapped_ids(db, account_id, product_id):
    rows = db.scalars(select(MarketplaceProductMapping).where(
        MarketplaceProductMapping.account_id == account_id,
        MarketplaceProductMapping.product_id == product_id,
        MarketplaceProductMapping.channel == "Wildberries",
    )).all()
    nm_ids, chrt_ids = set(), set()
    for row in rows:
        key = row.external_key
        if key.startswith("product:"):
            parts = key.split(":")
            if len(parts) >= 2 and parts[1].isdigit(): nm_ids.add(int(parts[1]))
            if len(parts) >= 3 and parts[2].isdigit(): chrt_ids.add(int(parts[2]))
        elif key.startswith("size:") and key[5:].isdigit():
            chrt_ids.add(int(key[5:]))
    return rows, sorted(nm_ids), sorted(chrt_ids)


def sync_wb_product(db, account, product):
    settings = account.integration_settings or {}
    warehouse_id = str(settings.get("wb_warehouse_id") or "")
    if not warehouse_id:
        raise HTTPException(409, "Сначала выберите склад продавца Wildberries в настройках интеграции")
    if not account.active or account.mode != "api":
        raise HTTPException(409, "Интеграция Wildberries выключена")
    if account.settlement_currency != product.currency:
        raise HTTPException(409, f"Валюта магазина WB ({account.settlement_currency}) не совпадает с валютой товара ({product.currency})")
    mappings, nm_ids, chrt_ids = _wb_mapped_ids(db, account.id, product.id)
    if not mappings or not nm_ids or not chrt_ids:
        raise HTTPException(409, "Товар не сопоставлен с карточкой и размерами Wildberries")
    allocation_map = (product.specifications or {}).get("marketplace_stock_allocations") or {}
    own = allocation_map.get(str(account.id)) or {}
    sizes = own.get("sizes") or {} if isinstance(own, dict) else {}
    if not sizes:
        raise HTTPException(409, "Укажите отдельное распределение остатка по размерам WB в карточке товара")
    physical = int(db.scalar(select(func.coalesce(func.sum(Batch.remaining), 0)).where(
        Batch.company_id == account.company_id, Batch.product_id == product.id)) or 0)
    reserved = int((product.specifications or {}).get("reserved_quantity") or 0)
    usable = max(0, physical - reserved)
    all_allocations = 0
    for account_alloc in allocation_map.values():
        if not isinstance(account_alloc, dict): continue
        try: all_allocations += max(0, int(account_alloc.get("quantity", 0) or 0))
        except (TypeError, ValueError): pass
        for amount in (account_alloc.get("sizes") or {}).values():
            try: all_allocations += max(0, int(amount or 0))
            except (TypeError, ValueError): pass
    if all_allocations > usable:
        raise HTTPException(409, f"Распределено {all_allocations} шт., доступно на складе {usable}. Исправьте лимиты площадок")
    stock_rows = []
    for chrt_id in chrt_ids:
        try: amount = int(sizes.get(str(chrt_id), 0) or 0)
        except (TypeError, ValueError): raise HTTPException(422, "Остаток WB должен быть целым неотрицательным числом")
        if amount < 0: raise HTTPException(422, "Остаток WB не может быть отрицательным")
        stock_rows.append({"chrtId": chrt_id, "amount": amount})
    if sum(row["amount"] for row in stock_rows) > usable:
        raise HTTPException(409, "Выделение WB превышает доступный физический остаток")
    try:
        client = connector(account, load_saved_credentials(db, account.id))
        price_rows = []
        if int(product.price or 0) <= 0:
            raise HTTPException(409, "Перед синхронизацией укажите цену товара больше нуля")
        for nm_id in nm_ids:
            payload = client.product_prices(nm_id)
            data = payload.get("data") or payload
            goods = data.get("listGoods") or data.get("goods") or [] if isinstance(data, dict) else []
            good = next((g for g in goods if int(g.get("nmID") or g.get("nmId") or 0) == nm_id), None)
            if not good:
                raise HTTPException(409, f"WB не вернул текущие скидки nmID {nm_id}; цена не отправлена")
            discounts = {int(size.get("discount") or 0) for size in (good.get("sizes") or [])}
            if len(discounts) > 1:
                raise HTTPException(409, f"У nmID {nm_id} разные скидки по размерам. Сначала выровняйте их в кабинете WB")
            discount = next(iter(discounts), 0)
            price_rows.append({"nmID": nm_id, "price": (int(product.price) + 50) // 100, "discount": discount})
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, "Wildberries: " + str(exc)[:500])
    price_result = stock_result = None
    sync_errors = {}
    try:
        price_result = client.upload_prices(price_rows)
    except Exception as exc:
        sync_errors["price"] = str(exc)[:400]
    try:
        stock_result = client.update_seller_stocks(warehouse_id, stock_rows)
    except Exception as exc:
        sync_errors["stock"] = str(exc)[:400]
    audit(db, None, "wildberries_product_outbound_sync", "accounts", account.id,
          {"product_id": product.id, "nm_ids": nm_ids, "sizes": len(stock_rows), "stock_total": sum(x["amount"] for x in stock_rows),
           "price_submitted": price_result is not None, "stock_updated": stock_result is not None, "errors": sync_errors})
    db.commit()
    if sync_errors:
        raise HTTPException(422, "WB: " + "; ".join(f"{key}: {value}" for key, value in sync_errors.items()))
    return {"ok": True, "channel": "Wildberries", "product_id": product.id,
            "price_task": price_result, "stock_result": stock_result,
            "price_kzt": (int(product.price) + 50) // 100,
            "stock_total": sum(x["amount"] for x in stock_rows)}


def _auto_sync_wb_product(account_id, product_id):
    """Best-effort background sync after ERP commits; failures stay in service logs."""
    with SessionLocal() as db:
        account = db.get(Account, account_id)
        product = db.get(Product, product_id)
        if not account or not product or not (account.integration_settings or {}).get("auto_sync"):
            return
        try:
            sync_wb_product(db, account, product)
        except Exception:
            db.rollback()
            logger.exception("Automatic Wildberries product sync failed for account=%s product=%s", account_id, product_id)


def _schedule_wb_sync_for_product(tasks, product_id, company_id=None):
    if tasks is None: return
    with SessionLocal() as db:
        query = select(Account.id).join(
            MarketplaceProductMapping, MarketplaceProductMapping.account_id == Account.id
        ).where(
            MarketplaceProductMapping.product_id == product_id,
            Account.channel == "Wildberries", Account.active.is_(True)
        )
        if company_id is not None: query = query.where(Account.company_id == company_id)
        account_ids = set(db.scalars(query.distinct()).all())
        if not account_ids: return
        accounts = db.scalars(select(Account).where(Account.id.in_(account_ids))).all()
        eligible = [a.id for a in accounts if (a.integration_settings or {}).get("auto_sync")]
    for account_id in eligible:
        tasks.add_task(_auto_sync_wb_product, account_id, product_id)

@app.get("/api/accounts/{id}/wb/warehouses")
def wb_seller_warehouses(id: int, u=Depends(current_user), db=Depends(get_db)):
    require(u, "accounts")
    account = db.get(Account, id)
    if not account or account.channel != "Wildberries": raise HTTPException(404, "Магазин Wildberries не найден")
    scope(u, account.company_id)
    try:
        client = connector(account, load_saved_credentials(db, account.id))
        result = client.list_seller_warehouses()
        return {"items": result if isinstance(result, list) else result.get("warehouses", [])}
    except Exception as exc:
        raise HTTPException(422, "Не удалось получить склады продавца WB: " + str(exc)[:400])


@app.post("/api/accounts/{id}/wb/products/{product_id}/sync")
def wb_product_sync(id: int, product_id: int, p: dict, u=Depends(current_user), db=Depends(get_db)):
    require(u, "accounts", True)
    account = db.get(Account, id)
    product = db.get(Product, product_id)
    if not account or account.channel != "Wildberries" or not product: raise HTTPException(404, "Магазин или товар не найден")
    scope(u, account.company_id)
    if p.get("confirm") is not True: raise HTTPException(428, "Подтвердите отправку текущей цены и распределённых остатков в Wildberries")
    return sync_wb_product(db, account, product)

@app.post("/api/accounts/{id}/kaspi/orders/{order_id}/status")
def kaspi_order_status(id: int, order_id: str, p: dict, u=Depends(current_user), db=Depends(get_db)):
    require(u, "accounts", True)
    acc=db.get(Account,id)
    if not acc or acc.channel != "Kaspi": raise HTTPException(404,"Kaspi аккаунт не найден")
    scope(u,acc.company_id)
    creds=load_saved_credentials(db,acc.id)
    try:
        result=connector(acc,creds).update_order(order_id=order_id,status=p.get("status"),code=p.get("code"),cancellation_reason=p.get("cancellation_reason"),cancellation_comment=p.get("cancellation_comment"),number_of_space=p.get("number_of_space"),security_code=p.get("security_code"),send_code=bool(p.get("send_code",False)))
    except Exception as exc:
        raise HTTPException(422,str(exc)[:1000])
    audit(db,u,"kaspi_order_status","accounts",id,{"order_id":order_id,"status":p.get("status")}); db.commit()
    return {"ok":True,"result":result}


def _kaspi_api_context(account_id, user, db, write=False):
    require(user, "accounts", write)
    account = db.get(Account, account_id)
    if not account or account.channel != "Kaspi":
        raise HTTPException(404, "Kaspi аккаунт не найден")
    scope(user, account.company_id)
    if account.mode != "api":
        raise HTTPException(409, "Для операции включите режим API у Kaspi-аккаунта")
    try:
        credentials = load_saved_credentials(db, account.id)
    except AccessRequired as exc:
        raise HTTPException(409, str(exc))
    return account, connector(account, credentials)


@app.get("/api/accounts/{id}/kaspi/orders/{order_id}")
def kaspi_get_order(id: int, order_id: str, u=Depends(current_user), db=Depends(get_db)):
    _, client = _kaspi_api_context(id, u, db)
    try:
        return client.get_order(order_id)
    except Exception as exc:
        raise HTTPException(422, str(exc)[:500])


@app.get("/api/accounts/{id}/kaspi/orders/{order_id}/entries")
def kaspi_get_order_entries(id: int, order_id: str, page_size: int = Query(100, ge=1, le=100),
                            u=Depends(current_user), db=Depends(get_db)):
    _, client = _kaspi_api_context(id, u, db)
    try:
        return {"data": client.get_order_entries(order_id, page_size)}
    except Exception as exc:
        raise HTTPException(422, str(exc)[:500])


@app.get("/api/accounts/{id}/kaspi/catalog/categories")
def kaspi_catalog_categories(id: int, u=Depends(current_user), db=Depends(get_db)):
    _, client = _kaspi_api_context(id, u, db)
    try:
        return client.list_categories()
    except Exception as exc:
        raise HTTPException(422, str(exc)[:500])


@app.get("/api/accounts/{id}/kaspi/catalog/attributes")
def kaspi_catalog_attributes(id: int, category: str = Query(..., min_length=1, max_length=240),
                             u=Depends(current_user), db=Depends(get_db)):
    _, client = _kaspi_api_context(id, u, db)
    try:
        return client.list_category_attributes(category)
    except Exception as exc:
        raise HTTPException(422, str(exc)[:500])


@app.get("/api/accounts/{id}/kaspi/catalog/attribute-values")
def kaspi_catalog_attribute_values(id: int, category: str = Query(..., min_length=1, max_length=240),
                                  attribute: str = Query(..., min_length=1, max_length=240),
                                  u=Depends(current_user), db=Depends(get_db)):
    _, client = _kaspi_api_context(id, u, db)
    try:
        return client.list_attribute_values(category, attribute)
    except Exception as exc:
        raise HTTPException(422, str(exc)[:500])


@app.post("/api/accounts/{id}/kaspi/products/import")
def kaspi_import_products(id: int, payload: dict, u=Depends(current_user), db=Depends(get_db)):
    account, client = _kaspi_api_context(id, u, db, write=True)
    products = payload.get("products")
    try:
        result = client.import_products(products)
    except Exception as exc:
        raise HTTPException(422, str(exc)[:500])
    audit(db, u, "kaspi_product_import", "accounts", id, {"count": len(products or [])})
    db.commit()
    return {"ok": True, "submitted": len(products or []), "result": result}


@app.get("/api/accounts/{id}/kaspi/products/import/{import_code}")
def kaspi_product_import_status(id: int, import_code: str, u=Depends(current_user), db=Depends(get_db)):
    _, client = _kaspi_api_context(id, u, db)
    try:
        return client.get_product_import_status(import_code)
    except Exception as exc:
        raise HTTPException(422, str(exc)[:500])


@app.get("/api/accounts/{id}/kaspi/products/import/{import_code}/result")
def kaspi_product_import_result(id: int, import_code: str, u=Depends(current_user), db=Depends(get_db)):
    _, client = _kaspi_api_context(id, u, db)
    try:
        return client.get_product_import_result(import_code)
    except Exception as exc:
        raise HTTPException(422, str(exc)[:500])


def _stage_marketplace_records(db, acc, result):
    """Persist each external record once. Never invent a financial sale from an order payload."""
    inserted = 0
    updated = 0
    for item in result.records:
        external_id = str(item.get("external_id") or "").strip()
        if not external_id:
            continue
        row = db.scalar(select(MarketplaceRecord).where(
            MarketplaceRecord.account_id == acc.id,
            MarketplaceRecord.source == result.source,
            MarketplaceRecord.external_id == external_id,
        ))
        payload = jsonable_encoder(item)
        if row:
            old_payload = {k:v for k,v in (row.payload or {}).items() if not k.startswith('_qor_')}
            # Kaspi list responses can omit event fields that were present in a
            # code lookup. Keep established source facts when a later response
            # sends null; never synthesize one event date from another.
            merged = dict(old_payload)
            merged.update(payload)
            for field in (
                "completion_date", "completed_date", "returned_date", "return_date",
                "creation_date", "order_creation_date", "entries_complete", "entries_count",
                "quantity", "total_price", "sku", "external_product_id", "product_name",
            ):
                if payload.get(field) is None and old_payload.get(field) is not None:
                    merged[field] = old_payload[field]
            old_raw = old_payload.get("raw_attributes")
            new_raw = payload.get("raw_attributes")
            if isinstance(old_raw, dict) and isinstance(new_raw, dict):
                raw = dict(old_raw)
                raw.update({k: v for k, v in new_raw.items() if v is not None})
                merged["raw_attributes"] = raw
            if old_payload != merged:
                row.payload = merged
                # Revalidate revisions through canonical identity checks, never
                # leave a changed source silently marked as projected.
                row.status = 'received'
            updated += 1
        else:
            db.add(MarketplaceRecord(account_id=acc.id, company_id=acc.company_id, source=result.source, external_id=external_id, payload=payload))
            inserted += 1
    return inserted, updated


def _stage_marketplace_catalog_records(db, acc, result):
    """Idempotently stage WB catalogue pages without changing financial states."""
    inserted = updated = 0
    for item in result.records:
        if str(item.get("record_type") or "").lower() != "product":
            continue
        external_id = str(item.get("external_id") or "").strip()
        if not external_id:
            continue
        row = db.scalar(select(MarketplaceRecord).where(
            MarketplaceRecord.account_id == acc.id,
            MarketplaceRecord.source == "wildberries",
            MarketplaceRecord.external_id == external_id,
        ))
        payload = jsonable_encoder(item)
        if row:
            if (row.payload or {}) != payload:
                row.payload = payload
                updated += 1
            continue
        try:
            with db.begin_nested():
                db.add(MarketplaceRecord(account_id=acc.id, company_id=acc.company_id,
                    source="wildberries", external_id=external_id, payload=payload))
                db.flush()
            inserted += 1
        except IntegrityError:
            # A normal marketplace sync can stage the same card concurrently.
            row = db.scalar(select(MarketplaceRecord).where(
                MarketplaceRecord.account_id == acc.id,
                MarketplaceRecord.source == "wildberries",
                MarketplaceRecord.external_id == external_id,
            ))
            if row is None:
                raise
            if (row.payload or {}) != payload:
                row.payload = payload
                updated += 1
    return inserted, updated


def _background_kaspi_period_sync(account_id, user_id, start, end, extra_order_ids, report_mode="issued", status_only=False, expected_lease=None):
    """Refresh only a requested Kaspi creation-date window plus old open orders."""
    with SessionLocal() as db:
        acc = db.get(Account, account_id)
        user = db.get(User, user_id)
        if not acc:
            return
        warnings = []
        try:
            if not user or not user.active:
                raise AccessRequired("Пользователь отключён")
            require(user, "accounts", True)
            scope(user, acc.company_id)
            worker = connector(acc, load_saved_credentials(db, acc.id))
            lease = acc.sync_error
            if expected_lease is not None and lease != expected_lease:
                return
            job = dict((acc.integration_settings or {}).get("kaspi_period_job") or {})
            if job and (job.get("start") != str(start) or job.get("end") != str(end)):
                return
            today_local = datetime.now(ZoneInfo("Asia/Almaty"))
            sync_start = _kaspi_issued_scan_start(today_local.date()) if report_mode == "issued" else start
            local_start = datetime.combine(sync_start, datetime.min.time(), ZoneInfo("Asia/Almaty"))
            local_end = (today_local if end >= today_local.date() else
                         datetime.combine(end + timedelta(days=1), datetime.min.time(), ZoneInfo("Asia/Almaty")))
            event_start = datetime.combine(start, datetime.min.time(), ZoneInfo("Asia/Almaty"))
            cached_children_by_order = {}
            known_orders = {}
            order_ids_with_entries = set()
            if status_only:
                # Known orders with cached line items need only a status refresh.
                for cached in db.scalars(select(MarketplaceRecord).where(
                    MarketplaceRecord.account_id == acc.id,
                    MarketplaceRecord.source == "kaspi",
                )).all():
                    child = cached.payload or {}
                    record_type = child.get("record_type")
                    order_id = str(child.get("order_id") or "")
                    if record_type == "order" and child.get("code"):
                        known_orders[order_id] = str(child["code"])
                    elif record_type in {"order_entry", "delivery_charge", "return_delivery_charge"}:
                        cached_children_by_order.setdefault(order_id, []).append(child)
                        if record_type == "order_entry":
                            order_ids_with_entries.add(order_id)
            known_order_ids = [code for oid, code in known_orders.items() if oid in order_ids_with_entries]
            known_entry_counts = {}
            for cached in db.scalars(select(MarketplaceRecord).where(
                MarketplaceRecord.account_id == acc.id, MarketplaceRecord.source == "kaspi",
                MarketplaceRecord.payload["record_type"].as_string() == "order",
            )):
                p = cached.payload or {}
                oid = str(p.get("order_id") or "")
                count = sum(x.get("record_type") == "order_entry" for x in cached_children_by_order.get(oid, []))
                if p.get("entries_complete") is True and count > 0 and count == p.get("entries_count"):
                    known_entry_counts[str(p.get("code") or "")] = count
            progress_key = _kaspi_period_sync_key(report_mode, start, end)
            settings = dict(acc.integration_settings or {})
            progress = dict(settings.get("kaspi_period_progress") or {})
            cursor = progress[progress_key] if progress_key in progress else json.dumps({"_qor_kaspi_cursor":1,"scope":"status_only" if status_only else "period",
                "scan_version":3,"state_index":0,"known_entry_counts":known_entry_counts,
                "start_ms":int(local_start.timestamp()*1000),"until_ms":int(local_end.timestamp()*1000)-1,
                "event_start_ms":int(event_start.timestamp()*1000),
                "event_end_ms":int(local_end.timestamp()*1000)-1,
                "page":0,"offset":0,"extra_ids":([] if status_only else extra_order_ids),"extra_offset":0,
                "known_order_ids":known_order_ids}, separators=(",", ":"))
            pages = records = inserted_total = updated_total = 0
            while cursor:
                # Do not occupy a pooled DB connection while waiting on Kaspi.
                db.commit()
                result = worker.fetch(cursor)
                db.refresh(acc)
                if acc.sync_error != lease or not acc.active:
                    db.rollback()
                    return
                if status_only:
                    present = {str(row.get("external_id") or "") for row in result.records}
                    propagated = []
                    for order in result.records:
                        op = order or {}
                        if op.get("record_type") != "order":
                            continue
                        oid = str(op.get("order_id") or "")
                        attrs = op.get("raw_attributes") or {}
                        for prior in cached_children_by_order.get(oid, []):
                            child = dict(prior)
                            child["order_status"] = op.get("status")
                            child["order_state"] = op.get("state")
                            child["completed_date"] = op.get("completion_date") or child.get("completed_date")
                            child["returned_date"] = op.get("returned_date") or child.get("returned_date")
                            if child.get("record_type") == "delivery_charge":
                                child["completed_date"] = op.get("completion_date")
                                if op.get("delivery_cost_for_seller") is not None:
                                    child["amount"] = op.get("delivery_cost_for_seller")
                            elif child.get("record_type") == "return_delivery_charge":
                                if op.get("return_delivery_cost_for_seller") is not None:
                                    child["amount"] = op.get("return_delivery_cost_for_seller")
                            if child.get("external_id") not in present:
                                propagated.append(child)
                                present.add(child.get("external_id"))
                        return_cost = op.get("return_delivery_cost_for_seller")
                        if return_cost not in (None, "", 0, 0.0, "0"):
                            ext = "return-delivery:" + oid
                            if ext not in present:
                                propagated.append({
                                    "external_id": ext, "record_type": "return_delivery_charge",
                                    "order_id": oid, "amount": return_cost, "currency": "KZT",
                                    "order_status": op.get("status"), "returned_date": op.get("returned_date"),
                                    "order_creation_date": op.get("creation_date"),
                                    "completion_date": op.get("completion_date"), "raw_attributes": attrs,
                                })
                                present.add(ext)
                    result.records.extend(propagated)
                inserted, updated = _stage_marketplace_records(db, acc, result)
                record_ids = [str(x.get("external_id")) for x in result.records if x.get("external_id")]
                projection = marketplace_pipeline.project_account(db, user, acc, external_ids=record_ids, reconcile_pending=False)
                pages += 1
                records += len(result.records)
                inserted_total += inserted
                updated_total += updated
                warnings.extend(result.warnings or [])
                cursor = result.next_cursor
                acc.last_sync = now()
                acc.sync_status = "running"
                acc.sync_error = lease
                progress[progress_key] = cursor
                settings = dict(acc.integration_settings or {})
                settings["kaspi_period_progress"] = dict(progress)
                position = json.loads(cursor) if cursor else {}
                type_counts = {}
                for record in result.records:
                    record_type = str(record.get("record_type") or "other")
                    type_counts[record_type] = type_counts.get(record_type, 0) + 1
                job.update(status="running", batches=int(job.get("batches") or 0)+1,
                           records=int(job.get("records") or 0)+len(result.records),
                           order_list_pages=int(job.get("order_list_pages") or 0)+1,
                           order_records_received=int(job.get("order_records_received") or 0)+type_counts.get("order", 0),
                           order_entries=int(job.get("order_entries") or 0)+type_counts.get("order_entry", 0),
                           scan_from=str(sync_start),
                           scan_through=datetime.fromtimestamp(position.get("until_ms", int(local_end.timestamp()*1000))/1000, ZoneInfo("Asia/Almaty")).date().isoformat(),
                           state=("ARCHIVE", "NEW", "SIGN_REQUIRED", "PICKUP", "DELIVERY", "KASPI_DELIVERY")[int(position.get("state_index") or 0)],
                           page=position.get("page", 0), page_count=position.get("page_count"),
                           offset=position.get("offset", 0), page_size=position.get("page_size"))
                settings["kaspi_period_job"] = dict(job)
                acc.integration_settings = settings
                db.commit()
                if pages >= 20000:
                    raise AccessRequired("Kaspi: превышен защитный лимит страниц диапазона")
            marketplace_pipeline.project_account(db, user, acc, external_ids=[])
            acc = db.get(Account, account_id)
            acc.last_sync = now()
            acc.sync_status = "partial" if warnings else "done"
            acc.sync_error = "; ".join(warnings)[:500]
            settings = dict(acc.integration_settings or {})
            progress = dict(settings.get("kaspi_period_progress") or {})
            progress.pop(progress_key, None)
            settings["kaspi_period_progress"] = progress
            job["status"] = "partial" if warnings else "done"
            settings["kaspi_period_job"] = dict(job)
            acc.integration_settings = settings
            if not warnings:
                coverage = dict(settings.get("kaspi_period_sync") or {})
                coverage[progress_key] = now().isoformat()
                settings["kaspi_period_sync"] = dict(list(coverage.items())[-100:])
                acc.integration_settings = settings
            audit(db, user, "sync_kaspi_period", "accounts", acc.id,
                  {"start":str(start),"end":str(end),"report_mode":report_mode,"status_only":status_only,"records":records,"pages":pages,
                   "inserted":inserted_total,"updated":updated_total})
            db.commit()
        except Exception as exc:
            logger.exception("Kaspi period sync failed for account %s", account_id)
            db.rollback()
            acc = db.get(Account, account_id)
            if acc and acc.sync_error == locals().get("lease"):
                acc.last_sync = now()
                acc.sync_status = "failed"
                acc.sync_error = str(getattr(exc, "detail", exc))[:500]
                settings = dict(acc.integration_settings or {})
                job = dict(settings.get("kaspi_period_job") or {})
                attempts = int(job.get("failures") or 0) + 1
                job.update(status="retry" if attempts < 3 else "failed", failures=attempts,
                           retry_at=(now()+timedelta(seconds=60)).isoformat())
                settings["kaspi_period_job"] = job
                acc.integration_settings = settings
                db.commit()


@app.post("/api/accounts/{id}/sync-kaspi-period")
def sync_kaspi_period(id: int, tasks: BackgroundTasks, start: date, end: date,
                      report_mode: str = Query(default="issued", regex="^(created|issued)$"),
                      force: bool = False,
                      u=Depends(current_user), db=Depends(get_db)):
    """Schedule a bounded Kaspi sync and refresh previously open/very recent orders."""
    require(u, "accounts", True)
    acc = db.get(Account, id)
    if not acc:
        raise HTTPException(404, "Интеграция не найдена")
    scope(u, acc.company_id)
    if acc.channel != "Kaspi" or acc.mode != "api" or not acc.active:
        raise HTTPException(409, "Активная API-интеграция Kaspi не найдена")
    if end < start or (end - start).days > 366:
        raise HTTPException(422, "Укажите корректный период не длиннее 367 дней")
    sync_is_running = sync_status_for_display(acc)[0] == "running"
    abandoned_kaspi_period_sync = (
        sync_is_running
        and acc.channel == "Kaspi"
        and str(acc.sync_error or "").startswith("KASPI_PERIOD_SYNC")
        and (acc.last_sync is None or acc.last_sync < now()-timedelta(minutes=5))
    )
    if sync_is_running and not abandoned_kaspi_period_sync:
        return {"ok":True,"status":"running","already_running":True,"account_id":acc.id,
                "start":str(start),"end":str(end),"report_mode":report_mode}
    local_today = datetime.now(ZoneInfo("Asia/Almaty")).date()
    if start > local_today:
        raise HTTPException(422, "Начало периода не может быть в будущем")
    bounded_end = min(end, local_today)
    if report_mode == "issued" and bounded_end < _kaspi_issued_scan_start(local_today):
        raise HTTPException(422, "Этот период старше настроенного окна загрузки; увеличьте глубину истории")
    period_key = _kaspi_period_sync_key(report_mode, start, bounded_end)
    settings = dict(acc.integration_settings or {})
    cache = dict(settings.get("kaspi_period_sync") or {})
    cached_at = None
    if cache.get(period_key):
        try: cached_at = datetime.fromisoformat(str(cache[period_key]))
        except ValueError: cached_at = None
    age = (now() - cached_at).total_seconds() if cached_at else None
    cache_ttl = 5 * 60 if start == bounded_end == local_today else (10 * 60 if bounded_end == local_today - timedelta(days=1) else 3600)
    cache_fresh = cached_at is not None and age is not None and age < cache_ttl
    if cache_fresh and not force:
        return {"ok":True,"status":"cached","cached":True,"account_id":acc.id,"start":str(start),"end":str(bounded_end),"report_mode":report_mode}
    refresh_ids = set()
    status_only = report_mode == "issued"
    stale_period_claim = and_(
        Account.channel == "Kaspi",
        Account.sync_error.like("KASPI_PERIOD_SYNC%"),
        or_(Account.last_sync < now()-timedelta(minutes=5), Account.last_sync.is_(None)),
    )
    job = dict(settings.get("kaspi_period_job") or {})
    if job.get("start") != str(start) or job.get("end") != str(bounded_end) or job.get("status") == "done":
        job = {"start":str(start), "end":str(bounded_end), "mode":report_mode,
               "user_id":u.id, "batches":0, "records":0}
    job.update(status="running", failures=0)
    settings["kaspi_period_job"] = job
    lease = "KASPI_PERIOD_SYNC:" + uuid.uuid4().hex
    claimed = db.execute(update(Account).where(Account.id == id,
        or_(Account.sync_status != "running", Account.last_sync < now()-timedelta(seconds=SYNC_RUNNING_TIMEOUT_SECONDS), Account.last_sync.is_(None), stale_period_claim))
        .values(sync_status="running", sync_error=lease, last_sync=now(), integration_settings=settings))
    if claimed.rowcount != 1:
        return {"ok":True,"status":"running","already_running":True,"account_id":acc.id,
                "start":str(start),"end":str(bounded_end),"report_mode":report_mode}
    db.commit()
    tasks.add_task(_background_kaspi_period_sync, acc.id, u.id, start, bounded_end, ([] if status_only else sorted(refresh_ids)), report_mode, status_only, lease)
    return {"ok":True,"status":"running","account_id":acc.id,"start":str(start),"end":str(bounded_end),
            "report_mode":report_mode,"old_orders_to_refresh":len(refresh_ids),"status_only":status_only}


@app.post("/api/accounts/{id}/sync-now")
def account_sync_now(id: int, tasks: BackgroundTasks, background: bool = False, sections: str | None = None,
                     u=Depends(current_user), db=Depends(get_db)):
    """Run marketplace sync immediately in the web process.

    Render free/web deployments do not require a separate queue worker for marketplace checks.
    The endpoint always finishes as success or error and persists the account status.
    """
    require(u, "accounts", True)
    acc = db.get(Account, id)
    if not acc:
        raise HTTPException(404, "Интеграция не найдена")
    scope(u, acc.company_id)
    if not acc.active:
        raise HTTPException(409, "Аккаунт отключён. Подключите его заново перед синхронизацией")
    claimed = db.execute(update(Account).where(Account.id == id,
        Account.active.is_(True),
        or_(Account.sync_status != 'running', Account.last_sync < now()-timedelta(seconds=SYNC_RUNNING_TIMEOUT_SECONDS), Account.last_sync.is_(None)))
        .values(sync_status='running', sync_error='', last_sync=now()))
    if claimed.rowcount != 1:
        raise HTTPException(409, 'Синхронизация уже выполняется')
    db.commit()
    selected_sections = None
    if acc.channel == "Wildberries" and acc.mode == "api":
        allowed = WB_ALLOWED_SYNC_SECTIONS
        selected_sections = [x.strip() for x in sections.split(",") if x.strip()] if sections else list(WB_DEFAULT_SYNC_SECTIONS)
        unknown = set(selected_sections) - allowed
        if not selected_sections or unknown:
            acc.sync_status = "failed"
            acc.sync_error = "Неизвестные разделы WB: " + ", ".join(sorted(unknown or {"пустой список"}))
            db.commit()
            raise HTTPException(422, acc.sync_error)
    if background and acc.mode == 'api':
        tasks.add_task(_background_marketplace_sync, acc.id, u.id, selected_sections)
        return {'ok':True, 'status':'running', 'account_id':acc.id,
                **({'sections': selected_sections, 'catalog_included': 'products' in selected_sections} if selected_sections is not None else {})}
    return _perform_marketplace_sync(db, u, acc, sections=selected_sections)


def _perform_marketplace_sync(db, u, acc, sections=None):
    account_id = acc.id
    try:
        creds = load_saved_credentials(db, acc.id)
        worker = connector(acc, creds)
        if acc.channel == "Wildberries" and acc.mode == "api":
            # WB is synchronized as independent durable streams. Each page is staged
            # and committed before the next remote request; no page payloads are kept
            # in memory after staging.
            import json as _json
            state = _json.loads(acc.sync_cursor) if acc.sync_cursor else {}
            if not isinstance(state, dict): state = {}
            section_state = state.setdefault('_sections', {})
            warnings, counts = [], {}
            total_records = total_inserted = total_updated = 0
            successful_sections = 0
            selected_sections = list(sections or worker.sections)
            for section in selected_sections:
                section_inserted = section_updated = section_pages = section_records = 0
                section_state[section] = {'status': 'running'}
                acc.sync_cursor = _json.dumps(state, separators=(',', ':'))
                db.commit()
                try:
                    for piece in worker.fetch_section_pages(section, _json.dumps(state, separators=(',', ':'))):
                        added, changed = _stage_marketplace_records(db, acc, piece)
                        section_inserted += added; section_updated += changed
                        section_pages += 1; section_records += len(piece.records)
                        total_records += len(piece.records); total_inserted += added; total_updated += changed
                        if piece.next_cursor:
                            next_state = _json.loads(piece.next_cursor)
                            if isinstance(next_state, dict):
                                # Preserve local durable section metadata while accepting
                                # remote cursors advanced by the connector.
                                next_state['_sections'] = section_state
                                state = next_state
                        section_state = state.setdefault('_sections', section_state)
                        section_state[section] = {'status':'running','pages':section_pages,'records':section_records}
                        acc.sync_cursor = _json.dumps(state, separators=(',', ':'))
                        db.commit()
                    successful_sections += 1
                    section_state[section] = {'status':'done','pages':section_pages,'records':section_records}
                    counts[section] = {'pages':section_pages,'records':section_records,'inserted':section_inserted,'updated':section_updated}
                    acc.sync_cursor = _json.dumps(state, separators=(',', ':'))
                    db.commit()
                except Exception as section_exc:
                    db.rollback()
                    acc = db.get(Account, account_id)
                    # Reload the last committed cursor so a failed page is retried rather
                    # than skipped. Persist a terminal state for this individual stream.
                    try:
                        state = _json.loads(acc.sync_cursor) if acc.sync_cursor else state
                    except (ValueError, TypeError):
                        pass
                    if not isinstance(state, dict): state = {}
                    section_state = state.setdefault('_sections', {})
                    message = str(section_exc)[:180]
                    section_state[section] = {'status':'failed','pages':section_pages,'records':section_records,'error':message}
                    acc.sync_cursor = _json.dumps(state, separators=(',', ':'))
                    warnings.append(f"{section}: {message}")
                    counts[section] = {'pages':section_pages,'records':section_records,'inserted':section_inserted,'updated':section_updated,'error':message}
                    db.commit()
                    logger.exception("WB section %s failed for account %s", section, account_id)
            # One idempotent ERP projection after all remote streams have finished.
            # This avoids scanning the full staging history six times per sync.
            financial_sections = any(section != 'order_statuses' for section in selected_sections)
            projection = marketplace_pipeline.project_account(db, u, acc) if total_records and financial_sections else {'projected':0,'pending_erp':0}
            acc = db.get(Account, account_id)
            acc.last_sync = now()
            acc.sync_status = 'failed' if successful_sections == 0 else ('partial' if warnings else 'done')
            acc.sync_error = '; '.join(warnings)[:500]
            audit(db, u, 'sync_now', 'accounts', account_id, {'channel':'Wildberries','sections':counts,'projection':projection})
            db.commit()
            return {'ok': successful_sections > 0, 'status':acc.sync_status,'account_id':account_id,
                    'sections':counts,'warnings':warnings,'records':total_records,
                    'staged_inserted':total_inserted,'staged_updated':total_updated,'erp_projection':projection}
        if acc.channel == "Kaspi" and acc.mode == "api":
            import json as _json
            cursor = acc.sync_cursor or None
            total_records = total_inserted = total_updated = 0
            incremental_projected = 0
            warnings = []
            # Project already-staged events from the active 7-day window first.
            # This recovers the recent rows saved before the incremental projector
            # was deployed, without waiting for the 90-day backfill to finish.
            try:
                cursor_state = _json.loads(cursor) if cursor else {}
            except (ValueError, TypeError):
                cursor_state = {}
            range_start = datetime.fromtimestamp(cursor_state["start_ms"] / 1000, timezone.utc).date() if cursor_state.get("start_ms") else (now() - timedelta(days=7)).date()
            range_end = datetime.fromtimestamp(cursor_state["until_ms"] / 1000, timezone.utc).date() if cursor_state.get("until_ms") else now().date()
            recent_projection = marketplace_pipeline.project_account(
                db, u, acc, date_from=range_start, date_to=range_end)
            incremental_projected += recent_projection.get("projected", 0)
            db.commit()
            recovery_attempt = 2 if "attempt=2" in (acc.sync_error or "") else (1 if "SYNC_RECOVERY_STARTED" in (acc.sync_error or "") else 0)
            recovery_in_progress = recovery_attempt > 0
            chunks = 0
            while True:
                result = worker.fetch(cursor)
                inserted, updated = _stage_marketplace_records(db, acc, result)
                batch_external_ids = [str(item.get("external_id") or "") for item in result.records
                                      if item.get("external_id")]
                if batch_external_ids:
                    batch_projection = marketplace_pipeline.project_account(
                        db, u, acc, external_ids=batch_external_ids)
                    incremental_projected += batch_projection.get("projected", 0)
                total_records += len(result.records)
                total_inserted += inserted
                total_updated += updated
                warnings.extend(result.warnings or [])
                cursor = result.next_cursor
                acc.sync_cursor = cursor or ""
                acc.last_sync = now()
                acc.sync_status = "running"
                acc.sync_error = (f"SYNC_RECOVERY_STARTED: attempt={recovery_attempt}" if recovery_in_progress
                                  else "; ".join(warnings)[-500:])
                db.commit()
                chunks += 1
                if not cursor:
                    break
                if chunks >= 20000:
                    raise AccessRequired("Kaspi: достигнут защитный лимит пачек синхронизации; повторите позже")
            projection = marketplace_pipeline.project_account(db, u, acc) if total_records else {"projected":0,"pending_erp":0}
            acc = db.get(Account, account_id)
            acc.last_sync = now()
            acc.sync_status = "partial" if warnings else "done"
            acc.sync_error = "; ".join(warnings)[:500]
            acc.sync_cursor = ""
            audit(db, u, "sync_now", "accounts", acc.id, {"channel":"Kaspi","records":total_records,"windows":chunks})
            db.commit()
            return {"ok":True,"status":acc.sync_status,"account_id":acc.id,"records":total_records,
                    "windows":chunks,"staged_inserted":total_inserted,"staged_updated":total_updated,
                    "incremental_projected":incremental_projected,
                    "warnings":warnings,"erp_projection":projection}
        result = worker.fetch(acc.sync_cursor or None)
        inserted, updated = _stage_marketplace_records(db, acc, result)
        db.flush()
        projection = marketplace_pipeline.project_account(db, u, acc)
        payload = __import__("fastapi").encoders.jsonable_encoder(asdict(result))
        payload["staged_inserted"] = inserted
        payload["staged_updated"] = updated
        payload["erp_projection"] = projection
        acc.last_sync = now()
        acc.sync_status = "partial" if result.warnings else "done"
        acc.sync_error = "; ".join(result.warnings or [])[:500]
        acc.sync_cursor = result.next_cursor or acc.sync_cursor or ""
        audit(db, u, "sync_now", "accounts", acc.id, {"channel": acc.channel, "records": len(result.records)})
        db.commit()
        return {"ok": True, "status": acc.sync_status, "account_id": acc.id, **payload}
    except Exception as exc:
        logger.exception("Marketplace sync failed for account %s", account_id)
        db.rollback()
        acc = db.get(Account, account_id)
        message = str(getattr(exc, "detail", exc))[:1000]
        if acc:
            acc.sync_status = "failed"
            acc.sync_error = message
            acc.last_sync = now()
            db.commit()
        raise HTTPException(422, message)



def _background_marketplace_sync(account_id, user_id, sections=None):
    with SessionLocal() as db:
        user = db.get(User, user_id)
        acc = db.get(Account, account_id)
        if not acc: return
        try:
            if not user or not user.active: raise HTTPException(403, 'Пользователь отключён')
            require(user, 'accounts', True)
            scope(user, acc.company_id)
            _perform_marketplace_sync(db, user, acc, sections=sections)
        except Exception as exc:
            logger.exception("Marketplace background sync failed for account %s", account_id)
            db.rollback()
            acc = db.get(Account, account_id)
            if acc:
                acc.sync_status = 'failed'
                acc.sync_error = str(getattr(exc, 'detail', exc))[:500]
                db.commit()


@app.get('/api/accounts/{id}/sync-status')
def account_sync_status(id: int, background_tasks: BackgroundTasks, u=Depends(current_user), db=Depends(get_db)):
    require(u, 'accounts')
    acc = db.get(Account, id)
    if not acc: raise HTTPException(404, 'Интеграция не найдена')
    scope(u, acc.company_id)
    recovery_attempt = 2 if "attempt=2" in (acc.sync_error or "") else (1 if "SYNC_RECOVERY_STARTED" in (acc.sync_error or "") else 0)
    last_sync = acc.last_sync
    if last_sync is not None and last_sync.tzinfo is None:
        last_sync = last_sync.replace(tzinfo=timezone.utc)
    if acc.sync_status == 'running' and (not last_sync or last_sync < now()-timedelta(seconds=SYNC_RUNNING_TIMEOUT_SECONDS)):
        # A worker can be killed by a deploy/restart. Mark the stale run and allow
        # at most two bounded recovery attempts for Kaspi.
        acc.sync_status = 'failed'
        prefix = f'SYNC_RECOVERY_STARTED: attempt={recovery_attempt}; ' if recovery_attempt else ''
        acc.sync_error = prefix + 'SYNC_TIMEOUT: синхронизация превысила лимит времени'
        db.commit()
    if (acc.channel == 'Kaspi' and acc.mode == 'api' and acc.sync_status == 'failed'
            and 'SYNC_TIMEOUT:' in (acc.sync_error or '') and recovery_attempt < 2):
        recovery_attempt += 1
        acc.sync_status = 'running'
        acc.last_sync = now()
        acc.sync_error = f'SYNC_RECOVERY_STARTED: attempt={recovery_attempt}'
        db.commit()
        background_tasks.add_task(_background_marketplace_sync, acc.id, u.id)
    counts = dict(db.execute(select(MarketplaceRecord.status, func.count()).where(
        MarketplaceRecord.account_id == id, MarketplaceRecord.company_id == acc.company_id).group_by(MarketplaceRecord.status)).all())
    return {'account_id':id, 'status':acc.sync_status, 'error':acc.sync_error,
            'last_sync':acc.last_sync, 'total_records':sum(counts.values()),
            'erp_projection':{'total_projected':counts.get('projected',0),
                              'pending_erp':sum(v for k,v in counts.items() if k.startswith('pending'))}}


@app.post("/api/accounts/{id}/test-connection")
def account_test_connection(id: int, u=Depends(current_user), db=Depends(get_db)):
    """Validate credentials through the same connector without staging ERP records."""
    require(u, "accounts", True)
    acc = db.get(Account, id)
    if not acc:
        raise HTTPException(404, "Интеграция не найдена")
    scope(u, acc.company_id)
    if not acc.active:
        raise HTTPException(409, "Аккаунт отключён. Подключите его заново")
    try:
        result = connector(acc, load_saved_credentials(db, acc.id)).test_connection()
        return {
            "ok": True,
            "account_id": acc.id,
            "channel": acc.channel,
            "message": result["message"],
            "records_checked": result["records_checked"],
        }
    except Exception as exc:
        message = str(getattr(exc, "detail", exc))[:500]
        raise HTTPException(422, message)


@app.get("/api/accounts/{id}/connection-status")
def account_connection_status(id: int, u=Depends(current_user), db=Depends(get_db)):
    require(u, "accounts")
    acc = db.get(Account, id)
    if not acc:
        raise HTTPException(404, "Аккаунт не найден")
    scope(u, acc.company_id)
    # Return secret NAMES only, never values.
    row = db.scalar(select(IntegrationSecret).where(IntegrationSecret.account_id == acc.id))
    credential_error = ""
    try:
        creds = load_saved_credentials(db, acc.id)
    except AccessRequired as exc:
        creds = {}
        credential_error = str(exc)
    status = connector(acc, creds).connection_status()
    return {**status, "account_id": acc.id, "channel": acc.channel, "sync_status": acc.sync_status,
            "last_sync": acc.last_sync, "sync_error": acc.sync_error, "saved_on_site": bool(row),
            "credential_hint": row.key_hint if row else "", "credential_error": credential_error}


def _validate_product_card(db, card):
    product = db.get(Product, card.product_id)
    errors=[]; warnings=[]
    if not product: errors.append("Товар не найден")
    if not card.title.strip(): errors.append("Название обязательно")
    if not card.description.strip(): errors.append("Описание обязательно")
    if not card.image_document_ids: errors.append("Нужно минимум одно изображение товара")
    facts={"brand": getattr(product,"brand",None), "country": getattr(product,"country",None), "weight": getattr(product,"weight_g",None), "volume": getattr(product,"volume_cm3",None)} if product else {}
    # Claims explicitly marked as sourced must match ERP facts. Unknown claims are rejected.
    for key,val in (card.characteristics or {}).items():
        if key in facts and facts[key] not in (None, "") and str(val).strip()!=str(facts[key]).strip(): errors.append(f"{key}: значение не совпадает с ERP")
        elif key in facts and facts[key] in (None, ""): errors.append(f"{key}: характеристика не подтверждена ERP")
    if card.marketplace not in {"kaspi","wildberries","ozon"}: errors.append("Неподдерживаемый marketplace")
    return {"marketplace":card.marketplace,"passed":not errors,"errors":errors,"warnings":warnings,"checked_at":datetime.now(timezone.utc).isoformat(),"rule_set":"qor-marketplace-safe-v1","note":"Финальная модерация выполняется самим marketplace; QOR не гарантирует принятие площадкой."}

@app.get("/api/product-studio/cards")
def product_cards(product_id: int | None = None, u=Depends(current_user), db=Depends(get_db)):
    require(u,"marketing")
    q=select(ProductCardVersion).order_by(ProductCardVersion.id.desc())
    if product_id: q=q.where(ProductCardVersion.product_id==product_id)
    rows=db.scalars(q).all(); out=[]
    for x in rows:
        scope(u,x.company_id); out.append(dump(x))
    return {"items":out}

@app.post("/api/product-studio/cards")
def product_card_create(p: dict, u=Depends(current_user), db=Depends(get_db)):
    require(u,"marketing",True)
    product=db.get(Product,int(p.get("product_id") or 0))
    if not product: raise HTTPException(404,"Товар не найден")
    cid=int(p.get("company_id") or 0)
    if not cid: raise HTTPException(422,"company_id обязателен для карточки")
    scope(u,cid)
    mp=str(p.get("marketplace","kaspi")).lower()
    prev=db.scalar(select(func.max(ProductCardVersion.version)).where(ProductCardVersion.product_id==product.id,ProductCardVersion.marketplace==mp)) or 0
    obj=ProductCardVersion(company_id=cid,product_id=product.id,marketplace=mp,version=prev+1,title=str(p.get("title") or product.name)[:240],description=str(p.get("description") or ""),characteristics=p.get("characteristics") or {},image_document_ids=[int(x) for x in p.get("image_document_ids",[])],parent_version_id=p.get("parent_version_id"))
    db.add(obj); db.flush(); obj.validation=_validate_product_card(db,obj); obj.status="validated" if obj.validation["passed"] else "needs_fix"; audit(db,u,"create","product_card_versions",obj.id,{"marketplace":mp,"version":obj.version}); db.commit(); return dump(obj)

@app.post("/api/product-studio/cards/{id}/validate")
def product_card_validate(id:int,u=Depends(current_user),db=Depends(get_db)):
    require(u,"marketing",True); obj=db.get(ProductCardVersion,id)
    if not obj: raise HTTPException(404,"Карточка не найдена")
    scope(u,obj.company_id); obj.validation=_validate_product_card(db,obj); obj.status="validated" if obj.validation["passed"] else "needs_fix"; db.commit(); return dump(obj)

@app.post("/api/product-studio/cards/{id}/publish")
def product_card_publish(id:int,p:dict,u=Depends(current_user),db=Depends(get_db)):
    require(u,"marketing",True); obj=db.get(ProductCardVersion,id)
    if not obj: raise HTTPException(404,"Карточка не найдена")
    scope(u,obj.company_id); obj.validation=_validate_product_card(db,obj)
    if not obj.validation["passed"]: raise HTTPException(422,"Карточка не прошла QOR Validator")
    # Never fake marketplace publication. Until a live write adapter is connected, mark ready_for_publish only.
    obj.status="ready_for_publish"; obj.review_after=datetime.now(timezone.utc)+timedelta(days=max(1,int(p.get("review_days",7)))); obj.review_state="scheduled"; obj.optimization_policy={"auto_optimize":bool(p.get("auto_optimize",True)),"auto_publish":bool(p.get("auto_publish",False)),"rollback":bool(p.get("rollback",True)),"review_days":max(1,int(p.get("review_days",7))),"max_generations":max(1,min(10,int(p.get("max_generations",3))))}
    audit(db,u,"ready_for_publish","product_card_versions",obj.id,{"review_after":obj.review_after.isoformat()}); db.commit(); return dump(obj)

@app.post("/api/product-studio/cards/{id}/metrics")
def product_card_metrics(id:int,p:dict,u=Depends(current_user),db=Depends(get_db)):
    require(u,"marketing",True); obj=db.get(ProductCardVersion,id)
    if not obj: raise HTTPException(404,"Карточка не найдена")
    scope(u,obj.company_id); allowed={k:p[k] for k in ("impressions","views","clicks","orders","revenue","returns","ad_spend","conversion") if k in p}; obj.metrics={**(obj.metrics or {}),**allowed,"observed_at":datetime.now(timezone.utc).isoformat()}; db.commit(); return dump(obj)

@app.post("/api/product-studio/cards/{id}/optimize")
def product_card_optimize(id:int,u=Depends(current_user),db=Depends(get_db)):
    require(u,"marketing",True); obj=db.get(ProductCardVersion,id)
    if not obj: raise HTTPException(404,"Карточка не найдена")
    scope(u,obj.company_id); m=obj.metrics or {}
    if not m or not any(k in m for k in ("views","clicks","orders","conversion")): raise HTTPException(422,"Данные недостаточны для точного вывода")
    assessment=assess_performance(m)
    if assessment["decision"] == "insufficient_data": raise HTTPException(422,assessment["reason"])
    reason=assessment["reason"]; experiment=next_experiment(obj.version)
    prev=db.scalar(select(func.max(ProductCardVersion.version)).where(ProductCardVersion.product_id==obj.product_id,ProductCardVersion.marketplace==obj.marketplace)) or obj.version
    new=ProductCardVersion(company_id=obj.company_id,product_id=obj.product_id,marketplace=obj.marketplace,version=prev+1,title=obj.title,description=obj.description,characteristics=obj.characteristics,image_document_ids=obj.image_document_ids,status="optimization_draft",parent_version_id=obj.id,optimization_policy=obj.optimization_policy or {},review_state="idle",validation={"passed":False,"errors":[],"warnings":[reason],"requires_ai_regeneration":True,"experiment":experiment})
    db.add(new); db.flush(); audit(db,u,"optimize_proposal","product_card_versions",new.id,{"parent":obj.id,"reason":reason}); db.commit(); return dump(new)

@app.get("/api/ai/threads")
def ai_threads(u=Depends(current_user), db=Depends(get_db)):
    require(u, "ai")
    rows = db.scalars(select(AIThread).where(AIThread.user_id == u.id, AIThread.archived == False).order_by(AIThread.id.desc())).all()
    return {"items": [dump(x) for x in rows]}


@app.post("/api/ai/threads")
def ai_thread_create(p: dict, u=Depends(current_user), db=Depends(get_db)):
    require(u, "ai", True)
    agent = str(p.get("agent", "rick"))
    if agent not in ai.AI_AGENTS:
        raise HTTPException(422, "Неизвестный AI-сотрудник")
    cid = p.get("company_id")
    if cid:
        f.get_company(db, u, int(cid))
    obj = AIThread(user_id=u.id, company_id=int(cid) if cid else None, agent=agent, title=str(p.get("title") or "Новый чат")[:180])
    db.add(obj); db.flush(); audit(db, u, "create", "ai_threads", obj.id); db.commit()
    return dump(obj)


@app.get("/api/ai/threads/{id}/messages")
def ai_thread_messages(id: int, u=Depends(current_user), db=Depends(get_db)):
    require(u, "ai")
    t = db.get(AIThread, id)
    if not t or t.user_id != u.id or t.archived:
        raise HTTPException(404, "Чат не найден")
    rows = db.scalars(select(AIMessage).where(AIMessage.thread_id == id).order_by(AIMessage.id)).all()
    return {"thread": dump(t), "items": [dump(x) for x in rows]}


@app.post("/api/ai/threads/{id}/messages")
def ai_thread_message(id: int, p: dict, u=Depends(current_user), db=Depends(get_db)):
    require(u, "ai", True)
    t = db.get(AIThread, id)
    if not t or t.user_id != u.id or t.archived:
        raise HTTPException(404, "Чат не найден")
    role = str(p.get("role", "user"))
    if role not in ("user", "assistant"):
        raise HTTPException(422, "Некорректная роль")
    content = str(p.get("content", "")).strip()[:20000]
    attachments = p.get("attachments") or []
    if not content and not attachments:
        raise HTTPException(422, "Пустое сообщение")
    safe_attachments = []
    for item in attachments[:8]:
        did = int(item.get("document_id"))
        d = db.get(Document, did)
        if not d:
            raise HTTPException(404, "Вложение не найдено")
        scope(u, d.company_id)
        safe_attachments.append({"document_id": d.id, "name": d.name, "mime": d.mime})
    obj = AIMessage(thread_id=id, role=role, content=content, attachments=safe_attachments, meta=p.get("meta") or {})
    db.add(obj)
    if role == "user" and t.title == "Новый чат" and content:
        t.title = content.replace("\n", " ")[:70]
    db.commit()
    return dump(obj)


@app.patch("/api/ai/threads/{id}")
def ai_thread_update(id: int, p: dict, u=Depends(current_user), db=Depends(get_db)):
    require(u, "ai", True)
    t = db.get(AIThread, id)
    if not t or t.user_id != u.id:
        raise HTTPException(404, "Чат не найден")
    if "title" in p: t.title = str(p["title"]).strip()[:180] or "Новый чат"
    if "agent" in p:
        if p["agent"] not in ai.AI_AGENTS: raise HTTPException(422, "Неизвестный AI-сотрудник")
        t.agent = p["agent"]
    if "archived" in p: t.archived = bool(p["archived"])
    db.commit(); return dump(t)


@app.get("/api/ai/actions")
def ai_actions(status: str | None = None, u=Depends(current_user), db=Depends(get_db)):
    require(u, "ai")
    q=select(AIAction).where(AIAction.user_id==u.id)
    if status: q=q.where(AIAction.status==status)
    rows=db.scalars(q.order_by(AIAction.id.desc()).limit(100)).all()
    return {"items":[dump(x) for x in rows],"catalog":action_engine.catalog()}

@app.post("/api/ai/actions/propose")
def ai_action_propose(p: dict, u=Depends(current_user), db=Depends(get_db)):
    require(u,"ai",True)
    row=action_engine.propose(db,u,str(p.get("action_type","")),p.get("payload") or {},p.get("thread_id"),str(p.get("agent","rick")))
    return dump(row)

@app.post("/api/ai/actions/{id}/confirm")
def ai_action_confirm(id:int,tasks:BackgroundTasks,u=Depends(current_user),db=Depends(get_db)):
    require(u,"ai",True); row=db.get(AIAction,id)
    if not row: raise HTTPException(404,"Действие не найдено")
    return dump(action_engine.execute(db,u,row,tasks=tasks))

@app.post("/api/ai/actions/{id}/cancel")
def ai_action_cancel(id:int,u=Depends(current_user),db=Depends(get_db)):
    require(u,"ai",True); row=db.get(AIAction,id)
    if not row: raise HTTPException(404,"Действие не найдено")
    return dump(action_engine.cancel(db,u,row))

@app.get("/api/ai/plans")
def ai_plans(status: str | None = None, u=Depends(current_user), db=Depends(get_db)):
    require(u,"ai")
    q=select(AIActionPlan).where(AIActionPlan.user_id==u.id)
    if status: q=q.where(AIActionPlan.status==status)
    return {"items":[dump(x) for x in db.scalars(q.order_by(AIActionPlan.id.desc()).limit(50)).all()]}

@app.post("/api/ai/plans/propose")
def ai_plan_propose(p: dict, u=Depends(current_user), db=Depends(get_db)):
    require(u,"ai",True)
    return dump(action_engine.propose_plan(db,u,p.get("title"),p.get("goal"),p.get("steps") or [],p.get("thread_id"),str(p.get("agent","rick"))))

@app.post("/api/ai/plans/{id}/confirm")
def ai_plan_confirm(id:int,tasks:BackgroundTasks,u=Depends(current_user),db=Depends(get_db)):
    require(u,"ai",True); row=db.get(AIActionPlan,id)
    if not row: raise HTTPException(404,"План не найден")
    return dump(action_engine.execute_plan(db,u,row,tasks=tasks))

@app.post("/api/ai/plans/{id}/cancel")
def ai_plan_cancel(id:int,u=Depends(current_user),db=Depends(get_db)):
    require(u,"ai",True); row=db.get(AIActionPlan,id)
    if not row: raise HTTPException(404,"План не найден")
    return dump(action_engine.cancel_plan(db,u,row))

@app.post("/api/jobs")
def enqueue(p: dict, tasks: BackgroundTasks, u=Depends(current_user), db=Depends(get_db)):
    kind = p.get("kind")
    payload = p.get("payload", {})
    if kind == "sync":
        raise HTTPException(410, "Marketplace sync перенесён в /api/accounts/{id}/sync-now")
    elif kind in ("assistant", "content", "designer"):
        require(u, "ai" if kind == "assistant" else "marketing", True)
        if kind == "assistant":
            require(u, "analytics")
        cid = payload.get("company_id")
        if cid:
            f.get_company(db, u, int(cid))
        if kind == "designer" and not cid:
            raise HTTPException(422, "Выберите компанию")
    else:
        raise HTTPException(422, "Неизвестное задание")
    job = Job(user_id=u.id, company_id=cid, kind=kind, payload=payload)
    db.add(job)
    db.flush()
    audit(db, u, "enqueue", "jobs", job.id)
    db.commit()
    if kind in ("assistant", "content", "designer"):
        # Return the job immediately; provider calls can take minutes. The worker
        # uses a fresh DB session after the HTTP response has been sent.
        from .worker import run_one
        tasks.add_task(run_one, job.id)
    return dump(job)


@app.get("/api/jobs/{id}")
def job_status(id: int, u=Depends(current_user), db=Depends(get_db)):
    job = db.get(Job, id)
    if not job or job.user_id != u.id:
        raise HTTPException(404, "Задание не найдено")
    stale_after = timedelta(hours=2) if job.kind == "marketplace_catalog_refresh" else timedelta(minutes=15)
    if (job.status == "running" and job.started_at
            and job.started_at < now() - stale_after):
        job.status = "failed"
        job.error = ("Загрузка каталога прервана сервером. Повторите задание: прогресс страниц сохранён."
                     if job.kind == "marketplace_catalog_refresh" else
                     "AI-задание прервано сервером. Проверьте подключение AI и повторите запрос.")
        db.commit()
    return dump(job)


@app.post("/api/jobs/{id}/retry")
def retry_job(id: int, tasks: BackgroundTasks, u=Depends(current_user), db=Depends(get_db)):
    job = db.get(Job, id)
    if not job or job.user_id != u.id:
        raise HTTPException(404, "Не найдено")
    if job.kind == "marketplace_catalog_refresh" and job.status in {"failed", "partial"} and job.attempts < 3:
        job.status = "queued"
        job.error = ""
        db.commit()
        tasks.add_task(_run_wb_product_catalog_refresh, job.id)
        return dump(job)
    if job.status != "failed" or job.attempts >= 3:
        raise HTTPException(409, "Повтор недоступен")
    if job.kind == "card_review":
        job.status = "queued"
        job.error = ""
        db.commit()
        return dump(job)
    payload = dict(job.payload or {})
    kind = job.kind
    db.delete(job)
    db.commit()
    return enqueue({"kind": kind, "payload": payload}, tasks, u, db)


@app.post("/api/documents/upload")
async def upload(
    company_id: int = Form(...),
    file: UploadFile = File(...),
    u=Depends(current_user),
    db=Depends(get_db),
):
    require(u, "documents", True)
    f.get_company(db, u, company_id)
    raw = await file.read(5 * 1024 * 1024 + 1)
    if len(raw) > 5 * 1024 * 1024:
        raise HTTPException(413, "Максимум 5 МБ")
    ext = Path(file.filename or "").suffix.lower()
    allowed_types = {
        ".pdf": "application/pdf",
        ".png": "image/png",
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".webp": "image/webp",
        ".csv": "text/csv",
        ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    }
    if ext not in allowed_types:
        raise HTTPException(422, "Разрешены PDF, PNG, JPEG, WebP, CSV, XLSX")
    if ext in (".png", ".jpg", ".jpeg", ".webp"):
        from PIL import Image

        try:
            im = Image.open(io.BytesIO(raw))
            im.verify()
            if im.width * im.height > 25_000_000:
                raise ValueError()
        except Exception:
            raise HTTPException(422, "Некорректное или слишком большое изображение")
    if ext == ".pdf" and not raw.startswith(b"%PDF"):
        raise HTTPException(422, "Некорректный PDF")
    root = Path(os.getenv("UPLOAD_DIR", "./uploads"))
    root.mkdir(parents=True, exist_ok=True)
    path = root / (uuid.uuid4().hex + ext)
    path.write_bytes(raw)
    obj = Document(
        company_id=company_id,
        name=Path(file.filename).name[:200],
        path=str(path.resolve()),
        mime=allowed_types[ext],
        size=len(raw),
    )
    db.add(obj)
    db.flush()
    audit(db, u, "upload", "documents", obj.id)
    db.commit()
    return dump(obj)


@app.get("/api/documents/{id}/download")
def download(id: int, u=Depends(current_user), db=Depends(get_db)):
    require(u, "documents")
    d = db.get(Document, id)
    if not d:
        raise HTTPException(404, "Не найдено")
    scope(u, d.company_id)
    if not Path(d.path).exists():
        raise HTTPException(404, "Файл отсутствует в хранилище")
    return FileResponse(d.path, filename=d.name, media_type=d.mime)


@app.get("/api/sales/{id}/invoice")
def invoice(id: int, u=Depends(current_user), db=Depends(get_db)):
    require(u, "sales")
    sale = db.get(Sale, id)
    if not sale:
        raise HTTPException(404, "Не найдено")
    scope(u, sale.company_id)
    import html

    company = db.get(Company, sale.company_id)
    product = db.get(Product, sale.product_id)
    customer = db.get(Customer, sale.customer_id) if sale.customer_id else None
    esc = html.escape
    body = f"""<!doctype html><html lang="ru"><meta charset="utf-8"><title>Счёт {esc(sale.reference)}</title><style>body{{font:16px Arial;max-width:800px;margin:60px auto}}table{{border-collapse:collapse;width:100%}}td,th{{border:1px solid #aaa;padding:14px}}</style><h1>Счёт на оплату № {esc(sale.reference)}</h1><p>Дата: {sale.date}</p><p>Поставщик: {esc(company.name)} · БИН/ИИН {esc(company.bin)}</p><p>Покупатель: {esc(customer.name) if customer else "Розничный покупатель"}</p><table><tr><th>Товар</th><th>Количество</th><th>Сумма KZT</th></tr><tr><td>{esc(product.name)}</td><td>{sale.quantity}</td><td>{sale.revenue / 100:,.2f}</td></tr></table><p>Управленческий документ. Банковские реквизиты и налоговые реквизиты необходимо дополнить перед отправкой.</p></html>"""
    return Response(
        body,
        media_type="text/html",
        headers={"Content-Disposition": f'attachment; filename="invoice-{id}.html"'},
    )


@app.get("/api/import/fields/{kind}")
def import_fields(kind: str, u=Depends(current_user)):
    if kind not in importing.KINDS:
        raise HTTPException(404, "Нет такого типа")
    require(u, importing.KINDS[kind][0], True)
    return {
        "fields": importing.FIELDS[kind],
        "money_unit": "KZT (decimal); dates YYYY-MM-DD; IDs from directories",
    }


@app.post("/api/import/parse")
async def import_parse(file: UploadFile = File(...), sheet: str | None = Form(None), header_row: int | None = Form(None), u=Depends(current_user)):
    if not any(
        allowed(u, c, True) for c in ("catalog", "inventory", "sales", "finance")
    ):
        raise HTTPException(403, "Нет прав импорта")
    raw = await file.read(5 * 1024 * 1024 + 1)
    try:
        metadata = {}
        headers, rows = importing.parse_file(raw, file.filename, sheet, header_row, metadata)
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(
            422, "Не удалось прочитать файл; проверьте формат и кодировку UTF-8"
        )
    return {"headers": headers, "rows": rows, "count": len(rows), **metadata}


@app.post("/api/import/preview")
def import_preview(p: dict, u=Depends(current_user), db=Depends(get_db)):
    if len(p.get("rows", [])) > 5000:
        raise HTTPException(413, "Максимум 5000 строк")
    run, normalized = importing.preview(db, u, p["kind"], p["rows"], p["mapping"], p.get("defaults"))
    db.commit()
    return {**dump(run), "preview": normalized[:30], "count": len(run.rows)}


@app.post("/api/import/{id}/confirm")
def import_confirm(id: int, u=Depends(current_user), db=Depends(get_db)):
    run = db.scalar(select(ImportRun).where(ImportRun.id == id).with_for_update())
    if not run:
        raise HTTPException(404, "Не найдено")
    count = importing.confirm(db, u, run)
    db.commit()
    return {"count": count}


@app.get("/api/import/{id}/errors")
def import_errors(id: int, u=Depends(current_user), db=Depends(get_db)):
    run = db.get(ImportRun, id)
    if not run or run.user_id != u.id:
        raise HTTPException(404, "Не найдено")
    return Response(
        importing.export_bytes(run.errors, ["row", "error"], "csv"),
        media_type="text/csv",
        headers={"Content-Disposition": 'attachment; filename="errors.csv"'},
    )


@app.get("/api/export/{entity}")
def export(
    entity: str, format: str = "csv", start: date | None = None, end: date | None = None,
    company_id: int | None = None, u=Depends(current_user), db=Depends(get_db)
):
    if format not in ("csv", "xlsx"):
        raise HTTPException(422, "Формат csv или xlsx")
    aliases = {"bank": "cash", "reports": "sales", "b2b": "sales", "stock": "movements"}
    resolved = aliases.get(entity, entity)
    if start and end and end < start:
        raise HTTPException(422, "Дата окончания раньше даты начала")
    q = access_query(db, u, resolved, company_id=company_id)
    model = MODELS[resolved][0]
    if start and hasattr(model, "date"):
        q = q.where(model.date >= start)
    if end and hasattr(model, "date"):
        q = q.where(model.date <= end)
    if entity == "b2b":
        q = q.where(
            Sale.account_id.in_(select(Account.id).where(Account.channel == "B2B"))
        )
    data = [dump(x) for x in db.scalars(q.order_by(MODELS[resolved][0].id))]
    if len(data) > 50000:
        raise HTTPException(
            413, "Слишком большой экспорт; используйте резервную копию БД"
        )
    for row in data:
        for k, v in list(row.items()):
            if k in importing.MONEY | {
                "revenue",
                "cogs",
                "cost",
                "total",
                "total_cost",
                "remaining_cost",
                "returned_cost",
            } and isinstance(v, int):
                row[k] = str(__import__("decimal").Decimal(v) / 100)
    fields = (
        list(data[0])
        if data
        else [
            c.name
            for c in MODELS[resolved][0].__table__.columns
            if c.name not in ("password_hash", "token_hash", "path", "rows")
        ]
    )
    return Response(
        importing.export_bytes(data, fields, format),
        media_type="application/octet-stream",
        headers={"Content-Disposition": f'attachment; filename="{entity}.{format}"'},
    )


@app.get("/api/templates/{kind}")
def template(kind: str, format: str = "csv", u=Depends(current_user)):
    if kind not in importing.FIELDS:
        raise HTTPException(404, "Не найдено")
    return Response(
        importing.export_bytes([], importing.FIELDS[kind], format),
        media_type="application/octet-stream",
        headers={
            "Content-Disposition": f'attachment; filename="template-{kind}.{format}"'
        },
    )


static = Path(os.getenv("FRONTEND_DIST", "../frontend/dist"))
if static.exists() and (static / "assets").exists():
    app.mount("/assets", StaticFiles(directory=static / "assets"), name="assets")

    @app.head("/")
    def frontend_head():
        return Response(status_code=200)

    @app.get("/{path:path}")
    def frontend(path: str):
        if path.startswith("api/"):
            raise HTTPException(404, "API route not found")
        if path.startswith("public/kaspi-feeds/"):
            raise HTTPException(404, "Прайс-лист не найден")
        return FileResponse(static / "index.html")
