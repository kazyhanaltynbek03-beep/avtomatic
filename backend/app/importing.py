import csv, io, hashlib, json
from decimal import Decimal, InvalidOperation
from openpyxl import load_workbook, Workbook
from fastapi import HTTPException
from sqlalchemy import select
from . import schemas as s, finance as f, canonical
from .models import Product, ImportRun, Sale, Account, MarketplaceProductMapping
from .security import require, audit

KINDS = {
    "opening": ("inventory", s.OpeningInventoryIn),
    "products": ("catalog", s.ProductIn),
    "sales": ("sales", s.SaleIn),
    "returns": ("sales", s.ReturnIn),
    "purchases": ("inventory", s.PurchaseIn),
    "stock": ("inventory", s.AdjustIn),
    "expenses": ("finance", s.ExpenseIn),
    "b2b": ("sales", s.SaleIn),
    "reports": ("sales", s.SaleIn),
    "marketplace_reports": ("sales", dict),
    "bank": ("finance", s.CashIn),
}
MONEY = {
    "price",
    "min_price",
    "recommended_price",
    "unit_price",
    "unit_cost",
    "extra_cost",
    "paid",
    "commission",
    "logistics",
    "advertising",
    "discount",
    "refund",
    "amount",
}
FIELDS = {
    "opening":["company_id","product_id","sku","reference","date","quantity","unit_cost","currency","fx_rate","evidence"],
    "products": ["sku", "name", "category", "article", "brand", "supplier", "country", "media", "currency", "fx_rate", "weight_g", "volume_cm3", "min_price", "recommended_price", "price", "min_stock", "active", "specifications"],
    "sales": [
        "reference",
        "date",
        "company_id",
        "account_id",
        "product_id",
        "sku",
        "customer_id",
        "quantity",
        "unit_price",
        "discount",
        "commission",
        "logistics",
        "advertising",
        "paid",
        "due_date",
    ],
    "returns": ["sale_id", "reference", "date", "quantity", "refund"],
    "purchases": [
        "reference",
        "date",
        "company_id",
        "supplier",
        "product_id",
        "quantity",
        "unit_cost",
        "sku",
        "currency",
        "fx_rate",
        "extra_cost",
        "paid",
    ],
    "stock": ["date", "company_id", "sku", "product_id", "quantity", "unit_cost", "reason"],
    "expenses": [
        "reference",
        "date",
        "company_id",
        "account_id",
        "category",
        "project",
        "payment_method",
        "description",
        "amount",
        "paid",
    ],
    "bank": [
        "reference",
        "date",
        "company_id",
        "sale_id",
        "purchase_id",
        "expense_id",
        "amount",
        "category",
        "description",
    ],
}
FIELDS["b2b"] = FIELDS["reports"] = FIELDS["sales"]
FIELDS["marketplace_reports"] = [
    "event_type", "external_id", "date", "company_id", "account_id", "sku",
    "quantity", "unit_price", "discount", "commission", "logistics", "advertising", "paid",
    "sale_reference", "refund", "stock_delta", "unit_cost", "reason", "amount"
]

_MARKETPLACE_EVENTS = {
    "order": "order", "заказ": "order",
    "sale": "sale", "sales": "sale", "продажа": "sale",
    "return": "return", "refund": "return", "возврат": "return",
    "stock": "stock", "остаток": "stock", "остатки": "stock", "inventory": "stock",
    "payment": "payment", "payout": "payment", "выплата": "payment",
    "commission": "commission", "комиссия": "commission",
    "logistics": "logistics", "логистика": "logistics",
    "return_logistics": "return_logistics", "return logistics": "return_logistics",
    "логистика возврата": "return_logistics", "доставка возврата": "return_logistics",
    "advertising": "advertising", "ads": "advertising", "реклама": "advertising",
}

def marketplace_event(value):
    key = str(value or "").strip().lower()
    if key not in _MARKETPLACE_EVENTS:
        raise HTTPException(422, "event_type: order / sale / payment / return / commission / logistics / return_logistics / advertising / stock")
    return _MARKETPLACE_EVENTS[key]

def _mapped(row, mapping):
    return {dest: row.get(source) for dest, source in mapping.items() if row.get(source) not in ("", None)}

def _money(value):
    return int((Decimal(str(value or 0).replace("\u00a0", "").replace("\u202f", "").replace(" ", "").replace(",", ".")) * 100).quantize(Decimal("1")))

def normalize_marketplace(db, user, row, mapping):
    r = _mapped(row, mapping)
    event = marketplace_event(r.get("event_type"))
    try:
        company_id, account_id = int(r["company_id"]), int(r["account_id"])
    except (KeyError, ValueError):
        raise HTTPException(422, "Нужны company_id и account_id")
    account = db.get(Account, account_id)
    if not account or account.company_id != company_id or account.channel not in ("Kaspi", "WB", "Ozon", "Wildberries"):
        raise HTTPException(422, "Неверный marketplace account")
    sku = str(r.get("sku") or "").strip()
    product = None
    if event in {"sale", "return", "stock"}:
        product = db.scalar(select(Product).where(Product.sku == sku))
        if not product:
            raise HTTPException(422, f"SKU не найден: {sku}")
    external = str(r.get("external_id") or "").strip()
    if not external:
        raise HTTPException(422, "Нужен external_id")
    dt = _date(r.get("date"))
    if event == "order":
        # A marketplace order is not a revenue or cash event. It is retained in
        # the import run/audit only until a confirmed SALE is supplied.
        return event, {"external_id": external, "account_id": account_id, "company_id": company_id}
    if event == "sale":
        payload = {
            "company_id": company_id, "account_id": account_id, "product_id": product.id,
            "reference": f"{account.channel}:{account_id}:{external}", "date": dt,
            "quantity": int(r.get("quantity") or 0), "unit_price": _money(r.get("unit_price")),
            "discount": _money(r.get("discount")),
            # Costs and cash are separate source events; never infer them from a SALE.
            "commission": 0, "logistics": 0, "advertising": 0, "paid": 0,
        }
        return event, s.SaleIn.model_validate(payload)
    if event == "return":
        sale_ref = str(r.get("sale_reference") or "").strip()
        sale = db.scalar(select(Sale).where(Sale.reference == sale_ref))
        if not sale or sale.company_id != company_id or sale.account_id != account_id or sale.product_id != product.id:
            raise HTTPException(422, "Исходная продажа для возврата не найдена")
        return event, s.ReturnIn.model_validate({
            "sale_id": sale.id, "reference": f"{account.channel}:return:{account_id}:{external}",
            "date": dt, "quantity": int(r.get("quantity") or 0), "refund": _money(r.get("refund")),
        })
    if event == "payment":
        sale_ref = str(r.get("sale_reference") or "").strip()
        sale = db.scalar(select(Sale).where(Sale.reference == sale_ref))
        if not sale or sale.company_id != company_id or sale.account_id != account_id:
            raise HTTPException(422, "Продажа для выплаты не найдена")
        payment_amount = _money(r.get("amount") or r.get("paid"))
        if payment_amount <= 0:
            raise HTTPException(422, "payment: нужна положительная сумма")
        return event, s.CashIn.model_validate({"company_id": company_id, "sale_id": sale.id,
            "reference": f"{account.channel}:payment:{account_id}:{external}", "date": dt,
            "amount": payment_amount, "category": "Marketplace payout"})
    if event in {"commission", "logistics", "return_logistics", "advertising"}:
        amount = _money(r.get("amount") or r.get(event))
        if amount <= 0:
            raise HTTPException(422, f"{event}: нужна положительная сумма")
        labels = {"commission": "Marketplace commission", "logistics": "Marketplace logistics", "return_logistics": "Marketplace return logistics", "advertising": "Реклама"}
        return event, s.ExpenseIn.model_validate({"company_id": company_id, "account_id": account_id,
            "reference": f"{account.channel}:{event}:{account_id}:{external}", "date": dt,
            "category": labels[event], "description": str(r.get("sale_reference") or external), "amount": amount, "paid": False})
    delta = int(r.get("stock_delta") or 0)
    if delta == 0:
        raise HTTPException(422, "stock_delta не может быть 0")
    return event, s.AdjustIn.model_validate({
        "company_id": company_id, "product_id": product.id, "date": dt, "quantity": delta,
        "unit_cost": _money(r.get("unit_cost")), "kind": "adjustment",
        "reason": str(r.get("reason") or f"Marketplace stock {account.channel} {external}"),
    })



HEADER_WORDS = set("sku name date quantity amount reference company_id account_id product_id unit_price unit_cost price event_type external_id артикул наименование дата количество сумма цена товар өнім саны күн баға сома".split())


def parse_file(raw, name, sheet=None, header_row=None, metadata=None):
    if len(raw) > 5 * 1024 * 1024:
        raise HTTPException(413, "Максимум 5 МБ")
    matrix = []
    if name.lower().endswith('.xlsx'):
        import zipfile
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            if sum(i.file_size for i in z.infolist()) > 40 * 1024 * 1024:
                raise HTTPException(413, "Слишком большой распакованный XLSX")
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
        try:
            selected = sheet or wb.active.title
            if selected not in wb.sheetnames:
                raise HTTPException(422, 'Лист не найден')
            if metadata is not None:
                metadata.update(sheets=wb.sheetnames, selected_sheet=selected)
            for idx, cells in enumerate(wb[selected].iter_rows(), 1):
                if idx > 5020:
                    raise HTTPException(413, 'Максимум 5000 строк данных и 20 строк заголовка')
                values = []
                for cell in cells:
                    if cell.data_type == 'f':
                        raise HTTPException(422, f'XLSX содержит формулу в {cell.coordinate}; вставьте рассчитанные значения')
                    value = cell.value
                    values.append(str(value.date() if hasattr(value, 'date') else value) if value is not None else '')
                matrix.append(values)
        finally:
            wb.close()
    elif name.lower().endswith('.csv'):
        text = raw.decode('utf-8-sig')
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t')
        except csv.Error:
            dialect = csv.excel
            # Stable fallback for title/preamble lines that defeat Sniffer.
            delimiter = ';' if text.count(';') > text.count(',') else ','
            dialect = type('ImportDialect', (csv.excel,), {'delimiter':delimiter})
        for i, row in enumerate(csv.reader(io.StringIO(text), dialect=dialect)):
            if i >= 5020:
                raise HTTPException(413, 'Максимум 5000 строк данных и 20 строк заголовка')
            matrix.append(row)
    else:
        raise HTTPException(422, 'Поддерживаются UTF-8 CSV и XLSX')
    nonempty = [(i,row) for i,row in enumerate(matrix[:20]) if any(str(x).strip() for x in row)]
    if not nonempty:
        raise HTTPException(422, 'Нет заголовков')
    if header_row is not None:
        index = int(header_row)-1
        if index < 0 or index >= min(20,len(matrix)):
            raise HTTPException(422, 'Номер строки заголовка: 1–20')
    else:
        scored = [(sum(str(x).strip().lower() in HEADER_WORDS for x in row), i) for i,row in nonempty]
        best = max(score for score,i in scored)
        candidates = [i for score,i in scored if score == best]
        index = candidates[0] if best >= 2 else nonempty[0][0]
    headers = [str(x).strip() for x in matrix[index]]
    while headers and not headers[-1]:
        headers.pop()
    if not headers or any(not h for h in headers) or len(headers) != len(set(headers)):
        raise HTTPException(422, 'Пустые или повторяющиеся заголовки; укажите строку заголовков')
    rows = []
    for row in matrix[index+1:]:
        if not any(str(x).strip() for x in row):
            continue
        if len(row) > len(headers) and any(str(x).strip() for x in row[len(headers):]):
            raise HTTPException(422, 'Строка содержит данные вне заголовков')
        if len(rows) >= 5000:
            raise HTTPException(413, 'Максимум 5000 строк')
        rows.append(dict(zip(headers, list(row[:len(headers)]) + [''] * max(0,len(headers)-len(row)))))
    if metadata is not None:
        metadata['header_row'] = index+1
    return headers, rows


def _date(value):
    from datetime import datetime, date
    if isinstance(value, date): return value.isoformat()
    text = str(value).strip()
    for fmt in ('%d.%m.%Y', '%d/%m/%Y', '%Y/%m/%d', '%Y-%m-%d', '%Y-%m-%d %H:%M:%S'):
        try: return datetime.strptime(text, fmt).date().isoformat()
        except ValueError: pass
    return text


def normalize(kind, row, mapping):
    result = {
        dest: row.get(source)
        for dest, source in mapping.items()
        if row.get(source) not in ("", None)
    }
    for key in ("date", "due_date"):
        if result.get(key): result[key] = _date(result[key])
    for k, v in list(result.items()):
        if k in MONEY and not (kind == "expenses" and k == "paid"):
            result[k] = int(
                (Decimal(str(v).replace("\u00a0", "").replace("\u202f", "").replace(" ", "").replace(",", ".")) * 100).quantize(
                    Decimal("1")
                )
            )
    if kind == "products":
        # Spreadsheet catalog rows may carry human-entered media lists and JSON
        # specifications. Normalize these before strict schema validation.
        if "fx_rate" in result:
            result["fx_rate"] = int((Decimal(str(result["fx_rate"]).replace(",", ".")) * 10000).quantize(Decimal("1")))
        if "media" in result and isinstance(result["media"], str):
            raw_media = result["media"].strip()
            if raw_media.startswith("["):
                try:
                    parsed_media = json.loads(raw_media)
                except (ValueError, TypeError):
                    raise HTTPException(422, "media: укажите список URL через ; или корректный JSON-массив")
                if not isinstance(parsed_media, list):
                    raise HTTPException(422, "media: ожидается список URL")
                result["media"] = [str(item).strip() for item in parsed_media if str(item).strip()]
            else:
                result["media"] = [item.strip() for item in raw_media.split(";") if item.strip()]
        if "specifications" in result and isinstance(result["specifications"], str):
            try:
                result["specifications"] = json.loads(result["specifications"] or "{}")
            except (ValueError, TypeError):
                raise HTTPException(422, "specifications: укажите корректный JSON")
    if kind == 'opening' and not result.get('evidence'):
        # The import run itself is auditable evidence; do not reject otherwise
        # valid opening inventory merely because a supplier sheet has no evidence column.
        result['evidence'] = 'Excel/CSV opening inventory import'
    if kind in {'purchases','opening'} and 'fx_rate' in result:
        result['fx_rate']=int((Decimal(str(result['fx_rate']).replace(',','.'))*10000).quantize(Decimal('1')))
    if kind == "purchases":
        line = {k: result.pop(k) for k in ("product_id", "quantity", "unit_cost")}
        for k in ("currency", "fx_rate"):
            if k in result: line[k] = result.pop(k)
        result["lines"] = [line]
    return KINDS[kind][1].model_validate(result)


def apply(db, user, kind, p):
    require(user, KINDS[kind][0], True)
    if kind == "marketplace_reports":
        event, payload = p
        if event == "order":
            return payload
        kind = 'expense' if event in {'commission','logistics','advertising'} else event
        return canonical.post_event(db, user, kind, payload, "stock-report:" + str(payload.company_id) + ":" + payload.reason if kind == "stock" else None)
    if kind == "products":
        values = p.model_dump(exclude={"version", "description", "barcode", "purchase_cost", "custom_fields", "marketplace_stock_allocations"})
        specs = dict(values.get("specifications") or {})
        specs.update({"description": p.description, "barcode": p.barcode,
                      "purchase_cost": p.purchase_cost, "custom_fields": p.custom_fields})
        values["specifications"] = specs
        obj = Product(**values)
        db.add(obj)
        db.flush()
        audit(db, user, "import", "products", obj.id)
        return obj
    if kind == "b2b":
        a = f.check_account(db, p.company_id, p.account_id)
        if a.channel != "B2B":
            f.fail("Ожидается аккаунт B2B")
    event = {"opening":"opening", "sales":"sale", "returns":"return", "b2b":"sale", "reports":"sale",
             "purchases":"purchase", "stock":"stock", "expenses":"expense", "bank":"payment"}[kind]
    identity = None
    if event == 'stock':
        identity = 'stock-import:' + hashlib.sha256(json.dumps(p.model_dump(mode='json'), sort_keys=True).encode()).hexdigest()
    return canonical.post_event(db, user, event, p, identity)


def _resolve_cost_product(db, sku, company_id):
    """Resolve a supplier Excel identifier to the exact product sold by a marketplace.

    Supplier files commonly call the WB seller article an SKU, while the API can
    keep it in ``article`` or in a marketplace mapping.  Treat all three as the
    same identity, but never choose arbitrarily if a value identifies two goods.
    """
    value = str(sku or "").strip()
    if not value:
        return None
    products = {
        product.id
        for product in db.scalars(
            select(Product).where((Product.sku == value) | (Product.article == value))
        )
    }
    mapping_keys = (f"sku:{value}", f"product:{value}", f"barcode:{value}")
    products.update(
        db.scalars(
            select(MarketplaceProductMapping.product_id).where(
                MarketplaceProductMapping.company_id == company_id,
                MarketplaceProductMapping.external_key.in_(mapping_keys),
            )
        )
    )
    if len(products) > 1:
        raise HTTPException(422, f"Артикул/SKU '{value}' сопоставлен с несколькими товарами; выберите точный SKU")
    return next(iter(products), None)


def prepare_rows(db, user, kind, rows, mapping, defaults=None):
    """Resolve human-friendly spreadsheet fields to ERP IDs without guessing financial values."""
    defaults = defaults or {}
    rows = [dict(r) for r in rows]
    mapping = dict(mapping)
    if kind in ("sales", "b2b", "reports", "purchases", "opening", "stock", "marketplace_reports", "expenses", "bank"):
        if "company_id" not in mapping and defaults.get("company_id"):
            for r in rows: r["__qor_company_id"] = str(defaults["company_id"])
            mapping["company_id"] = "__qor_company_id"
        if kind in ("sales", "b2b", "reports", "marketplace_reports", "expenses") and "account_id" not in mapping and defaults.get("account_id"):
            for r in rows: r["__qor_account_id"] = str(defaults["account_id"])
            mapping["account_id"] = "__qor_account_id"
        if kind in ("sales", "b2b", "reports", "purchases", "opening", "stock") and "product_id" not in mapping and mapping.get("sku"):
            sku_col = mapping.pop("sku")
            for r in rows:
                sku = str(r.get(sku_col) or "").strip()
                # Cost files must join the product already created/mapped from
                # Wildberries; do not create a second product with zero history.
                product_id = _resolve_cost_product(db, sku, int(defaults.get("company_id") or 0)) if sku else None
                r["__qor_product_id"] = str(product_id or "")
            mapping["product_id"] = "__qor_product_id"
    return rows, mapping

def preview(db, user, kind, rows, mapping, defaults=None):
    if kind not in KINDS:
        raise HTTPException(422, "Неизвестный тип импорта")
    require(user, KINDS[kind][0], True)
    rows, mapping = prepare_rows(db, user, kind, rows, mapping, defaults)
    if not rows:
        raise HTTPException(422, "Нет строк для импорта")
    fp = hashlib.sha256(
        json.dumps([kind, rows, mapping], sort_keys=True, ensure_ascii=False).encode()
    ).hexdigest()
    previous = db.scalar(
        select(ImportRun).where(
            ImportRun.fingerprint == fp, ImportRun.status == "committed"
        )
    )
    if previous:
        raise HTTPException(409, "Этот набор данных уже импортирован")
    errors = []
    normalized = []
    save = db.begin_nested()
    try:
        for i, row in enumerate(rows):
            try:
                with db.begin_nested():
                    p = normalize_marketplace(db, user, row, mapping) if kind == "marketplace_reports" else normalize(kind, row, mapping)
                    apply(db, user, kind, p)
                    db.flush()
                    if kind == "marketplace_reports":
                        body = p[1] if isinstance(p[1], dict) else p[1].model_dump(mode="json")
                        normalized.append({"event_type": p[0], **body})
                    else:
                        normalized.append(p.model_dump(mode="json"))
            except Exception as e:
                msg = getattr(e, "detail", None) or (
                    "Дубликат или ссылка на отсутствующую запись"
                    if "IntegrityError" in type(e).__name__
                    else str(e)
                )
                errors.append({"row": i + 2, "error": str(msg)[:400]})
    finally:
        save.rollback()
    run = ImportRun(
        user_id=user.id,
        kind=kind,
        fingerprint=fp,
        mapping=mapping,
        rows=rows,
        errors=errors,
    )
    db.add(run)
    db.flush()
    return run, normalized


def confirm(db, user, run):
    if run.user_id != user.id:
        raise HTTPException(403, "Чужой импорт")
    if run.status != "preview" or run.errors:
        raise HTTPException(409, "Импорт уже завершён или содержит ошибки")
    if db.scalar(
        select(ImportRun).where(
            ImportRun.fingerprint == run.fingerprint, ImportRun.status == "committed"
        )
    ):
        raise HTTPException(409, "Данные уже импортированы")
    for row in run.rows:
        p = normalize_marketplace(db, user, row, run.mapping) if run.kind == "marketplace_reports" else normalize(run.kind, row, run.mapping)
        apply(db, user, run.kind, p)
    run.status = "committed"
    audit(db, user, "commit", "imports", run.id, {"rows": len(run.rows)})
    return len(run.rows)


def export_bytes(rows, fields, fmt):
    # Prevent spreadsheet formula injection for all untrusted strings.
    def safe(x):
        if x is None:
            return ""
        if isinstance(x, (dict, list)):
            x = json.dumps(x, ensure_ascii=False)
        if isinstance(x, str) and x.lstrip().startswith(("=", "+", "-", "@")):
            return "'" + x
        return x

    matrix = [[safe(r.get(k)) for k in fields] for r in rows]
    if fmt == "xlsx":
        w = Workbook()
        ws = w.active
        ws.append(fields)
        for row in matrix:
            ws.append(row)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 22
        out = io.BytesIO()
        w.save(out)
        return out.getvalue()
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(fields)
    writer.writerows(matrix)
    return ("\ufeff" + out.getvalue()).encode()
